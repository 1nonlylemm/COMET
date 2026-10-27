from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
import json
import hashlib
from pathlib import Path

from sqlalchemy import false

from app.extensions import db
from app.models import (
    ImportBatch,
    ImportRow,
    LguProfile,
    ProgramCycle,
    Project,
    ProjectRequirement,
    WorkbookCell,
    WorkbookCellChange,
    WorkbookColumn,
    WorkbookConflict,
    WorkbookRow,
    WorkbookSheet,
    WorkbookWorkspace,
)
from app.services import access_service, audit_service, project_service, settings_service
from app.services.import_field_catalog import normalize_label
from app.services.import_mapping_service import normalize_requirement_status
from app.services.requirement_matching_service import find_requirement_match
from app.services.workbook_profile_service import WorkbookProfileError, extract_rows


class WorkbookWorkspaceError(Exception):
    def __init__(self, message, status_code=400):
        self.message = message
        self.status_code = status_code


READ_ONLY_BINDINGS = {
    'import_only',
    'calculated_only',
    'project.source_project_code',
    'lgu.province',
    'lgu.name',
    'lgu.municipality',
    'lgu.level',
    'program_cycle.fiscal_year',
}
EDITABLE_PROJECT_FIELDS = {
    'project.project_name',
    'project.budget_amount',
    'project.overall_status',
    'project.physical_completion_pct',
}
DISPLAY_FORMATS = {'auto', 'money', 'year', 'number'}
COLUMN_DATA_TYPES = {
    'text',
    'number',
    'currency',
    'percent',
    'date',
    'datetime',
    'boolean',
    'status',
    'formula',
    'json',
}
FORMULA_POLICIES = {'manual', 'preserve_only', 'system_calculated', 'manual_override_allowed'}


def create_or_update_workspace_from_batch(batch_id, user_id=None, commit=True):
    batch = db.session.get(ImportBatch, batch_id)
    if not batch:
        raise WorkbookWorkspaceError('Import batch not found.', 404)
    if not batch.program_cycle_id:
        return None

    workbook_reader = _WorkbookReader(batch.stored_file_path, batch.file_type)
    import_rows = _import_rows_by_source(batch.id)
    extracted_rows = _extracted_rows_by_sheet(batch)
    mappings = _field_mappings_by_source(batch.mapping_preview or {})

    existing = WorkbookWorkspace.query.filter_by(
        program_cycle_id=batch.program_cycle_id,
        status='active',
    ).first()
    if existing:
        _capture_reimport_conflicts(existing, batch)
        existing.source_import_batch_id = batch.id
        existing.active_import_template_id = batch.import_template_id
        existing.schema_signature = (batch.workbook_profile or {}).get('schema_signature')
        _sync_workspace_from_batch(existing, batch, workbook_reader, import_rows, extracted_rows, mappings)
        if commit:
            db.session.commit()
        else:
            db.session.flush()
        return existing

    workspace = WorkbookWorkspace(
        program_id=batch.program_id,
        program_cycle_id=batch.program_cycle_id,
        source_import_batch_id=batch.id,
        active_import_template_id=batch.import_template_id,
        name=_workspace_name(batch),
        status='active',
        schema_signature=(batch.workbook_profile or {}).get('schema_signature'),
        created_by=user_id,
    )
    db.session.add(workspace)
    db.session.flush()

    for sheet_index, sheet_profile in enumerate((batch.workbook_profile or {}).get('sheets', []), start=1):
        sheet = WorkbookSheet(
            workspace_id=workspace.id,
            source_sheet_name=sheet_profile.get('name') or f'Sheet {sheet_index}',
            display_name=sheet_profile.get('name') or f'Sheet {sheet_index}',
            sheet_type=sheet_profile.get('sheet_type'),
            sheet_order=sheet_index,
            max_row=sheet_profile.get('max_row'),
            max_column=sheet_profile.get('max_column'),
            merged_ranges=sheet_profile.get('merged_ranges') or [],
        )
        db.session.add(sheet)
        db.session.flush()

        columns = _create_columns(sheet, sheet_profile, mappings, workbook_reader)
        _create_rows_and_cells(sheet, sheet_profile, columns, batch, workbook_reader, import_rows, extracted_rows)

    if commit:
        db.session.commit()
    else:
        db.session.flush()
    return workspace


def list_workspaces(program_cycle_id=None, program_id=None, fiscal_year=None, user_id=None):
    query = WorkbookWorkspace.query
    if program_cycle_id is not None:
        query = query.filter_by(program_cycle_id=program_cycle_id)
    if program_id is not None:
        query = query.filter_by(program_id=program_id)
    if fiscal_year is not None:
        query = query.filter(WorkbookWorkspace.program_cycle.has(ProgramCycle.fiscal_year == fiscal_year))
    workspaces = query.order_by(WorkbookWorkspace.updated_at.desc(), WorkbookWorkspace.id.desc()).all()
    return [_serialize_workspace(workspace) for workspace in workspaces]


def list_sheets(workspace_id, user_id=None):
    workspace = _get_workspace(workspace_id)
    return [_serialize_sheet(sheet) for sheet in workspace.sheets.all()]


def get_sheet_grid(workspace_id, sheet_id, user_id=None, page=1, per_page=100):
    workspace = _get_workspace(workspace_id)
    sheet = WorkbookSheet.query.filter_by(id=sheet_id, workspace_id=workspace.id).first()
    if not sheet:
        raise WorkbookWorkspaceError('Workbook sheet not found.', 404)

    query = sheet.rows.filter(WorkbookRow.row_state == 'active')
    query = _restrict_row_query(query, user_id)
    pagination = query.order_by(WorkbookRow.source_row_number.asc()).paginate(
        page=page,
        per_page=min(per_page, 300),
        error_out=False,
    )
    columns = sheet.columns.order_by(WorkbookColumn.column_index.asc()).all()
    _ensure_editable_cells(pagination.items, columns)
    return {
        'workspace': _serialize_workspace(workspace),
        'sheet': _serialize_sheet(sheet),
        'columns': [_serialize_column(column) for column in columns],
        'rows': [_serialize_row(row, columns) for row in pagination.items],
        'page': pagination.page,
        'per_page': pagination.per_page,
        'total': pagination.total,
        'pages': pagination.pages,
    }


def list_conflicts(workspace_id, user_id=None, status='open', page=1, per_page=50):
    workspace = _get_workspace(workspace_id)
    query = (
        WorkbookConflict.query
        .join(WorkbookCell, WorkbookConflict.cell_id == WorkbookCell.id)
        .join(WorkbookRow, WorkbookCell.row_id == WorkbookRow.id)
        .join(WorkbookSheet, WorkbookRow.sheet_id == WorkbookSheet.id)
        .filter(WorkbookSheet.workspace_id == workspace.id)
    )
    if status and status != 'all':
        query = query.filter(WorkbookConflict.status == status)
    query = _restrict_conflict_query(query, user_id)
    pagination = query.order_by(WorkbookConflict.created_at.desc(), WorkbookConflict.id.desc()).paginate(
        page=max(page, 1),
        per_page=min(max(per_page, 1), 100),
        error_out=False,
    )
    return {
        'items': [_serialize_conflict(conflict) for conflict in pagination.items],
        'page': pagination.page,
        'per_page': pagination.per_page,
        'total': pagination.total,
        'pages': pagination.pages,
    }


def resolve_conflict(workspace_id, conflict_id, decision, user_id=None, expected_revision=None):
    try:
        return _resolve_conflict(workspace_id, conflict_id, decision, user_id, expected_revision)
    except Exception:
        db.session.rollback()
        raise


def _resolve_conflict(workspace_id, conflict_id, decision, user_id=None, expected_revision=None):
    workspace = _get_workspace(workspace_id)
    conflict = (
        WorkbookConflict.query
        .join(WorkbookCell, WorkbookConflict.cell_id == WorkbookCell.id)
        .join(WorkbookRow, WorkbookCell.row_id == WorkbookRow.id)
        .join(WorkbookSheet, WorkbookRow.sheet_id == WorkbookSheet.id)
        .filter(WorkbookConflict.id == conflict_id, WorkbookSheet.workspace_id == workspace.id)
        .with_for_update(of=WorkbookConflict)
        .populate_existing()
        .first()
    )
    if not conflict:
        raise WorkbookWorkspaceError('Workbook conflict not found.', 404)
    if conflict.status != 'open':
        raise WorkbookWorkspaceError('This workbook conflict is already resolved.', 409)

    decision = normalize_label(decision or '').replace(' ', '_')
    if decision not in {'keep_comet', 'accept_import'}:
        raise WorkbookWorkspaceError('Conflict decision must be keep_comet or accept_import.', 422)

    cell = _get_workspace_cell(workspace_id, conflict.cell_id, for_update=True)
    if expected_revision is not None and cell_revision(cell, conflict) != expected_revision:
        raise WorkbookWorkspaceError('The workbook conflict changed after review. Inspect it and prepare a new confirmation.', 409)
    row = cell.row
    if row.project_id and user_id is not None:
        try:
            access_service.ensure_project_access(row.project_id, user_id)
        except access_service.AccessError as exc:
            raise WorkbookWorkspaceError(exc.message, exc.status_code) from exc

    old_value = cell.current_value
    imported_value = _incoming_conflict_value(conflict, strict=True) if decision == 'accept_import' else None
    if decision == 'accept_import':
        coerced = _coerce_cell_value(imported_value, cell.column.data_type)
        _apply_bound_update(cell, coerced, user_id=user_id)
        cell.current_value = coerced
        cell.typed_value = coerced
        cell.display_value = _display_value(coerced)
        cell.source_value = coerced
        cell.raw_value = coerced
        cell.is_dirty = False
        cell.last_import_batch_id = conflict.import_batch_id
        cell.last_edited_by = user_id
        cell.last_edited_at = datetime.now(timezone.utc)
    else:
        coerced = cell.current_value
        # Re-import may have refreshed the bound project before the dirty
        # workbook cell was preserved. Keeping COMET must preserve both.
        target = cell.column.binding_target or ''
        if target in EDITABLE_PROJECT_FIELDS or target.startswith('project.custom_fields.'):
            _apply_bound_update(cell, coerced, user_id=user_id)

    conflict.status = decision
    conflict.resolved_by = user_id
    conflict.resolved_at = datetime.now(timezone.utc)
    change = WorkbookCellChange(
        cell_id=cell.id,
        workspace_id=workspace.id,
        project_id=row.project_id,
        binding_target=cell.column.binding_target,
        old_value=old_value,
        new_value=coerced,
        change_source='conflict_resolution',
        changed_by=user_id,
        reason='Accepted imported value' if decision == 'accept_import' else 'Kept COMET-edited value',
    )
    try:
        db.session.add(change)
        if user_id is not None:
            audit_service.log_action(user_id=user_id, action='resolve_conflict', entity_type='workbook_conflict',
                entity_id=conflict.id, old_value={'value': old_value, 'status': 'open'},
                new_value={'value': coerced, 'decision': decision, 'cell_id': cell.id}, commit=False)
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return _serialize_conflict(conflict)


def list_cell_history(workspace_id, cell_id, user_id=None, page=1, per_page=20):
    cell = _get_workspace_cell(workspace_id, cell_id)
    row = cell.row
    if row.project_id and user_id is not None:
        try:
            access_service.ensure_project_access(row.project_id, user_id)
        except access_service.AccessError as exc:
            raise WorkbookWorkspaceError(exc.message, exc.status_code) from exc

    pagination = cell.changes.order_by(WorkbookCellChange.changed_at.desc(), WorkbookCellChange.id.desc()).paginate(
        page=max(page, 1),
        per_page=min(max(per_page, 1), 100),
        error_out=False,
    )
    return {
        'items': [_serialize_cell_change(change) for change in pagination.items],
        'page': pagination.page,
        'per_page': pagination.per_page,
        'total': pagination.total,
        'pages': pagination.pages,
    }


def export_workspace(workspace_id, user_id=None):
    workspace = _get_workspace(workspace_id)
    scoped = False
    def visible_rows(query):
        query = _restrict_row_query(query, user_id)
        if scoped:
            query = query.join(Project, WorkbookRow.project_id == Project.id)
            query = access_service.restrict_project_query(query, user_id)
        return query

    if user_id is not None:
        try:
            scope = access_service.get_access_scope(user_id)
        except access_service.AccessError as exc:
            raise WorkbookWorkspaceError(exc.message, exc.status_code) from exc
        if not scope.get('unrestricted', False):
            scoped = True
            has_accessible_rows = visible_rows(
                WorkbookRow.query.join(WorkbookSheet, WorkbookRow.sheet_id == WorkbookSheet.id)
                .filter(WorkbookSheet.workspace_id == workspace.id, WorkbookRow.row_state == 'active'),
            ).first() is not None
            if not has_accessible_rows:
                raise WorkbookWorkspaceError('You do not have access to any records in this workbook workspace.', 403)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise WorkbookWorkspaceError('openpyxl is required to export workbook workspaces.', 500) from exc

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    for sheet in workspace.sheets.order_by(WorkbookSheet.sheet_order.asc()).all():
        rows_query = visible_rows(sheet.rows.filter(WorkbookRow.row_state == 'active'))
        rows = rows_query.order_by(WorkbookRow.source_row_number.asc()).all()
        if scoped and not rows:
            continue
        export_sheet = workbook.create_sheet(_safe_sheet_title(sheet.display_name or sheet.source_sheet_name))
        columns = sheet.columns.order_by(WorkbookColumn.column_index.asc()).all()

        export_sheet.append([column.display_label or column.source_header or column.source_column for column in columns])
        for header_cell in export_sheet[1]:
            header_cell.font = Font(bold=True)
            header_cell.fill = PatternFill('solid', fgColor='D9EAF7')
        export_sheet.freeze_panes = 'A2'

        row_ids = [row.id for row in rows]
        cells_by_row = defaultdict(dict)
        if row_ids:
            for i in range(0, len(row_ids), 1000):
                chunk = row_ids[i:i + 1000]
                cells = WorkbookCell.query.filter(WorkbookCell.row_id.in_(chunk)).all()
                for cell in cells:
                    cells_by_row[cell.row_id][cell.column_id] = cell

        for row in rows:
            cells_by_column = cells_by_row.get(row.id, {})
            export_sheet.append([
                _export_cell_value(cells_by_column.get(column.id))
                for column in columns
            ])

        for index, column in enumerate(columns, start=1):
            export_sheet.column_dimensions[export_sheet.cell(row=1, column=index).column_letter].width = min(
                max(len(str(column.display_label or column.source_header or column.source_column)) + 4, 14),
                42,
            )

    metadata = workbook.create_sheet('_COMET Metadata')
    metadata.append(['Field', 'Value'])
    metadata.append(['Workspace ID', workspace.id])
    metadata.append(['Workspace Name', workspace.name])
    metadata.append(['Program Cycle', workspace.program_cycle.name if workspace.program_cycle else ''])
    metadata.append(['Fiscal Year', workspace.program_cycle.fiscal_year if workspace.program_cycle else ''])
    metadata.append(['Source Import Batch', workspace.source_import_batch_id])
    conflicts = workspace.conflicts.filter_by(status='open')
    if scoped:
        visible_ids = visible_rows(WorkbookRow.query.join(WorkbookSheet)
            .filter(WorkbookSheet.workspace_id == workspace.id)).with_entities(WorkbookRow.id)
        conflicts = conflicts.join(WorkbookCell).filter(WorkbookCell.row_id.in_(visible_ids))
    metadata.append(['Open Conflicts', conflicts.count()])
    metadata.append(['Exported At', datetime.now(timezone.utc).isoformat()])
    metadata.append(['Export Note', f'Formula cells preserve formula text when available; Excel recalculation is not performed by {settings_service.get_app_name()}.'])
    for cell in metadata[1]:
        cell.font = Font(bold=True)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return {
        'content': buffer.getvalue(),
        'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'file_name': _safe_export_filename(workspace),
    }


def update_column_settings(workspace_id, column_id, data=None, is_editable=None, display_format=None,
                           update_display_format=False, user_id=None):
    try:
        return _update_column_settings(workspace_id, column_id, data, is_editable, display_format,
                                       update_display_format, user_id)
    except Exception:
        db.session.rollback()
        raise


def _update_column_settings(workspace_id, column_id, data=None, is_editable=None, display_format=None,
                            update_display_format=False, user_id=None):
    workspace = _get_workspace(workspace_id)
    column = (
        WorkbookColumn.query
        .join(WorkbookSheet, WorkbookColumn.sheet_id == WorkbookSheet.id)
        .filter(WorkbookColumn.id == column_id, WorkbookSheet.workspace_id == workspace.id)
        .first()
    )
    if not column:
        raise WorkbookWorkspaceError('Workbook column not found.', 404)

    before = _serialize_column(column)
    updates = dict(data or {})
    if is_editable is not None and 'is_editable' not in updates:
        updates['is_editable'] = is_editable
    if update_display_format and 'display_format' not in updates:
        updates['display_format'] = display_format

    if 'display_label' in updates:
        column.display_label = str(updates.get('display_label') or column.source_column).strip() or column.source_column
        column.normalized_header = normalize_label(column.display_label)
    if 'canonical_label' in updates:
        column.canonical_label = str(updates.get('canonical_label') or '').strip() or None
    if 'classification' in updates:
        column.classification = str(updates.get('classification') or '').strip() or None
    if 'data_type' in updates:
        column.data_type = _normalize_column_data_type(updates.get('data_type'))
        if column.data_type == 'formula':
            column.is_formula = True
            if not column.formula_policy or column.formula_policy == 'manual':
                column.formula_policy = 'preserve_only'
    if 'number_format' in updates:
        number_format = updates.get('number_format')
        column.number_format = str(number_format).strip() if number_format not in (None, '') else None
    if 'binding_target' in updates:
        binding_target = str(updates.get('binding_target') or 'import_only').strip() or 'import_only'
        column.binding_target = binding_target
    if 'binding_config' in updates:
        if not isinstance(updates.get('binding_config') or {}, dict):
            raise WorkbookWorkspaceError('Column binding config must be an object.', 422)
        column.binding_config = updates.get('binding_config') or {}
    if 'validation_rules' in updates:
        if not isinstance(updates.get('validation_rules') or {}, dict):
            raise WorkbookWorkspaceError('Column validation rules must be an object.', 422)
        column.validation_rules = updates.get('validation_rules') or {}
    if 'validation_options' in updates:
        if not isinstance(updates.get('validation_options') or [], list):
            raise WorkbookWorkspaceError('Column validation options must be a list.', 422)
        column.validation_options = updates.get('validation_options') or []
    if 'is_required' in updates:
        column.is_required = bool(updates.get('is_required'))
    if 'is_formula' in updates:
        column.is_formula = bool(updates.get('is_formula'))
        if column.is_formula and column.data_type != 'formula' and 'data_type' not in updates:
            column.data_type = 'formula'
        if not column.is_formula and column.data_type == 'formula':
            column.data_type = 'text'
    if 'formula_policy' in updates:
        column.formula_policy = _normalize_formula_policy(updates.get('formula_policy'))
    if 'is_editable' in updates:
        column.is_editable = bool(updates.get('is_editable'))
    if 'display_format' in updates:
        _set_column_display_format(column, updates.get('display_format'))

    if column.is_formula and not column.formula_policy:
        column.formula_policy = 'preserve_only'
    column.updated_at = datetime.now(timezone.utc)
    created_cells = 0
    if 'is_editable' in updates and column.is_editable:
        created_cells = _ensure_editable_cells(column.sheet.rows.all(), [column], commit=False)
    if 'data_type' in updates:
        column.cells.update({WorkbookCell.value_type: column.data_type}, synchronize_session=False)
    db.session.flush()
    serialized = _serialize_column(column)
    serialized['created_blank_cells'] = created_cells
    _audit_column_change(user_id, workspace.id, column.id, before, serialized, commit=False)
    db.session.commit()
    return serialized


def update_column_editability(workspace_id, column_id, is_editable, user_id=None):
    return update_column_settings(workspace_id, column_id, is_editable=is_editable, user_id=user_id)


def update_cell(workspace_id, cell_id, value, user_id=None, reason=None, expected_revision=None):
    try:
        return _update_cell(workspace_id, cell_id, value, user_id, reason, expected_revision)
    except Exception:
        db.session.rollback()
        raise


def _update_cell(workspace_id, cell_id, value, user_id=None, reason=None, expected_revision=None):
    workspace = _get_workspace(workspace_id)
    cell = _get_workspace_cell(workspace_id, cell_id, for_update=True)
    _assert_cell_editable(cell)
    if expected_revision is not None and cell_revision(cell) != expected_revision:
        raise WorkbookWorkspaceError('The workbook cell changed after review. Inspect it and prepare a new confirmation.', 409)

    row = cell.row
    if row.project_id:
        try:
            if user_id is not None:
                access_service.ensure_project_access(row.project_id, user_id)
            access_service.assert_project_movement_allowed(row.project_id)
        except access_service.AccessError as exc:
            raise WorkbookWorkspaceError(exc.message, exc.status_code) from exc

    old_value = cell.current_value
    coerced = _coerce_cell_value(value, cell.column.data_type)
    _apply_bound_update(cell, coerced, user_id=user_id)

    cell.current_value = coerced
    cell.typed_value = coerced
    cell.display_value = _display_value(coerced)
    cell.value_type = cell.column.data_type
    cell.is_dirty = True
    cell.last_edited_by = user_id
    cell.last_edited_at = datetime.now(timezone.utc)
    change = WorkbookCellChange(
        cell_id=cell.id,
        workspace_id=workspace.id,
        project_id=row.project_id,
        binding_target=cell.column.binding_target,
        old_value=old_value,
        new_value=coerced,
        change_source='staff_edit',
        changed_by=user_id,
        reason=reason,
    )
    try:
        db.session.add(change)
        if user_id is not None:
            audit_service.log_action(user_id=user_id, action='update', entity_type='workbook_cell',
                entity_id=cell.id, old_value={'value': old_value},
                new_value={'value': coerced, 'reason': reason, 'workspace_id': workspace.id}, commit=False)
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return _serialize_cell(cell)


def _workspace_name(batch):
    if batch.program_cycle:
        return f'{batch.program_cycle.name} Workbook'
    return f'{Path(batch.original_file_name).stem} Workbook'


def _create_columns(sheet, sheet_profile, mappings, workbook_reader):
    columns = {}
    for column_profile in sheet_profile.get('columns') or []:
        source_column = column_profile.get('column')
        mapping = mappings.get((sheet.source_sheet_name, source_column), {})
        values = workbook_reader.column_values(
            sheet.source_sheet_name,
            source_column,
            start_row=sheet_profile.get('data_start_row') or 1,
            max_row=sheet_profile.get('max_row') or 0,
            data_only=True,
        )
        formulas = workbook_reader.column_values(
            sheet.source_sheet_name,
            source_column,
            start_row=sheet_profile.get('data_start_row') or 1,
            max_row=sheet_profile.get('max_row') or 0,
            data_only=False,
        )
        is_formula = any(_is_formula(value) for value in formulas)
        number_format = workbook_reader.first_number_format(
            sheet.source_sheet_name,
            source_column,
            start_row=sheet_profile.get('data_start_row') or 1,
            max_row=sheet_profile.get('max_row') or 0,
        )
        binding_target, binding_config = _binding_for_column(mapping, column_profile, is_formula)
        column = WorkbookColumn(
            sheet_id=sheet.id,
            source_column=source_column,
            column_index=column_profile.get('index') or 0,
            header_path=column_profile.get('header_path') or [],
            source_header=column_profile.get('header') or source_column,
            display_label=column_profile.get('header') or source_column,
            canonical_label=mapping.get('canonical_label') or column_profile.get('header') or source_column,
            normalized_header=column_profile.get('normalized_header') or normalize_label(column_profile.get('header')),
            classification=mapping.get('classification'),
            mapping_decision_id=mapping.get('mapping_decision_id'),
            data_type=_infer_data_type(column_profile, values, number_format, is_formula),
            number_format=number_format,
            binding_target=binding_target,
            binding_config=binding_config,
            is_editable=_is_editable_binding(binding_target, is_formula),
            is_required=binding_target in {'project.project_name', 'lgu.name', 'lgu.municipality'},
            is_formula=is_formula,
            formula_policy='preserve_only' if is_formula else 'manual',
        )
        db.session.add(column)
        db.session.flush()
        columns[source_column] = column
    return columns


def _create_rows_and_cells(sheet, sheet_profile, columns, batch, workbook_reader, import_rows, extracted_rows):
    data_start = sheet_profile.get('data_start_row') or 1
    max_row = sheet_profile.get('max_row') or data_start
    extracted_by_row = extracted_rows.get(sheet.source_sheet_name, {})

    for row_number in range(data_start, max_row + 1):
        source_values = (extracted_by_row.get(row_number) or {}).get('values') or {}
        formula_values = workbook_reader.row_values(sheet.source_sheet_name, row_number, data_only=False)
        cached_values = workbook_reader.row_values(sheet.source_sheet_name, row_number, data_only=True)
        if not _row_has_values(source_values, formula_values, columns):
            continue

        import_row = import_rows.get((sheet.source_sheet_name, row_number))
        project = db.session.get(Project, import_row.committed_entity_id) if import_row and import_row.committed_entity_id else None
        workbook_row = WorkbookRow(
            sheet_id=sheet.id,
            source_row_number=row_number,
            source_row_key=import_row.row_key if import_row else _fallback_row_key(source_values),
            project_id=project.id if project else None,
            lgu_id=project.lgu_id if project else None,
            row_state='active',
            last_import_batch_id=batch.id,
        )
        db.session.add(workbook_row)
        db.session.flush()

        for source_column, column in columns.items():
            formula_raw = formula_values.get(source_column)
            cached = cached_values.get(source_column)
            extracted = source_values.get(source_column)
            raw = formula_raw if _is_formula(formula_raw) else extracted
            if raw in (None, '') and cached in (None, '') and not column.is_editable:
                continue
            current = _json_safe(cached if _is_formula(formula_raw) else extracted)
            formula = formula_raw if _is_formula(formula_raw) else None
            cell = _new_cell(
                workbook_row,
                column,
                raw_value=_json_safe(raw),
                typed_value=current,
                display_value=_display_value(current),
                formula=formula,
                cached_formula_value=_json_safe(cached) if formula else None,
                source_value=current,
                current_value=current,
                last_import_batch_id=batch.id,
            )
            db.session.add(cell)


def _sync_workspace_from_batch(workspace, batch, workbook_reader, import_rows, extracted_rows, mappings):
    for sheet_index, sheet_profile in enumerate((batch.workbook_profile or {}).get('sheets', []), start=1):
        source_sheet_name = sheet_profile.get('name') or f'Sheet {sheet_index}'
        sheet = WorkbookSheet.query.filter_by(
            workspace_id=workspace.id,
            source_sheet_name=source_sheet_name,
        ).first()
        if not sheet:
            sheet = WorkbookSheet(
                workspace_id=workspace.id,
                source_sheet_name=source_sheet_name,
                display_name=source_sheet_name,
            )
            db.session.add(sheet)
            db.session.flush()

        sheet.display_name = source_sheet_name
        sheet.sheet_type = sheet_profile.get('sheet_type')
        sheet.sheet_order = sheet_index
        sheet.max_row = sheet_profile.get('max_row')
        sheet.max_column = sheet_profile.get('max_column')
        sheet.merged_ranges = sheet_profile.get('merged_ranges') or []
        sheet.is_visible = True

        columns = _upsert_columns(sheet, sheet_profile, mappings, workbook_reader)
        _sync_rows_and_cells(sheet, sheet_profile, columns, batch, workbook_reader, import_rows, extracted_rows)


def _upsert_columns(sheet, sheet_profile, mappings, workbook_reader):
    existing_columns = {
        column.source_column: column
        for column in sheet.columns.order_by(WorkbookColumn.column_index.asc()).all()
    }
    columns = {}
    for column_profile in sheet_profile.get('columns') or []:
        source_column = column_profile.get('column')
        mapping = mappings.get((sheet.source_sheet_name, source_column), {})
        values = workbook_reader.column_values(
            sheet.source_sheet_name,
            source_column,
            start_row=sheet_profile.get('data_start_row') or 1,
            max_row=sheet_profile.get('max_row') or 0,
            data_only=True,
        )
        formulas = workbook_reader.column_values(
            sheet.source_sheet_name,
            source_column,
            start_row=sheet_profile.get('data_start_row') or 1,
            max_row=sheet_profile.get('max_row') or 0,
            data_only=False,
        )
        is_formula = any(_is_formula(value) for value in formulas)
        number_format = workbook_reader.first_number_format(
            sheet.source_sheet_name,
            source_column,
            start_row=sheet_profile.get('data_start_row') or 1,
            max_row=sheet_profile.get('max_row') or 0,
        )
        binding_target, binding_config = _binding_for_column(mapping, column_profile, is_formula)
        column = existing_columns.get(source_column)
        if not column:
            column = WorkbookColumn(sheet_id=sheet.id, source_column=source_column)
            db.session.add(column)
            db.session.flush()

        column.column_index = column_profile.get('index') or 0
        column.header_path = column_profile.get('header_path') or []
        column.source_header = column_profile.get('header') or source_column
        column.display_label = column_profile.get('header') or source_column
        column.canonical_label = mapping.get('canonical_label') or column_profile.get('header') or source_column
        column.normalized_header = column_profile.get('normalized_header') or normalize_label(column_profile.get('header'))
        column.classification = mapping.get('classification')
        column.mapping_decision_id = mapping.get('mapping_decision_id')
        column.data_type = _infer_data_type(column_profile, values, number_format, is_formula)
        column.number_format = number_format
        column.binding_target = binding_target
        column.binding_config = binding_config
        column.is_editable = _is_editable_binding(binding_target, is_formula)
        column.is_required = binding_target in {'project.project_name', 'lgu.name', 'lgu.municipality'}
        column.is_formula = is_formula
        column.formula_policy = 'preserve_only' if is_formula else 'manual'
        columns[source_column] = column
    return columns


def _sync_rows_and_cells(sheet, sheet_profile, columns, batch, workbook_reader, import_rows, extracted_rows):
    data_start = sheet_profile.get('data_start_row') or 1
    max_row = sheet_profile.get('max_row') or data_start
    extracted_by_row = extracted_rows.get(sheet.source_sheet_name, {})
    existing_rows = {
        row.source_row_number: row
        for row in sheet.rows.order_by(WorkbookRow.source_row_number.asc()).all()
    }
    active_row_ids = set()

    for row_number in range(data_start, max_row + 1):
        source_values = (extracted_by_row.get(row_number) or {}).get('values') or {}
        formula_values = workbook_reader.row_values(sheet.source_sheet_name, row_number, data_only=False)
        cached_values = workbook_reader.row_values(sheet.source_sheet_name, row_number, data_only=True)
        if not _row_has_values(source_values, formula_values, columns):
            continue

        import_row = import_rows.get((sheet.source_sheet_name, row_number))
        project = db.session.get(Project, import_row.committed_entity_id) if import_row and import_row.committed_entity_id else None
        workbook_row = existing_rows.get(row_number)
        if not workbook_row:
            workbook_row = WorkbookRow(sheet_id=sheet.id, source_row_number=row_number)
            db.session.add(workbook_row)
            db.session.flush()

        workbook_row.source_row_key = import_row.row_key if import_row else _fallback_row_key(source_values)
        workbook_row.project_id = project.id if project else None
        workbook_row.lgu_id = project.lgu_id if project else None
        workbook_row.row_state = 'active'
        workbook_row.last_import_batch_id = batch.id
        active_row_ids.add(workbook_row.id)

        existing_cells = {
            cell.column_id: cell
            for cell in workbook_row.cells.all()
        }
        for source_column, column in columns.items():
            formula_raw = formula_values.get(source_column)
            cached = cached_values.get(source_column)
            extracted = source_values.get(source_column)
            raw = formula_raw if _is_formula(formula_raw) else extracted
            current = _json_safe(cached if _is_formula(formula_raw) else extracted)
            formula = formula_raw if _is_formula(formula_raw) else None
            cell = existing_cells.get(column.id)

            if raw in (None, '') and cached in (None, '') and not column.is_editable and cell is None:
                continue

            if cell is None:
                cell = _new_cell(
                    workbook_row,
                    column,
                    raw_value=_json_safe(raw),
                    typed_value=current,
                    display_value=_display_value(current),
                    formula=formula,
                    cached_formula_value=_json_safe(cached) if formula else None,
                    source_value=current,
                    current_value=current,
                    last_import_batch_id=batch.id,
                )
                db.session.add(cell)
                continue

            if cell.is_dirty:
                continue

            cell.cell_ref = f'{column.source_column}{workbook_row.source_row_number}'
            cell.raw_value = _json_safe(raw)
            cell.typed_value = current
            cell.display_value = _display_value(current)
            cell.formula = formula
            cell.cached_formula_value = _json_safe(cached) if formula else None
            cell.source_value = current
            cell.current_value = current
            cell.value_type = column.data_type
            cell.validation_state = 'valid'
            cell.validation_messages = []
            cell.last_import_batch_id = batch.id

    for workbook_row in existing_rows.values():
        if workbook_row.id not in active_row_ids:
            workbook_row.row_state = 'archived'


def _extracted_rows_by_sheet(batch):
    try:
        rows = extract_rows(batch.stored_file_path, batch.workbook_profile or {})
    except WorkbookProfileError as exc:
        raise WorkbookWorkspaceError(exc.message, exc.status_code) from exc
    grouped = {}
    for row in rows:
        grouped.setdefault(row['sheet_name'], {})[row['row_number']] = row
    return grouped


def _capture_reimport_conflicts(workspace, batch):
    dirty_cells = (
        WorkbookCell.query
        .join(WorkbookRow, WorkbookCell.row_id == WorkbookRow.id)
        .join(WorkbookSheet, WorkbookRow.sheet_id == WorkbookSheet.id)
        .filter(WorkbookSheet.workspace_id == workspace.id, WorkbookCell.is_dirty.is_(True))
        .all()
    )
    for cell in dirty_cells:
        existing = WorkbookConflict.query.filter_by(
            workspace_id=workspace.id,
            cell_id=cell.id,
            import_batch_id=batch.id,
            status='open',
        ).first()
        if existing:
            continue
        db.session.add(WorkbookConflict(
            workspace_id=workspace.id,
            cell_id=cell.id,
            import_batch_id=batch.id,
            source_value=cell.source_value,
            current_value=cell.current_value,
            conflict_type='dirty_cell_reimport',
        ))


def _field_mappings_by_source(mapping_preview):
    return {
        (item.get('source_sheet'), item.get('source_column')): item
        for item in mapping_preview.get('field_mappings', [])
        if item.get('source_sheet') and item.get('source_column')
    }


def _import_rows_by_source(batch_id):
    rows = ImportRow.query.filter_by(import_batch_id=batch_id).all()
    return {(row.sheet_name, row.row_number): row for row in rows}


def _binding_for_column(mapping, column_profile, is_formula):
    if is_formula:
        return 'calculated_only', {}
    target = mapping.get('target_field')
    if target in READ_ONLY_BINDINGS or target in EDITABLE_PROJECT_FIELDS:
        return target, {}
    if target in {'project.location', 'project.target_beneficiaries', 'project.mode_of_implementation', 'project.implementation_period'}:
        return f'project.custom_fields.{target.split(".", 1)[1]}', {}
    if target and target.startswith('monitoring.'):
        return f'project.custom_fields.{target.split(".", 1)[1]}', {'source_target': target}
    if target and target.startswith('project.'):
        return target, {}
    if target and target.startswith(('lgu.', 'program_cycle.')):
        return target, {}
    if mapping.get('suggested_role') == 'requirement':
        return f'project_requirement.status:{_requirement_name(mapping)}', {}
    if mapping.get('suggested_role'):
        return f'project.custom_fields.{_custom_key(mapping)}', {'suggested_role': mapping.get('suggested_role')}
    normalized = column_profile.get('normalized_header') or normalize_label(column_profile.get('header'))
    if normalized in {'remarks', 'remark', 'comments', 'comment'}:
        return 'project.custom_fields.remarks', {}
    return 'import_only', {}


def _is_editable_binding(binding_target, is_formula):
    if is_formula or binding_target in READ_ONLY_BINDINGS:
        return False
    return (
        binding_target in EDITABLE_PROJECT_FIELDS
        or binding_target.startswith('project.custom_fields.')
        or binding_target.startswith('project_requirement.status:')
    )


def _infer_data_type(column_profile, values, number_format, is_formula):
    if is_formula:
        return _numeric_type_from_format(number_format) or 'formula'
    format_type = _numeric_type_from_format(number_format)
    if format_type:
        return format_type
    samples = [value for value in values if value not in (None, '')]
    profile_samples = [value for value in column_profile.get('sample_values') or [] if value not in (None, '')]
    samples = samples or profile_samples
    if not samples:
        return 'text'
    if all(isinstance(value, bool) for value in samples):
        return 'boolean'
    if all(isinstance(value, (date, datetime)) for value in samples):
        return 'date'
    if all(isinstance(value, (int, float, Decimal)) and not isinstance(value, bool) for value in samples):
        return 'number'
    return 'text'


def _numeric_type_from_format(number_format):
    if not number_format:
        return None
    fmt = str(number_format).lower()
    if '%' in fmt:
        return 'percent'
    if '\\u20b1' in fmt.encode('unicode_escape').decode('ascii').lower() or '$' in fmt or 'php' in fmt:
        return 'currency'
    return None


def _row_has_values(source_values, formula_values, columns):
    for source_column in columns:
        if source_values.get(source_column) not in (None, ''):
            return True
        if formula_values.get(source_column) not in (None, ''):
            return True
    return False


def _fallback_row_key(values):
    signal = '|'.join(str(value) for value in values.values() if value not in (None, ''))[:255]
    return signal or None


def _get_workspace(workspace_id):
    workspace = db.session.get(WorkbookWorkspace, workspace_id)
    if not workspace:
        raise WorkbookWorkspaceError('Workbook workspace not found.', 404)
    return workspace


def _get_workspace_cell(workspace_id, cell_id, for_update=False):
    workspace = _get_workspace(workspace_id)
    cell = (
        WorkbookCell.query
        .join(WorkbookRow, WorkbookCell.row_id == WorkbookRow.id)
        .join(WorkbookSheet, WorkbookRow.sheet_id == WorkbookSheet.id)
        .filter(WorkbookCell.id == cell_id, WorkbookSheet.workspace_id == workspace.id)
    )
    if for_update:
        cell = cell.with_for_update(of=WorkbookCell).populate_existing()
    cell = cell.first()
    if not cell:
        raise WorkbookWorkspaceError('Workbook cell not found.', 404)
    if for_update:
        # Re-read and lock the binding and target as well as the cell, so
        # a reviewed edit cannot race a column or project update.
        WorkbookColumn.query.filter_by(id=cell.column_id).with_for_update().populate_existing().one()
        row = WorkbookRow.query.filter_by(id=cell.row_id).with_for_update().populate_existing().one()
        if row.project_id:
            Project.query.filter_by(id=row.project_id).with_for_update().populate_existing().one()
    return cell


def _restrict_row_query(query, user_id):
    if user_id is None:
        return query
    try:
        scope = access_service.get_access_scope(user_id)
    except access_service.AccessError as exc:
        raise WorkbookWorkspaceError(exc.message, exc.status_code) from exc
    if scope['unrestricted']:
        return query
    lgu_ids = scope['lgu_ids']
    if lgu_ids:
        return query.filter(WorkbookRow.lgu_id.in_(lgu_ids))
    cluster_ids = scope['cluster_ids']
    if not cluster_ids:
        return query.filter(false())
    return query.filter(WorkbookRow.lgu.has(LguProfile.cluster_id.in_(cluster_ids)))


def _restrict_conflict_query(query, user_id):
    if user_id is None:
        return query
    try:
        scope = access_service.get_access_scope(user_id)
    except access_service.AccessError as exc:
        raise WorkbookWorkspaceError(exc.message, exc.status_code) from exc
    if scope['unrestricted']:
        return query
    lgu_ids = scope['lgu_ids']
    if lgu_ids:
        return query.filter(WorkbookRow.lgu_id.in_(lgu_ids))
    cluster_ids = scope['cluster_ids']
    if not cluster_ids:
        return query.filter(false())
    return query.filter(WorkbookRow.lgu.has(LguProfile.cluster_id.in_(cluster_ids)))


def cell_revision(cell, conflict=None):
    """Revision for a reviewed edit, including its binding and target state."""
    def values(record):
        return {column.name: getattr(record, column.name) for column in record.__table__.columns} if record else None

    row, column = cell.row, cell.column
    workspace = row.sheet.workspace
    payload = {
        'cell': values(cell), 'column': values(column), 'row': values(row),
        'workspace': values(workspace), 'project': values(row.project),
        'program_active': workspace.program.is_active,
        'cycle_active': workspace.program_cycle.is_active,
        'conflicts': [values(item) for item in WorkbookConflict.query.filter_by(cell_id=cell.id).order_by(WorkbookConflict.id).all()],
    }
    if conflict:
        payload['conflict'] = values(conflict)
        payload['incoming_value'] = _incoming_conflict_value(conflict, strict=True)
        batch = conflict.import_batch
        try:
            with open(batch.stored_file_path, 'rb') as source:
                payload['incoming_file_hash'] = hashlib.file_digest(source, 'sha256').hexdigest()
        except (OSError, AttributeError) as exc:
            raise WorkbookWorkspaceError('The incoming workbook is unavailable. Inspect the conflict again.', 422) from exc
        if payload['incoming_file_hash'] != batch.file_hash:
            raise WorkbookWorkspaceError('The incoming workbook changed. Re-import it before resolving this conflict.', 409)
    return hashlib.sha256(json.dumps(payload, sort_keys=True, default=str, separators=(',', ':')).encode()).hexdigest()


def _assert_cell_editable(cell):
    if not cell.column.is_editable:
        raise WorkbookWorkspaceError('This workbook cell is read-only because it is calculated or source-controlled.', 422)


def _apply_bound_update(cell, value, user_id=None):
    target = cell.column.binding_target
    row = cell.row
    if target.startswith('project.custom_fields.'):
        if not row.project_id:
            raise WorkbookWorkspaceError('This workbook row is not linked to a project.', 422)
        key = target.split('project.custom_fields.', 1)[1]
        project = db.session.get(Project, row.project_id)
        custom_fields = dict(project.custom_fields or {})
        custom_fields[key] = value
        project_service.update_project(
            project.id,
            {'custom_fields': custom_fields},
            user_id=user_id,
            enforce_workbook_locks=False,
            sync_workbook=False,
            commit=False,
        )
        return
    if target in EDITABLE_PROJECT_FIELDS:
        if not row.project_id:
            raise WorkbookWorkspaceError('This workbook row is not linked to a project.', 422)
        field = target.split('.', 1)[1]
        payload = {field: value}
        project_service.update_project(
            row.project_id,
            payload,
            user_id=user_id,
            enforce_workbook_locks=False,
            sync_workbook=False,
            commit=False,
        )
        return
    if target.startswith('project_requirement.status:'):
        if not row.project_id:
            raise WorkbookWorkspaceError('This workbook row is not linked to a project.', 422)
        requirement_name = target.split(':', 1)[1]
        requirement = _find_project_requirement_for_binding(row.project_id, requirement_name)
        if not requirement:
            raise WorkbookWorkspaceError('This requirement is not linked to the workbook row project.', 422)
        project_service.update_project_requirement(
            requirement.id,
            {'status': normalize_requirement_status(value)},
            user_id=user_id,
            enforce_workbook_locks=False,
            sync_workbook=False,
        )
        return
    return


def _find_project_requirement_for_binding(project_id, requirement_name):
    requirement = ProjectRequirement.query.filter_by(
        project_id=project_id,
        requirement_name=requirement_name,
    ).first()
    if requirement:
        return requirement

    project = db.session.get(Project, project_id)
    if not project:
        return None
    canonical = find_requirement_match(project.program_id, requirement_name)
    if canonical:
        requirement = ProjectRequirement.query.filter_by(
            project_id=project_id,
            requirement_id=canonical.id,
        ).first()
        if requirement:
            return requirement

    normalized_name = normalize_label(requirement_name)
    for candidate in ProjectRequirement.query.filter_by(project_id=project_id).all():
        if normalize_label(candidate.requirement_name) == normalized_name:
            return candidate
    return None


def _ensure_editable_cells(rows, columns, commit=True):
    editable_columns = [column for column in columns if column.is_editable]
    if not rows or not editable_columns:
        return 0

    row_ids = [row.id for row in rows]
    column_ids = [column.id for column in editable_columns]
    existing = {
        (cell.row_id, cell.column_id)
        for cell in WorkbookCell.query.filter(
            WorkbookCell.row_id.in_(row_ids),
            WorkbookCell.column_id.in_(column_ids),
        ).all()
    }

    created = 0
    for row in rows:
        for column in editable_columns:
            if (row.id, column.id) in existing:
                continue
            db.session.add(_new_cell(row, column, last_import_batch_id=row.last_import_batch_id))
            created += 1
    if created:
        if commit:
            db.session.commit()
        else:
            db.session.flush()
    return created


def _new_cell(row, column, raw_value=None, typed_value=None, display_value=None, formula=None,
              cached_formula_value=None, source_value=None, current_value=None, last_import_batch_id=None):
    return WorkbookCell(
        row_id=row.id,
        column_id=column.id,
        cell_ref=f'{column.source_column}{row.source_row_number}',
        raw_value=raw_value,
        typed_value=typed_value,
        display_value=display_value,
        formula=formula,
        cached_formula_value=cached_formula_value,
        source_value=source_value,
        current_value=current_value,
        value_type=column.data_type,
        validation_state='valid',
        validation_messages=[],
        last_import_batch_id=last_import_batch_id,
    )


def _coerce_cell_value(value, data_type):
    if value in ('', None):
        return None
    if data_type in {'number', 'currency', 'percent'}:
        try:
            return float(str(value).replace(',', ''))
        except (TypeError, ValueError) as exc:
            raise WorkbookWorkspaceError('Cell value must be numeric.', 422) from exc
    if data_type == 'boolean':
        if isinstance(value, bool):
            return value
        normalized = normalize_label(value)
        if normalized in {'true', 'yes', '1', 'submitted', 'complete', 'completed'}:
            return True
        if normalized in {'false', 'no', '0', 'pending', 'not submitted'}:
            return False
        raise WorkbookWorkspaceError('Cell value must be true or false.', 422)
    return _json_safe(value)


def _serialize_workspace(workspace):
    return {
        'id': workspace.id,
        'program_id': workspace.program_id,
        'program_cycle_id': workspace.program_cycle_id,
        'source_import_batch_id': workspace.source_import_batch_id,
        'active_import_template_id': workspace.active_import_template_id,
        'name': workspace.name,
        'status': workspace.status,
        'schema_signature': workspace.schema_signature,
        'sheet_count': workspace.sheets.count(),
        'conflict_count': workspace.conflicts.filter_by(status='open').count(),
        'created_at': workspace.created_at.isoformat() if workspace.created_at else None,
        'updated_at': workspace.updated_at.isoformat() if workspace.updated_at else None,
    }


def _serialize_sheet(sheet):
    return {
        'id': sheet.id,
        'workspace_id': sheet.workspace_id,
        'source_sheet_name': sheet.source_sheet_name,
        'display_name': sheet.display_name,
        'sheet_type': sheet.sheet_type,
        'sheet_order': sheet.sheet_order,
        'max_row': sheet.max_row,
        'max_column': sheet.max_column,
        'merged_ranges': sheet.merged_ranges or [],
        'row_count': sheet.rows.count(),
        'column_count': sheet.columns.count(),
    }


def _serialize_column(column):
    binding_config = column.binding_config or {}
    return {
        'id': column.id,
        'sheet_id': column.sheet_id,
        'source_column': column.source_column,
        'column_index': column.column_index,
        'header_path': column.header_path or [],
        'source_header': column.source_header,
        'display_label': column.display_label,
        'canonical_label': column.canonical_label,
        'normalized_header': column.normalized_header,
        'classification': column.classification,
        'mapping_decision_id': column.mapping_decision_id,
        'data_type': column.data_type,
        'number_format': column.number_format,
        'validation_rules': column.validation_rules or {},
        'validation_options': column.validation_options or [],
        'binding_target': column.binding_target,
        'binding_config': binding_config,
        'display_format': binding_config.get('display_format') or 'auto',
        'is_editable': column.is_editable,
        'is_required': column.is_required,
        'is_formula': column.is_formula,
        'formula_policy': column.formula_policy,
    }


def _set_column_display_format(column, display_format):
    normalized = normalize_label(display_format or 'auto').replace(' ', '_')
    if normalized in {'normal', 'normal_number'}:
        normalized = 'number'
    if normalized not in DISPLAY_FORMATS:
        raise WorkbookWorkspaceError('Column display format must be auto, money, year, or number.', 422)

    binding_config = dict(column.binding_config or {})
    if normalized == 'auto':
        binding_config.pop('display_format', None)
    else:
        binding_config['display_format'] = normalized
    column.binding_config = binding_config


def _normalize_column_data_type(value):
    normalized = normalize_label(value or 'text').replace(' ', '_')
    if normalized in {'string', 'plain_text'}:
        normalized = 'text'
    if normalized in {'money', 'amount'}:
        normalized = 'currency'
    if normalized not in COLUMN_DATA_TYPES:
        raise WorkbookWorkspaceError(
            'Column data type must be text, number, currency, percent, date, datetime, boolean, status, formula, or json.',
            422,
        )
    return normalized


def _normalize_formula_policy(value):
    normalized = normalize_label(value or 'manual').replace(' ', '_')
    if normalized not in FORMULA_POLICIES:
        raise WorkbookWorkspaceError(
            'Column formula policy must be manual, preserve_only, system_calculated, or manual_override_allowed.',
            422,
        )
    return normalized


def _serialize_row(row, columns):
    cells_by_column = {cell.column_id: cell for cell in row.cells.all()}
    return {
        'id': row.id,
        'sheet_id': row.sheet_id,
        'source_row_number': row.source_row_number,
        'source_row_key': row.source_row_key,
        'project_id': row.project_id,
        'lgu_id': row.lgu_id,
        'row_state': row.row_state,
        'cells': {
            column.source_column: _serialize_cell(cells_by_column.get(column.id), column=column)
            for column in columns
            if cells_by_column.get(column.id) is not None
        },
    }


def _serialize_cell(cell, column=None):
    if cell is None:
        return None
    column = column or cell.column
    return {
        'id': cell.id,
        'row_id': cell.row_id,
        'column_id': cell.column_id,
        'cell_ref': cell.cell_ref,
        'raw_value': cell.raw_value,
        'typed_value': cell.typed_value,
        'display_value': cell.display_value,
        'formula': cell.formula,
        'cached_formula_value': cell.cached_formula_value,
        'source_value': cell.source_value,
        'current_value': cell.current_value,
        'value_type': cell.value_type,
        'validation_state': cell.validation_state,
        'validation_messages': cell.validation_messages or [],
        'is_dirty': cell.is_dirty,
        'is_editable': bool(column.is_editable),
        'last_edited_by': cell.last_edited_by,
        'last_edited_at': cell.last_edited_at.isoformat() if cell.last_edited_at else None,
    }


def _serialize_conflict(conflict):
    cell = conflict.cell
    row = cell.row
    column = cell.column
    return {
        'id': conflict.id,
        'workspace_id': conflict.workspace_id,
        'cell_id': conflict.cell_id,
        'cell_ref': cell.cell_ref,
        'row_id': row.id,
        'source_row_number': row.source_row_number,
        'source_row_key': row.source_row_key,
        'project_id': row.project_id,
        'lgu_id': row.lgu_id,
        'column_id': column.id,
        'source_column': column.source_column,
        'column_label': column.display_label or column.source_header or column.source_column,
        'binding_target': column.binding_target,
        'import_batch_id': conflict.import_batch_id,
        'source_value': conflict.source_value,
        'incoming_value': _incoming_conflict_value(conflict),
        'current_value': conflict.current_value,
        'conflict_type': conflict.conflict_type,
        'status': conflict.status,
        'resolved_by': conflict.resolved_by,
        'resolved_at': conflict.resolved_at.isoformat() if conflict.resolved_at else None,
        'created_at': conflict.created_at.isoformat() if conflict.created_at else None,
    }


def _serialize_cell_change(change):
    actor = change.actor
    return {
        'id': change.id,
        'cell_id': change.cell_id,
        'workspace_id': change.workspace_id,
        'project_id': change.project_id,
        'binding_target': change.binding_target,
        'old_value': change.old_value,
        'new_value': change.new_value,
        'change_source': change.change_source,
        'changed_by': change.changed_by,
        'changed_by_name': actor.full_name if actor else None,
        'changed_at': change.changed_at.isoformat() if change.changed_at else None,
        'reason': change.reason,
    }


def _incoming_conflict_value(conflict, strict=False):
    batch = conflict.import_batch
    cell = conflict.cell
    row = cell.row
    column = cell.column
    if not batch or not row or not column:
        return None
    try:
        workbook_reader = _WorkbookReader(batch.stored_file_path, batch.file_type)
        formula_value = workbook_reader.cell_value(row.sheet.source_sheet_name, column.source_column, row.source_row_number, data_only=False)
        cached_value = workbook_reader.cell_value(row.sheet.source_sheet_name, column.source_column, row.source_row_number, data_only=True)
    except Exception as exc:
        if strict:
            raise WorkbookWorkspaceError('The incoming workbook is unavailable. Re-import or inspect the conflict before accepting it.', 422) from exc
        return None
    if _is_formula(formula_value):
        return _json_safe(cached_value)
    return _json_safe(cached_value if cached_value not in (None, '') else formula_value)


def _audit_column_change(user_id, workspace_id, column_id, old_value, new_value, commit=True):
    if user_id is None:
        return
    audit_service.log_action(
        user_id=user_id,
        action='update_column_editability',
        entity_type='workbook_column',
        entity_id=column_id,
        old_value=old_value,
        new_value={**new_value, 'workspace_id': workspace_id},
        commit=commit,
    )


def _display_value(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    return str(value)


def _export_cell_value(cell):
    if cell is None:
        return None
    if cell.formula:
        return cell.formula
    value = cell.current_value if cell.current_value is not None else cell.typed_value
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and _is_spreadsheet_formula_like(value):
        return f"'{value}"
    return value


def _safe_sheet_title(value):
    title = ''.join('_' if char in r'[]:*?/\\' else char for char in str(value or 'Sheet')).strip() or 'Sheet'
    return title[:31]


def _safe_export_filename(workspace):
    base = normalize_label(workspace.name or f'workbook-{workspace.id}').replace(' ', '-')
    if not base:
        base = f'workbook-{workspace.id}'
    return f'{base}-{datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")}.xlsx'


def _json_safe(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _is_formula(value):
    return isinstance(value, str) and value.startswith('=')


def _is_spreadsheet_formula_like(value):
    if not isinstance(value, str) or not value:
        return False
    return value[0] in ('=', '+', '-', '@', '\t', '\r')


def _requirement_name(mapping_item):
    header = mapping_item.get('source_header') or mapping_item.get('header') or mapping_item.get('source_column')
    if header:
        parts = [part.strip() for part in str(header).split('/') if part.strip()]
        return parts[-1] if parts else str(header).strip()
    return 'Imported Requirement'


def _custom_key(mapping_item):
    return normalize_label(mapping_item.get('source_header') or mapping_item.get('source_column')).replace(' ', '_')


class _WorkbookReader:
    def __init__(self, file_path, file_type):
        self.file_path = file_path
        self.file_type = file_type
        self._formula_workbook = None
        self._value_workbook = None

    def column_values(self, sheet_name, source_column, start_row, max_row, data_only):
        return [
            self.cell_value(sheet_name, source_column, row_number, data_only=data_only)
            for row_number in range(start_row, max_row + 1)
        ]

    def row_values(self, sheet_name, row_number, data_only):
        if self.file_type == 'csv':
            return {}
        worksheet = self._worksheet(sheet_name, data_only=data_only)
        values = {}
        for cell in worksheet[row_number]:
            if cell.value not in (None, ''):
                values[cell.column_letter] = _json_safe(cell.value)
        return values

    def cell_value(self, sheet_name, source_column, row_number, data_only):
        if self.file_type == 'csv':
            return None
        worksheet = self._worksheet(sheet_name, data_only=data_only)
        return _json_safe(worksheet[f'{source_column}{row_number}'].value)

    def first_number_format(self, sheet_name, source_column, start_row, max_row):
        if self.file_type == 'csv':
            return None
        worksheet = self._worksheet(sheet_name, data_only=False)
        for row_number in range(start_row, max_row + 1):
            cell = worksheet[f'{source_column}{row_number}']
            if cell.value not in (None, '') and cell.number_format and cell.number_format != 'General':
                return cell.number_format
        return None

    def _worksheet(self, sheet_name, data_only):
        workbook = self._load_workbook(data_only=data_only)
        return workbook[sheet_name]

    def _load_workbook(self, data_only):
        if self.file_type not in {'xlsx', 'xls'}:
            return None
        from openpyxl import load_workbook

        if data_only:
            if self._value_workbook is None:
                self._value_workbook = load_workbook(self.file_path, data_only=True, read_only=False)
            return self._value_workbook
        if self._formula_workbook is None:
            self._formula_workbook = load_workbook(self.file_path, data_only=False, read_only=False)
        return self._formula_workbook
