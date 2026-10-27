import csv
import io
from pathlib import Path

from flask import current_app
from sqlalchemy import func

from app.extensions import db
from app.models import LguCluster, LguProfile, Program, User
from app.services import auth_service, lgu_service, program_service
from app.services.import_field_catalog import normalize_label
from app.services.upload_validation import ArchiveSecurityError, validate_ooxml_container, _measure_upload_size


class BatchImportError(Exception):
    def __init__(self, message, status_code=400, errors=None):
        self.message = message
        self.status_code = status_code
        self.errors = errors
        super().__init__(message)


SUPPORTED_TARGETS = {
    'clusters': {
        'label': 'LGU Clusters',
        'required': ['name'],
        'columns': [
            {'field': 'name', 'label': 'Cluster name', 'required': True, 'aliases': ['name', 'cluster', 'cluster name', 'lgu cluster']},
            {'field': 'description', 'label': 'Description', 'required': False, 'aliases': ['description', 'remarks', 'notes']},
        ],
    },
    'lgus': {
        'label': 'LGUs',
        'required': ['name', 'municipality'],
        'columns': [
            {'field': 'name', 'label': 'LGU name', 'required': True, 'aliases': ['name', 'lgu name', 'lgu', 'local government unit']},
            {'field': 'municipality', 'label': 'Municipality', 'required': True, 'aliases': ['municipality', 'city', 'city municipality', 'city/municipality']},
            {'field': 'province', 'label': 'Province', 'required': False, 'aliases': ['province']},
            {'field': 'cluster', 'label': 'Cluster name or ID', 'required': False, 'aliases': ['cluster', 'cluster name', 'cluster id', 'lgu cluster']},
            {'field': 'mayor_name', 'label': 'Mayor name', 'required': False, 'aliases': ['mayor', 'mayor name', 'local chief executive']},
            {'field': 'mayor_email', 'label': 'Mayor email', 'required': False, 'aliases': ['mayor email', 'lce email']},
            {'field': 'contact_person', 'label': 'Contact person', 'required': False, 'aliases': ['contact person', 'focal person', 'lgu focal']},
            {'field': 'contact_email', 'label': 'Contact email', 'required': False, 'aliases': ['contact email', 'email']},
            {'field': 'contact_phone', 'label': 'Contact phone', 'required': False, 'aliases': ['contact phone', 'phone', 'mobile', 'telephone']},
            {'field': 'is_active', 'label': 'Active', 'required': False, 'aliases': ['active', 'is active', 'status']},
        ],
    },
    'programs': {
        'label': 'Programs',
        'required': ['name', 'code', 'program_type'],
        'columns': [
            {'field': 'name', 'label': 'Program name', 'required': True, 'aliases': ['name', 'program', 'program name']},
            {'field': 'code', 'label': 'Program code', 'required': True, 'aliases': ['code', 'program code', 'short code', 'acronym']},
            {'field': 'program_type', 'label': 'Program type', 'required': True, 'aliases': ['program type', 'type', 'category']},
            {'field': 'description', 'label': 'Description', 'required': False, 'aliases': ['description', 'details', 'remarks']},
            {'field': 'funding_agency', 'label': 'Funding agency', 'required': False, 'aliases': ['funding agency', 'agency', 'source agency']},
            {'field': 'is_active', 'label': 'Active', 'required': False, 'aliases': ['active', 'is active', 'status']},
        ],
    },
    'users': {
        'label': 'Internal Users',
        'required': ['email', 'first_name', 'last_name', 'role'],
        'columns': [
            {'field': 'email', 'label': 'Email', 'required': True, 'aliases': ['email', 'email address', 'user email']},
            {'field': 'first_name', 'label': 'First name', 'required': True, 'aliases': ['first name', 'given name', 'firstname']},
            {'field': 'last_name', 'label': 'Last name', 'required': True, 'aliases': ['last name', 'surname', 'family name', 'lastname']},
            {'field': 'role', 'label': 'Role', 'required': True, 'aliases': ['role', 'account role', 'user role']},
            {'field': 'clusters', 'label': 'Cluster names or IDs', 'required': False, 'aliases': ['clusters', 'cluster', 'cluster names', 'cluster ids', 'assigned clusters', 'scope']},
            {'field': 'lgus', 'label': 'LGU names or IDs', 'required': False, 'aliases': ['lgus', 'lgu', 'lgu names', 'lgu ids', 'assigned lgus', 'field lgus', 'field scope']},
            {'field': 'is_active', 'label': 'Active', 'required': False, 'aliases': ['active', 'is active', 'status']},
        ],
    },
}

TARGET_ALIASES = {
    normalize_label(target): target
    for target in SUPPORTED_TARGETS
}
TARGET_ALIASES.update({
    'lgu': 'lgus',
    'local government units': 'lgus',
    'local government unit': 'lgus',
    'lgu profiles': 'lgus',
    'lgu clusters': 'clusters',
    'cluster': 'clusters',
    'program': 'programs',
    'internal users': 'users',
    'user': 'users',
    'accounts': 'users',
})


def list_targets():
    return [
        {
            'target': target,
            **schema,
        }
        for target, schema in SUPPORTED_TARGETS.items()
    ]


def process_upload(target, upload, user_id, dry_run=True):
    target = _normalize_target(target)
    if upload is None:
        raise BatchImportError('Choose a CSV or XLSX file to import.', 422)

    rows = _read_upload(upload)
    if not rows:
        raise BatchImportError('No importable rows were found in the uploaded file.', 422)

    results = []
    seen_keys = set()
    for source_row in rows:
        try:
            payload = _payload_for_target(target, source_row)
            duplicate_reason = _find_duplicate_reason(target, payload, seen_keys)
            if duplicate_reason:
                result = _row_result(source_row, payload, 'skipped', 'skip_duplicate', duplicate_reason)
            elif dry_run:
                result = _row_result(source_row, payload, 'ready', 'create', 'Ready to create on commit.')
            else:
                result = _create_record(target, payload, user_id, source_row)
            for key in _dedupe_keys(target, payload):
                seen_keys.add(key)
        except (BatchImportError, auth_service.AuthError, lgu_service.LguError, program_service.ProgramError) as exc:
            db.session.rollback()
            message = getattr(exc, 'message', str(exc))
            result = _row_result(source_row, {}, 'failed', 'invalid', message)
        except Exception as exc:
            db.session.rollback()
            result = _row_result(source_row, {}, 'failed', 'invalid', f'Unexpected row error: {exc}')
        results.append(result)

    return {
        'target': target,
        'target_label': SUPPORTED_TARGETS[target]['label'],
        'dry_run': bool(dry_run),
        'file': {
            'filename': upload.filename,
            'row_count': len(rows),
        },
        'summary': _summarize(results, dry_run=dry_run),
        'rows': results,
    }


def _normalize_target(target):
    normalized = normalize_label(target)
    resolved = TARGET_ALIASES.get(normalized)
    if not resolved:
        raise BatchImportError('Unsupported batch import target.', 404)
    return resolved


def _read_upload(upload):
    if upload is None or not getattr(upload, 'filename', '').strip():
        raise BatchImportError('Choose a CSV or XLSX file to import.', 422)
    size = _measure_upload_size(upload)
    if size == 0:
        raise BatchImportError('The uploaded file is empty (0 bytes).', 422)
    _validate_upload_size(upload)
    suffix = Path(upload.filename or '').suffix.lower().lstrip('.')
    if suffix == 'csv':
        return _read_csv_rows(upload)
    if suffix == 'xlsx':
        try:
            validate_ooxml_container(upload, expected_kind='xlsx')
        except ArchiveSecurityError as exc:
            raise BatchImportError(f'Invalid or corrupt Excel workbook: {exc}', 422) from exc
        return _read_xlsx_rows(upload)
    if suffix == 'xls':
        raise BatchImportError('Legacy .xls files are not supported here. Save the workbook as .xlsx or CSV before importing.', 422)
    raise BatchImportError('Unsupported file type. Upload a CSV or XLSX file.', 422)


def _read_csv_rows(upload):
    upload.stream.seek(0)
    raw = upload.stream.read()
    if b'\x00' in raw[:4096]:
        raise BatchImportError('CSV file contains invalid binary content.', 422)
    text = _decode_csv(raw)
    matrix = [row for row in csv.reader(io.StringIO(text))]
    _validate_matrix_bounds(matrix)
    return _matrix_to_rows(matrix)


def _read_xlsx_rows(upload):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise BatchImportError('openpyxl is required to import Excel workbooks.', 500) from exc

    upload.stream.seek(0)
    try:
        workbook = load_workbook(upload.stream, data_only=True, read_only=True)
    except Exception as exc:
        raise BatchImportError('Unable to read this Excel workbook. Save it as a valid .xlsx file and try again.', 422) from exc

    worksheet = workbook.active
    try:
        matrix = [[cell for cell in row] for row in worksheet.iter_rows(
            max_row=_max_import_rows() + 1,
            max_col=_max_import_columns(),
            values_only=True,
        )]
        _validate_matrix_bounds(matrix)
        return _matrix_to_rows(matrix)
    finally:
        workbook.close()


def _validate_upload_size(upload):
    max_size_bytes = int(current_app.config.get('MAX_IMPORT_FILE_BYTES') or 0)
    if not max_size_bytes:
        return
    stream = getattr(upload, 'stream', None)
    if stream is None:
        return
    position = stream.tell()
    stream.seek(0, 2)
    size = stream.tell()
    stream.seek(position)
    if size > max_size_bytes:
        raise BatchImportError(
            f'Batch import files must be {max_size_bytes // (1024 * 1024)} MB or smaller.',
            422,
        )


def _validate_matrix_bounds(matrix):
    if len(matrix) > _max_import_rows():
        raise BatchImportError(f'Batch imports are limited to {_max_import_rows()} rows.', 422)
    if any(len(row) > _max_import_columns() for row in matrix):
        raise BatchImportError(f'Batch imports are limited to {_max_import_columns()} columns.', 422)


def _max_import_rows():
    return max(int(current_app.config.get('MAX_IMPORT_PROFILE_ROWS') or 5000), 1)


def _max_import_columns():
    return max(int(current_app.config.get('MAX_IMPORT_PROFILE_COLUMNS') or 200), 1)


def _decode_csv(raw):
    for encoding in ('utf-8-sig', 'utf-8', 'cp1252'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise BatchImportError('Unable to decode this CSV file. Save it as UTF-8 and try again.', 422)


def _matrix_to_rows(matrix):
    header_index = None
    for index, row in enumerate(matrix):
        if sum(1 for value in row if _clean_value(value) is not None) >= 1:
            header_index = index
            break

    if header_index is None:
        return []

    headers = [_clean_value(value) for value in matrix[header_index]]
    rows = []
    for offset, row in enumerate(matrix[header_index + 1:], start=header_index + 2):
        values = {}
        normalized = {}
        for index, header in enumerate(headers):
            if header is None:
                continue
            value = _clean_value(row[index] if index < len(row) else None)
            values[str(header)] = value
            normalized[normalize_label(header)] = value
        if any(value is not None for value in values.values()):
            rows.append({
                'row_number': offset,
                'values': values,
                'normalized': normalized,
            })
    return rows


def _payload_for_target(target, source_row):
    if target == 'clusters':
        payload = {
            'name': _field(source_row, target, 'name'),
            'description': _field(source_row, target, 'description'),
        }
        _require_fields(payload, ['name'])
        payload['name'] = _collapse_spaces(payload['name'])
        return payload

    if target == 'lgus':
        cluster_ref = _field(source_row, target, 'cluster')
        payload = {
            'name': _field(source_row, target, 'name'),
            'municipality': _field(source_row, target, 'municipality'),
            'province': _field(source_row, target, 'province') or 'Laguna',
            'cluster_id': _resolve_single_cluster(cluster_ref) if cluster_ref else None,
            'mayor_name': _field(source_row, target, 'mayor_name'),
            'mayor_email': _field(source_row, target, 'mayor_email'),
            'contact_person': _field(source_row, target, 'contact_person'),
            'contact_email': _field(source_row, target, 'contact_email'),
            'contact_phone': _field(source_row, target, 'contact_phone'),
            'is_active': _parse_bool(_field(source_row, target, 'is_active'), default=True),
        }
        _require_fields(payload, ['name', 'municipality'])
        payload['name'] = _collapse_spaces(payload['name'])
        payload['municipality'] = _collapse_spaces(payload['municipality'])
        payload['province'] = _collapse_spaces(payload['province'] or 'Laguna')
        return payload

    if target == 'programs':
        payload = {
            'name': _field(source_row, target, 'name'),
            'code': _field(source_row, target, 'code'),
            'program_type': _normalize_program_type(_field(source_row, target, 'program_type')),
            'description': _field(source_row, target, 'description'),
            'funding_agency': _field(source_row, target, 'funding_agency'),
            'is_active': _parse_bool(_field(source_row, target, 'is_active'), default=True),
        }
        _require_fields(payload, ['name', 'code', 'program_type'])
        payload['name'] = _collapse_spaces(payload['name'])
        payload['code'] = _collapse_spaces(payload['code']).upper()
        if len(payload['code']) > 20:
            raise BatchImportError('Program code must be 20 characters or less.', 422)
        return payload

    if target == 'users':
        role = _normalize_role(_field(source_row, target, 'role'))
        cluster_ref = _field(source_row, target, 'clusters')
        lgu_ref = _field(source_row, target, 'lgus')
        payload = {
            'email': _normalize_email(_field(source_row, target, 'email')),
            'first_name': _field(source_row, target, 'first_name'),
            'last_name': _field(source_row, target, 'last_name'),
            'role': role,
            'cluster_ids': _resolve_cluster_list(cluster_ref, required=(role == 'staff')),
            'lgu_ids': _resolve_lgu_list(lgu_ref, required=(role == 'field_officer')),
            'is_active': _parse_bool(_field(source_row, target, 'is_active'), default=True),
        }
        _require_fields(payload, ['email', 'first_name', 'last_name', 'role'])
        if '@' not in payload['email']:
            raise BatchImportError('Invalid email format.', 422)
        payload['first_name'] = _collapse_spaces(payload['first_name'])
        payload['last_name'] = _collapse_spaces(payload['last_name'])
        return payload

    raise BatchImportError('Unsupported batch import target.', 404)


def _field(source_row, target, field_name):
    aliases = []
    for column in SUPPORTED_TARGETS[target]['columns']:
        if column['field'] == field_name:
            aliases = column.get('aliases', [])
            break
    normalized = source_row.get('normalized') or {}
    for alias in [field_name, *aliases]:
        value = normalized.get(normalize_label(alias))
        if value is not None:
            return value
    return None


def _require_fields(payload, fields):
    missing = [field for field in fields if payload.get(field) in (None, '')]
    if missing:
        readable = ', '.join(field.replace('_', ' ') for field in missing)
        raise BatchImportError(f'Missing required field(s): {readable}.', 422)


def _find_duplicate_reason(target, payload, seen_keys):
    for key in _dedupe_keys(target, payload):
        if key in seen_keys:
            return 'Duplicate row in uploaded file.'

    if target == 'clusters':
        if _cluster_by_name(payload['name']):
            return 'Cluster name already exists.'
    elif target == 'lgus':
        if _lgu_duplicate(payload):
            return 'LGU profile already exists.'
    elif target == 'programs':
        if _program_by_code(payload['code']):
            return 'Program code already exists.'
        if _program_by_name(payload['name']):
            return 'Program name already exists.'
    elif target == 'users':
        if _user_by_email(payload['email']):
            return 'Email already registered.'
    return None


def _create_record(target, payload, user_id, source_row):
    if target == 'clusters':
        record = lgu_service.create_cluster(
            name=payload['name'],
            description=payload.get('description'),
            user_id=user_id,
        )
        return _row_result(source_row, payload, 'created', 'create', 'Cluster created.', record=lgu_service._serialize_cluster(record))

    if target == 'lgus':
        record = lgu_service.create_lgu(
            name=payload['name'],
            municipality=payload['municipality'],
            province=payload.get('province') or 'Laguna',
            cluster_id=payload.get('cluster_id'),
            mayor_name=payload.get('mayor_name'),
            mayor_email=payload.get('mayor_email'),
            contact_person=payload.get('contact_person'),
            contact_email=payload.get('contact_email'),
            contact_phone=payload.get('contact_phone'),
            user_id=user_id,
        )
        if payload.get('is_active') is False:
            record = lgu_service.toggle_lgu_active(record.id, user_id=user_id)
        return _row_result(source_row, payload, 'created', 'create', 'LGU profile created.', record=lgu_service._serialize_lgu(record))

    if target == 'programs':
        record = program_service.create_program(
            name=payload['name'],
            code=payload['code'],
            program_type=payload['program_type'],
            description=payload.get('description'),
            funding_agency=payload.get('funding_agency'),
            user_id=user_id,
        )
        if payload.get('is_active') is False:
            record = program_service.update_program(record.id, {'is_active': False}, user_id=user_id)
        return _row_result(source_row, payload, 'created', 'create', 'Program created.', record=program_service._serialize_program(record))

    if target == 'users':
        result = auth_service.create_user_account(
            admin_id=user_id,
            email=payload['email'],
            first_name=payload['first_name'],
            last_name=payload['last_name'],
            role=payload['role'],
            cluster_ids=payload.get('cluster_ids') or [],
            lgu_ids=payload.get('lgu_ids') or [],
        )
        user = result['user']
        if payload.get('is_active') is False:
            user = auth_service.toggle_user_active(user_id, user['id'], False)
        return _row_result(
            source_row,
            payload,
            'created',
            'create',
            'User account created.',
            record=user,
            delivery=result.get('delivery'),
        )

    raise BatchImportError('Unsupported batch import target.', 404)


def _row_result(source_row, payload, status, action, message, record=None, delivery=None):
    identifier = _identifier(payload)
    result = {
        'row_number': source_row.get('row_number'),
        'status': status,
        'action': action,
        'identifier': identifier,
        'message': message,
        'data': payload,
    }
    if record is not None:
        result['record'] = record
    if delivery is not None:
        result['delivery'] = delivery
    return result


def _summarize(results, dry_run=False):
    counts = {
        'total_rows': len(results),
        'ready': 0,
        'created': 0,
        'skipped': 0,
        'failed': 0,
    }
    for result in results:
        status = result.get('status')
        if status in counts:
            counts[status] += 1
    counts['valid_for_commit'] = counts['ready'] if dry_run else counts['created']
    return counts


def _dedupe_keys(target, payload):
    if target == 'clusters':
        return [f"cluster:{_normalized_text(payload.get('name'))}"]
    if target == 'lgus':
        return [
            f"lgu-name:{_normalized_text(payload.get('name'))}",
            f"lgu-place:{_normalized_text(payload.get('municipality'))}:{_normalized_text(payload.get('province'))}",
        ]
    if target == 'programs':
        return [
            f"program-code:{_normalized_text(payload.get('code'))}",
            f"program-name:{_normalized_text(payload.get('name'))}",
        ]
    if target == 'users':
        return [f"user:{_normalize_email(payload.get('email'))}"]
    return []


def _identifier(payload):
    for key in ('email', 'code', 'name', 'municipality'):
        value = payload.get(key)
        if value:
            return str(value)
    return None


def _resolve_single_cluster(value):
    if value in (None, ''):
        return None
    items = _split_multi_value(value)
    if len(items) > 1:
        raise BatchImportError('LGU rows can only reference one cluster.', 422)
    return _resolve_cluster_token(items[0])


def _resolve_cluster_list(value, required=True):
    cluster_ids = [_resolve_cluster_token(item) for item in _split_multi_value(value)]
    cluster_ids = list(dict.fromkeys(cluster_ids))
    if required and not cluster_ids:
        raise BatchImportError('Staff rows must include at least one existing cluster.', 422)
    return cluster_ids


def _resolve_cluster_token(token):
    if token in (None, ''):
        raise BatchImportError('Cluster value is blank.', 422)
    text = str(token).strip()
    if text.isdigit():
        cluster = db.session.get(LguCluster, int(text))
        if not cluster:
            raise BatchImportError(f'Cluster {text} not found.', 422)
        return cluster.id

    cluster = _cluster_by_name(text)
    if not cluster:
        raise BatchImportError(f'Cluster "{text}" not found.', 422)
    return cluster.id


def _cluster_by_name(name):
    normalized = _normalized_text(name)
    if not normalized:
        return None
    return LguCluster.query.filter(func.lower(func.trim(LguCluster.name)) == normalized).first()


def _resolve_lgu_list(value, required=True):
    lgu_ids = [_resolve_lgu_token(item) for item in _split_multi_value(value)]
    lgu_ids = list(dict.fromkeys(lgu_ids))
    if required and not lgu_ids:
        raise BatchImportError('Field officer rows must include at least one existing LGU.', 422)
    return lgu_ids


def _resolve_lgu_token(token):
    if token in (None, ''):
        raise BatchImportError('LGU value is blank.', 422)
    text = str(token).strip()
    if text.isdigit():
        lgu = db.session.get(LguProfile, int(text))
        if not lgu:
            raise BatchImportError(f'LGU {text} not found.', 422)
        return lgu.id

    lgu = _lgu_by_name_or_municipality(text)
    if not lgu:
        raise BatchImportError(f'LGU "{text}" not found.', 422)
    return lgu.id


def _lgu_by_name_or_municipality(name):
    normalized = _normalized_text(name)
    if not normalized:
        return None
    matches = LguProfile.query.filter(
        db.or_(
            func.lower(func.trim(LguProfile.name)) == normalized,
            func.lower(func.trim(LguProfile.municipality)) == normalized,
        )
    ).all()
    if len(matches) > 1:
        raise BatchImportError(f'LGU "{name}" matched more than one profile. Use the LGU ID instead.', 422)
    return matches[0] if matches else None


def _lgu_duplicate(payload):
    name_match = LguProfile.query.filter(
        func.lower(func.trim(LguProfile.name)) == _normalized_text(payload.get('name')),
    ).first()
    if name_match:
        return name_match
    return LguProfile.query.filter(
        func.lower(func.trim(LguProfile.municipality)) == _normalized_text(payload.get('municipality')),
        func.lower(func.trim(LguProfile.province)) == _normalized_text(payload.get('province') or 'Laguna'),
    ).first()


def _program_by_code(code):
    return Program.query.filter(func.lower(func.trim(Program.code)) == _normalized_text(code)).first()


def _program_by_name(name):
    return Program.query.filter(func.lower(func.trim(Program.name)) == _normalized_text(name)).first()


def _user_by_email(email):
    return User.query.filter(func.lower(func.trim(User.email)) == _normalize_email(email)).first()


def _normalize_role(value):
    normalized = normalize_label(value)
    aliases = {
        'admin': 'admin',
        'administrator': 'admin',
        'pdmu admin': 'admin',
        'staff': 'staff',
        'pdmu staff': 'staff',
        'encoder': 'staff',
        'evaluator': 'staff',
        'field officer': 'field_officer',
        'field side officer': 'field_officer',
        'c mlgoo': 'field_officer',
        'clgoo': 'field_officer',
        'mlgoo': 'field_officer',
    }
    role = aliases.get(normalized)
    if not role:
        raise BatchImportError('Role must be admin, staff, or field_officer.', 422)
    return role


def _normalize_program_type(value):
    normalized = normalize_label(value)
    aliases = {
        'infrastructure': 'infrastructure',
        'infra': 'infrastructure',
        'non infrastructure': 'non_infrastructure',
        'non infra': 'non_infrastructure',
        'non_infrastructure': 'non_infrastructure',
        'hybrid': 'hybrid',
    }
    program_type = aliases.get(normalized)
    if not program_type:
        raise BatchImportError('Program type must be infrastructure, non_infrastructure, or hybrid.', 422)
    return program_type


def _parse_bool(value, default=None):
    if value in (None, ''):
        return default
    normalized = normalize_label(value)
    if normalized in {'true', 'yes', 'y', '1', 'active', 'enabled'}:
        return True
    if normalized in {'false', 'no', 'n', '0', 'inactive', 'disabled'}:
        return False
    return default


def _split_multi_value(value):
    if value in (None, ''):
        return []
    if isinstance(value, (list, tuple, set)):
        items = value
    else:
        text = str(value)
        for separator in ['|', ';']:
            text = text.replace(separator, ',')
        items = text.split(',')
    return [str(item).strip() for item in items if str(item).strip()]


def _clean_value(value):
    if value is None:
        return None
    if isinstance(value, str):
        text = _collapse_spaces(value)
        return text or None
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    return value


def _collapse_spaces(value):
    if value is None:
        return None
    return ' '.join(str(value).strip().split())


def _normalized_text(value):
    return _collapse_spaces(value or '').lower()


def _normalize_email(value):
    return _collapse_spaces(value or '').lower()
