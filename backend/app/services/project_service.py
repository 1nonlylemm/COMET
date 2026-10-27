from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
import html
import logging
import os
from pathlib import Path
import re
import zipfile

from flask import current_app

logger = logging.getLogger(__name__)

from sqlalchemy import func
from sqlalchemy.exc import IntegrityError

from app.extensions import db
from app.models import (
    AuditLog,
    ChangeStatusRequest,
    DocumentFile,
    GeneratedReport,
    GoogleDrivePush,
    LguProfile,
    NlpScanResult,
    Program,
    ProgramCycle,
    ProgramPhase,
    Project,
    ProjectPhaseLog,
    ProjectRequirement,
    Requirement,
    SlaDeadline,
    Submission,
    User,
)
from app.models.workflow import RequirementWorkflowTransition, RequirementDeficiency
from app.services import (
    access_service,
    audit_service,
    document_template_service,
    object_storage_service,
    requirement_workflow_service,
    search_service,
    settings_service,
    submission_service,
)
from app.services.fiscal_year_scope import apply_project_fiscal_year
from app.services.project_status import CANONICAL_PROJECT_STATUSES, normalize_project_status
from app.services.email_service import _format_requirement_status_label


class ProjectError(Exception):
    def __init__(self, message, status_code=400, errors=None):
        super().__init__(message)
        self.message = message
        self.status_code = status_code
        self.errors = errors


PROJECT_IDENTITY_CONSTRAINTS = {
    'uq_projects_null_source_identity',
    'uq_projects_source_code',
    'uq_project_cycle_source_code',
    'projects_program_cycle_id_source_project_code_key',
}


def _is_project_identity_conflict(exc):
    """Narrowly identifies whether an IntegrityError was caused by a project identity collision.

    Inspects PostgreSQL diagnostics when available (exc.orig.diag.constraint_name),
    or uses exact narrow SQLite constraint messages when running under SQLite.
    Never matches generic substrings or unrelated field names.
    """
    orig = getattr(exc, 'orig', None)
    if orig is not None:
        diag = getattr(orig, 'diag', None)
        if diag is not None:
            constraint_name = getattr(diag, 'constraint_name', None)
            if constraint_name:
                return constraint_name in PROJECT_IDENTITY_CONSTRAINTS

    # Fallback for SQLite / engines without psycopg2 diag
    msg = str(orig if orig is not None else exc).lower()
    if 'unique constraint failed:' in msg:
        if "index 'uq_projects_null_source_identity'" in msg or "index 'uq_projects_source_code'" in msg:
            return True
        if 'projects.program_cycle_id, projects.source_project_code' in msg:
            return True
        if 'projects.lgu_id, projects.fiscal_year, lower(trim(projects.project_name))' in msg:
            return True
    elif any(name in msg for name in PROJECT_IDENTITY_CONSTRAINTS):
        return True

    return False


VALID_PHASES = ('pre_implementation', 'implementation', 'completion')
VALID_STATUSES = CANONICAL_PROJECT_STATUSES
VALID_REQ_STATUSES = ('pending', 'under_review', 'compliant', 'non_compliant', 'waived')
REVIEWED_REQ_STATUSES = {'compliant', 'non_compliant', 'waived'}
FORMAL_COMPLIANCE_DEADLINE_TYPES = {'standard_15day', 'rectification_period'}
NON_COMPLIANT_REQUIRES_OVERDUE_DEADLINE_MESSAGE = (
    'A requirement can only be marked non-compliant after its formal compliance deadline is overdue. '
    'Custom tracking deadlines cannot authorize formal non-compliance. '
    'Use Return for Revision for documents that need correction before the deadline.'
)
MIN_FISCAL_YEAR = 2000
MAX_FISCAL_YEAR = 2100
NEXT_PHASE_BY_CURRENT = {'pre_implementation': 'implementation', 'implementation': 'completion'}
VALID_PROJECT_KINDS = ('infrastructure', 'goods', 'vehicle', 'financial_assistance')
NON_INFRA_PROJECT_KINDS = ('goods', 'vehicle', 'financial_assistance')

PROJECT_KIND_LABELS = {
    'infrastructure': 'Infrastructure',
    'goods': 'Goods',
    'vehicle': 'Vehicle',
    'financial_assistance': 'Financial Assistance',
}


def format_project_kind(project_kind: str | None, fallback: str | None = None) -> str:
    """Format a project kind into COMET canonical human-readable presentation."""
    if not project_kind:
        return fallback or 'Infrastructure'
    clean_kind = str(project_kind).strip()
    if not clean_kind:
        return fallback or 'Infrastructure'
    if clean_kind in PROJECT_KIND_LABELS:
        return PROJECT_KIND_LABELS[clean_kind]
    lower_kind = clean_kind.lower()
    if lower_kind in PROJECT_KIND_LABELS:
        return PROJECT_KIND_LABELS[lower_kind]
    if fallback is not None:
        return fallback
    return clean_kind.replace('_', ' ').title()



def _clean_source_value(value):
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def validate_project_kind_for_program_type(program_type, project_kind):
    cleaned = _clean_source_value(project_kind)
    if cleaned:
        cleaned = cleaned.casefold().replace('-', '_').replace(' ', '_')
    if not cleaned:
        if program_type in ('hybrid', 'non_infrastructure'):
            label = 'non-infrastructure' if program_type == 'non_infrastructure' else 'hybrid'
            raise ProjectError(
                f"Project classification (project_kind) is required for projects under a {label} program",
                422,
            )
        if program_type == 'infrastructure':
            return 'infrastructure'
        return None

    if cleaned not in VALID_PROJECT_KINDS:
        raise ProjectError(
            f"Invalid project_kind. Must be one of: {', '.join(VALID_PROJECT_KINDS)}",
            422,
        )

    if program_type == 'infrastructure' and cleaned != 'infrastructure':
        raise ProjectError(
            f"Infrastructure programs only accept project_kind='infrastructure', got '{cleaned}'",
            422,
        )
    if program_type == 'non_infrastructure' and cleaned not in NON_INFRA_PROJECT_KINDS:
        raise ProjectError(
            f"Non-infrastructure programs only accept project_kind of 'goods', 'vehicle', or 'financial_assistance', got '{cleaned}'",
            422,
        )
    return cleaned


_validate_project_kind_for_program_type = validate_project_kind_for_program_type


def resolve_effective_project_kind(project):
    """Canonical domain resolver for a project's effective project kind.
    
    Consumed by requirements, monitoring, evidence, completion, reporting,
    deadlines, and AI layers.
    """
    if not project:
        return 'infrastructure'

    kind = getattr(project, 'project_kind', None)
    if kind:
        cleaned = str(kind).strip().lower()
        if cleaned in VALID_PROJECT_KINDS:
            return cleaned

    program = getattr(project, 'program', None)
    if not program and getattr(project, 'program_id', None):
        program = db.session.get(Program, project.program_id)
    prog_type = getattr(program, 'program_type', None) if program else None

    if prog_type == 'infrastructure':
        return 'infrastructure'
    elif prog_type == 'non_infrastructure':
        return 'goods'
    elif prog_type == 'hybrid':
        return 'infrastructure'

    return 'infrastructure'


PHASE_NAME_BY_KEY = {
    'pre_implementation': 'Pre-Implementation',
    'implementation': 'Implementation',
    'completion': 'Completion',
}
DOCX_MIME_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
WORKBOOK_SYNC_PROJECT_FIELDS = {
    'project_name',
    'budget_amount',
    'overall_status',
    'physical_completion_pct',
    'source_project_code',
}


def _translate_access_error(exc):
    return ProjectError(exc.message, exc.status_code)


def _get_project_with_access(project_id, user_id=None):
    try:
        if user_id is not None:
            return access_service.ensure_project_access(project_id, user_id)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)

    project = db.session.get(Project, project_id)
    if not project:
        raise ProjectError('Project not found', 404)
    return project


def _get_project_requirement_with_access(requirement_id, user_id=None):
    try:
        if user_id is not None:
            return access_service.ensure_project_requirement_access(requirement_id, user_id)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)

    requirement = db.session.get(ProjectRequirement, requirement_id)
    if not requirement:
        raise ProjectError('Project requirement not found', 404)
    return requirement


def _require_role(user_id, *roles):
    role = access_service.get_user_role(user_id)
    if role not in roles:
        raise ProjectError(
            f"Role '{role}' is not authorized for this action. Requires one of: {', '.join(roles)}",
            403,
        )
    return role


def create_project(lgu_id, program_id, project_name, fiscal_year,
                   program_cycle_id=None, source_project_code=None,
                   source_row_key=None, source_sheet_name=None,
                   current_phase='pre_implementation', overall_status='not_yet_started',
                   budget_amount=None, target_start_date=None, target_end_date=None,
                   project_kind=None, monitoring_applicability=None, qaar_category=None,
                   subaybayan_project_code=None, subaybayan_url=None,
                   nadai_received_at=None, subaybayan_enrolled_at=None,
                   subaybayan_enrollment_status=None,
                   notes=None, custom_fields=None, user_id=None):
    try:
        if user_id is not None:
            access_service.ensure_lgu_access(lgu_id, user_id)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)

    lgu = db.session.get(LguProfile, lgu_id)
    if not lgu:
        raise ProjectError('LGU not found', 404)
    if not lgu.is_active:
        raise ProjectError('Cannot create project under an inactive LGU', 422, errors={'lgu_id': 'LGU is inactive'})

    program = db.session.get(Program, program_id)
    if not program:
        raise ProjectError('Program not found', 404)
    if not program.is_active:
        raise ProjectError('Cannot create project under an inactive program', 422, errors={'program_id': 'Program is inactive'})
    project_kind = _validate_project_kind_for_program_type(program.program_type, project_kind)

    normalized = _normalize_project_create_values(
        project_name=project_name,
        fiscal_year=fiscal_year,
        budget_amount=budget_amount,
        target_start_date=target_start_date,
        target_end_date=target_end_date,
    )
    project_name = normalized['project_name']
    fiscal_year = normalized['fiscal_year']
    budget_amount = normalized['budget_amount']
    target_start_date = normalized['target_start_date']
    target_end_date = normalized['target_end_date']

    program_cycle = None
    if program_cycle_id is not None:
        program_cycle = db.session.get(ProgramCycle, program_cycle_id)
        if not program_cycle:
            raise ProjectError('Program cycle not found', 404)
        if program_cycle.program_id != program_id:
            raise ProjectError('Program cycle does not belong to the selected program', 422)
        if fiscal_year != program_cycle.fiscal_year:
            raise ProjectError('Fiscal year must match the selected program cycle', 422)
    else:
        cycles_for_year = (
            ProgramCycle.query
            .filter_by(program_id=program_id, fiscal_year=fiscal_year)
            .order_by(ProgramCycle.name.asc(), ProgramCycle.id.asc())
            .all()
        )
        if cycles_for_year:
            active_cycles = [c for c in cycles_for_year if c.is_active]
            if len(active_cycles) == 1:
                program_cycle = active_cycles[0]
                program_cycle_id = program_cycle.id
            elif len(active_cycles) > 1:
                raise ProjectError(
                    f'Multiple active program cycles exist for FY {fiscal_year}. Please specify a program cycle.',
                    422,
                    errors={'program_cycle_id': f'Multiple active program cycles exist for FY {fiscal_year}. Please specify a program cycle.'},
                )
            else:
                raise ProjectError(
                    f'The program cycle for FY {fiscal_year} is inactive.',
                    422,
                    errors={'program_cycle_id': f'The program cycle for FY {fiscal_year} is inactive.'},
                )

    if current_phase not in VALID_PHASES:
        raise ProjectError(f'Invalid current_phase. Must be one of: {", ".join(VALID_PHASES)}', 422)
    overall_status = normalize_project_status(overall_status, default=overall_status)
    if overall_status not in VALID_STATUSES:
        raise ProjectError(f'Invalid overall_status. Must be one of: {", ".join(VALID_STATUSES)}', 422)

    _assert_project_identity_available(
        lgu_id=lgu_id,
        program_id=program_id,
        fiscal_year=fiscal_year,
        project_name=project_name,
        program_cycle_id=program_cycle_id,
        source_project_code=source_project_code,
    )

    project = Project(
        lgu_id=lgu_id,
        program_id=program_id,
        program_cycle_id=program_cycle_id,
        project_name=project_name,
        fiscal_year=fiscal_year,
        source_project_code=_clean_source_value(source_project_code),
        source_row_key=_clean_source_value(source_row_key),
        source_sheet_name=_clean_source_value(source_sheet_name),
        current_phase=current_phase,
        overall_status=overall_status,
        budget_amount=budget_amount,
        target_start_date=target_start_date,
        target_end_date=target_end_date,
        project_kind=_clean_source_value(project_kind),
        monitoring_applicability=_clean_source_value(monitoring_applicability),
        qaar_category=_clean_source_value(qaar_category),
        subaybayan_project_code=_clean_source_value(subaybayan_project_code),
        subaybayan_url=_clean_source_value(subaybayan_url),
        nadai_received_at=_parse_date(nadai_received_at),
        subaybayan_enrolled_at=_parse_date(subaybayan_enrolled_at),
        subaybayan_enrollment_status=_clean_source_value(subaybayan_enrollment_status),
        notes=notes,
        custom_fields=custom_fields or {},
    )
    try:
        db.session.add(project)
        db.session.flush()

        initialized_requirement_count = _create_project_requirement_snapshots(project)

        if user_id is not None:
            audit_service.log_action(
                user_id=user_id,
                action='create',
                entity_type='project',
                entity_id=project.id,
                old_value=None,
                new_value={
                    'project_name': project.project_name,
                    'lgu_id': project.lgu_id,
                    'program_id': project.program_id,
                    'program_cycle_id': project.program_cycle_id,
                    'fiscal_year': project.fiscal_year,
                    'initialized_requirement_count': initialized_requirement_count,
                },
                commit=False,
            )

        db.session.commit()
    except IntegrityError as exc:
        db.session.rollback()
        if _is_project_identity_conflict(exc):
            raise _project_identity_error(source_project_code) from exc
        raise
    except Exception:
        db.session.rollback()
        raise

    return project


def get_project(project_id, user_id=None):
    return _get_project_with_access(project_id, user_id=user_id)


def list_projects(search=None, lgu_id=None, program_id=None, fiscal_year=None,
                 program_cycle_id=None,
                 current_phase=None, overall_status=None, has_non_compliant_requirements=None, page=1, per_page=20,
                 user_id=None, project_ids=None):
    from app.models import ProjectRequirement
    query = Project.query
    if user_id is not None:
        query = access_service.restrict_project_query(query, user_id)

    if project_ids is not None:
        query = query.filter(Project.id.in_(project_ids))
    if lgu_id is not None:
        query = query.filter(Project.lgu_id == lgu_id)
    if program_id is not None:
        query = query.filter(Project.program_id == program_id)
    if program_cycle_id is not None:
        query = query.filter(Project.program_cycle_id == program_cycle_id)
    if fiscal_year is not None:
        query = query.filter(Project.fiscal_year == fiscal_year)
    if current_phase:
        query = query.filter(Project.current_phase == current_phase)
    if overall_status:
        overall_status = normalize_project_status(overall_status, default=overall_status)
        query = query.filter(Project.overall_status == overall_status)
    if has_non_compliant_requirements is not None:
        has_nc = ProjectRequirement.query.filter(
            ProjectRequirement.project_id == Project.id,
            ProjectRequirement.status == 'non_compliant'
        ).exists()
        if has_non_compliant_requirements:
            query = query.filter(has_nc)
        else:
            query = query.filter(~has_nc)

    if search:
        return search_service.search_projects(
            query,
            search,
            page=page,
            per_page=per_page,
            serialize=_serialize_project,
        )

    pagination = query.order_by(Project.created_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    return {
        'projects': [_serialize_project(project) for project in pagination.items],
        'total': pagination.total,
        'page': pagination.page,
        'per_page': pagination.per_page,
        'pages': pagination.pages,
    }


def update_project(project_id, data, user_id=None, enforce_workbook_locks=True, sync_workbook=True, commit=True):
    project = _get_project_with_access(project_id, user_id=user_id)
    try:
        access_service.assert_project_movement_allowed(project)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)
    before = _serialize_project(project)
    if 'current_phase' in data and data['current_phase'] != project.current_phase:
        raise ProjectError('Use the phase transition endpoint to change project phases.', 422)
    data = {**data, **_normalize_project_update_values(project, data)}

    target_lgu = db.session.get(LguProfile, data.get('lgu_id', project.lgu_id))
    if not target_lgu:
        raise ProjectError('LGU not found', 404)
    if target_lgu.id != project.lgu_id and not target_lgu.is_active:
        raise ProjectError('Cannot reassign project to an inactive LGU', 422, errors={'lgu_id': 'LGU is inactive'})
    try:
        if user_id is not None:
            access_service.ensure_lgu_access(target_lgu, user_id)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)

    new_program_id = data.get('program_id', project.program_id)
    target_program = db.session.get(Program, new_program_id)
    if not target_program:
        raise ProjectError('Program not found', 404)
    if new_program_id != project.program_id and not target_program.is_active:
        raise ProjectError('Cannot reassign project to an inactive program', 422,
                           errors={'program_id': 'Program is inactive'})
    new_fiscal_year = data.get('fiscal_year', project.fiscal_year)
    scope_changed = new_program_id != project.program_id or new_fiscal_year != project.fiscal_year
    new_program_cycle_id = data.get('program_cycle_id', None if scope_changed else project.program_cycle_id)
    if new_program_cycle_id is None and (scope_changed or 'program_cycle_id' in data):
        cycles = ProgramCycle.query.filter_by(
            program_id=new_program_id, fiscal_year=new_fiscal_year, is_active=True,
        ).all()
        if len(cycles) == 1:
            new_program_cycle_id = cycles[0].id
        elif len(cycles) > 1:
            raise ProjectError('Please select a program cycle for this program and fiscal year', 422,
                               errors={'program_cycle_id': 'Please select a program cycle.'})
        elif ProgramCycle.query.filter_by(program_id=new_program_id, fiscal_year=new_fiscal_year).first():
            raise ProjectError('No active program cycle exists for the selected fiscal year', 422,
                               errors={'program_cycle_id': 'No active cycle is available.'})
    target_cycle = db.session.get(ProgramCycle, new_program_cycle_id) if new_program_cycle_id else None
    if new_program_cycle_id is not None:
        if not target_cycle:
            raise ProjectError('Program cycle not found', 404)
        if target_cycle.program_id != new_program_id:
            raise ProjectError('Program cycle does not belong to the selected program', 422,
                               errors={'program_cycle_id': 'Select a cycle belonging to the selected program.'})
        if new_fiscal_year != target_cycle.fiscal_year:
            raise ProjectError('Fiscal year must match the selected program cycle', 422)
        if new_program_cycle_id != project.program_cycle_id and not target_cycle.is_active:
            raise ProjectError('Cannot assign an inactive program cycle', 422)
    target_kind = project.project_kind
    if new_program_id != project.program_id or 'project_kind' in data:
        target_kind = _validate_project_kind_for_program_type(
            target_program.program_type, data.get('project_kind', project.project_kind),
        )
    checklist_changed = (new_program_id != project.program_id
                         or new_program_cycle_id != project.program_cycle_id
                         or target_kind != project.project_kind)
    new_source_project_code = data.get('source_project_code', project.source_project_code)
    _assert_project_identity_available(
        lgu_id=target_lgu.id, program_id=new_program_id, fiscal_year=new_fiscal_year,
        project_name=data.get('project_name', project.project_name),
        program_cycle_id=new_program_cycle_id, source_project_code=new_source_project_code,
        exclude_project_id=project.id,
    )
    if 'overall_status' in data:
        raw_status = data['overall_status']
        if raw_status is None:
            raise ProjectError('overall_status cannot be null', 422)
        normalized_status = normalize_project_status(raw_status, default=raw_status)
        if normalized_status == 'completed':
            raise ProjectError('Project completion must be processed through the completion closeout service.', 422)
        if normalized_status not in VALID_STATUSES:
            raise ProjectError(f'Invalid overall_status. Must be one of: {", ".join(VALID_STATUSES)}', 422)
        data['overall_status'] = normalized_status

    workbook_changed_fields = _changed_workbook_project_fields(project, data)
    if enforce_workbook_locks and workbook_changed_fields:
        try:
            from app.services import workbook_sync_service
            workbook_sync_service.assert_project_fields_unlocked(project.id, workbook_changed_fields)
        except workbook_sync_service.WorkbookSyncError as exc:
            raise ProjectError(exc.message, exc.status_code) from exc

    try:
        project.lgu = target_lgu
        project.lgu_id = target_lgu.id
        project.program = target_program
        project.program_id = new_program_id
        project.program_cycle = target_cycle
        project.program_cycle_id = new_program_cycle_id
        project.fiscal_year = new_fiscal_year
        project.project_kind = target_kind
        if target_kind != before['project_kind'] and 'monitoring_applicability' not in data:
            project.monitoring_applicability = None
        if target_kind != 'infrastructure' and target_kind != before['project_kind'] and 'qaar_category' not in data:
            project.qaar_category = None
        if data.get('project_name'):
            project.project_name = data['project_name']
        for field in ('overall_status', 'budget_amount', 'physical_completion_pct', 'notes', 'custom_fields'):
            if field in data:
                setattr(project, field, data[field])
        for field in ('target_start_date', 'target_end_date', 'actual_start_date', 'actual_end_date',
                      'nadai_received_at', 'subaybayan_enrolled_at'):
            if field in data:
                setattr(project, field, _parse_date(data[field]))
        for field in ('monitoring_applicability', 'qaar_category', 'subaybayan_project_code',
                      'subaybayan_url', 'subaybayan_enrollment_status', 'source_project_code',
                      'source_row_key', 'source_sheet_name'):
            if field in data:
                setattr(project, field, _clean_source_value(data[field]))
        if checklist_changed:
            db.session.flush()
            reconcile_project_requirements(
                project.id, user_id=user_id, commit=False,
                retire_unmatched=scope_changed or new_program_cycle_id != before['program_cycle_id'],
            )
        if sync_workbook and workbook_changed_fields:
            try:
                from app.services import workbook_sync_service
                workbook_sync_service.sync_project_fields(project, workbook_changed_fields, user_id=user_id, commit=False)
            except workbook_sync_service.WorkbookSyncError as exc:
                raise ProjectError(exc.message, exc.status_code) from exc
        if user_id is not None:
            audit_service.log_action(
                user_id=user_id, action='update', entity_type='project', entity_id=project.id,
                old_value=before, new_value=_serialize_project(project), commit=False,
            )
        if commit:
            db.session.commit()
        else:
            db.session.flush()
    except IntegrityError as exc:
        db.session.rollback()
        if _is_project_identity_conflict(exc):
            raise _project_identity_error(new_source_project_code) from exc
        raise
    except Exception:
        db.session.rollback()
        raise
    return project


def _changed_workbook_project_fields(project, data):
    changed = []
    for field in WORKBOOK_SYNC_PROJECT_FIELDS:
        if field not in data:
            continue
        if not _project_sync_values_equal(getattr(project, field, None), data.get(field)):
            changed.append(field)
    return changed


def _project_sync_values_equal(current, incoming):
    if current in (None, '') and incoming in (None, ''):
        return True
    try:
        if isinstance(current, (int, float, Decimal)) or isinstance(incoming, (int, float, Decimal, str)):
            return float(current or 0) == float(incoming or 0)
    except (TypeError, ValueError):
        pass
    return current == incoming


def transition_project_phase(project_id, to_phase, reason, user_id=None):
    project = _get_project_with_access(project_id, user_id=user_id)
    try:
        access_service.assert_project_movement_allowed(project)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)
    if to_phase not in VALID_PHASES:
        raise ProjectError(f'Invalid phase. Must be one of: {", ".join(VALID_PHASES)}', 422)

    from_phase = project.current_phase
    if from_phase == to_phase:
        raise ProjectError(f'Project is already in {to_phase} phase', 409)

    expected_next_phase = NEXT_PHASE_BY_CURRENT.get(from_phase)
    if expected_next_phase != to_phase:
        if expected_next_phase is None:
            raise ProjectError('Project is already in the final phase and cannot advance further.', 422)
        raise ProjectError(
            f'Project can only transition from {from_phase} to {expected_next_phase}.',
            422,
        )

    _assert_phase_gate(project.id, from_phase)

    phase_log = ProjectPhaseLog(
        project_id=project.id,
        from_phase=from_phase,
        to_phase=to_phase,
        transitioned_by=user_id,
        notes=reason,
    )
    project.current_phase = to_phase
    if to_phase == 'implementation' and project.actual_start_date is None:
        project.actual_start_date = date.today()
    if to_phase == 'implementation' and project.overall_status == 'not_yet_started':
        project.overall_status = 'ongoing'
    if to_phase == 'completion' and project.actual_end_date is None:
        project.actual_end_date = date.today()
    # NOTE: overall_status is intentionally NOT set to 'completed' here.
    # Entering the completion phase only means closeout work (Completion
    # Matrix: Final SWA, FIR, COFA, Final SORD, Proof of Return of Funds)
    # can now begin -- it does not mean the project is actually done.
    # `overall_status` only becomes 'completed' via
    # completion_service.update_completion_detail()'s final-approval path,
    # which is gated by completion_documents_service.assert_ready_for_final_closeout().
    # Setting it here would let a project appear "completed" before any
    # closeout document is even uploaded, bypassing that guard entirely.
    if to_phase == 'completion' and project.overall_status == 'not_yet_started':
        project.overall_status = 'ongoing'

    try:
        db.session.add(phase_log)

        if user_id is not None:
            audit_service.log_action(
                user_id=user_id,
                action='status_change',
                entity_type='project',
                entity_id=project.id,
                old_value={'current_phase': from_phase},
                new_value={'current_phase': to_phase, 'reason': reason},
                commit=False,
            )

        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    return project, phase_log


def get_project_activity(project_id, user_id=None, limit=100):
    project = _get_project_with_access(project_id, user_id=user_id)
    try:
        limit = max(1, min(int(limit or 100), 200))
    except (TypeError, ValueError):
        limit = 100

    requirement_ids = [
        row[0]
        for row in db.session.query(ProjectRequirement.id)
        .filter(ProjectRequirement.project_id == project.id)
        .all()
    ]
    submission_ids = []
    scan_ids = []
    push_ids = []

    if requirement_ids:
        submission_ids = [
            row[0]
            for row in db.session.query(Submission.id)
            .filter(Submission.project_requirement_id.in_(requirement_ids))
            .all()
        ]
        push_ids = [
            row[0]
            for row in db.session.query(GoogleDrivePush.id)
            .filter(GoogleDrivePush.project_requirement_id.in_(requirement_ids))
            .all()
        ]

    if submission_ids:
        scan_ids = [
            row[0]
            for row in db.session.query(NlpScanResult.id)
            .filter(NlpScanResult.submission_id.in_(submission_ids))
            .all()
        ]

    filters = [
        db.and_(AuditLog.entity_type == 'project', AuditLog.entity_id == project.id),
    ]
    if requirement_ids:
        filters.append(db.and_(AuditLog.entity_type == 'project_requirement', AuditLog.entity_id.in_(requirement_ids)))
    if submission_ids:
        filters.append(db.and_(AuditLog.entity_type == 'submission', AuditLog.entity_id.in_(submission_ids)))
    if scan_ids:
        filters.append(db.and_(AuditLog.entity_type == 'nlp_scan_result', AuditLog.entity_id.in_(scan_ids)))
    if push_ids:
        filters.append(db.and_(AuditLog.entity_type == 'google_drive_push', AuditLog.entity_id.in_(push_ids)))

    logs = (
        AuditLog.query
        .filter(db.or_(*filters))
        .order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
        .limit(limit)
        .all()
    )
    return [audit_service._serialize_audit_log(log) for log in logs]


def generate_revision_report(project_id, user_id=None, template_id=None, placeholder_overrides=None):
    if user_id is not None:
        _require_role(user_id, 'admin', 'staff', 'field_officer')

    project = _get_project_with_access(project_id, user_id=user_id)
    lgu = db.session.get(LguProfile, project.lgu_id)
    program = db.session.get(Program, project.program_id)
    generated_at = datetime.now(timezone.utc)
    rows = _revision_report_rows(project)
    context = _revision_report_context(project, lgu, program, rows, generated_at)
    template_info = _get_revision_report_template(template_id)
    context = _apply_revision_report_placeholders(context, template_info, placeholder_overrides)
    content = _render_revision_report_docx(context, template_info=template_info, template_id=template_id)
    filename = f'project-{project.id}-correction-report.docx'

    with object_storage_service.storage_transaction() as storage_tx:
        stored_object = object_storage_service.store_bytes(
            content,
            source_module='revision_report',
            original_filename=filename,
            content_type=DOCX_MIME_TYPE,
            uploaded_by=user_id,
            metadata={
                'project_id': project.id,
                'lgu_id': project.lgu_id,
                'program_id': project.program_id,
            },
            tracker=storage_tx,
        )

        report = GeneratedReport(
            generated_by=user_id,
            report_type='custom',
            title=f'Correction Report - {project.project_name}',
            parameters={
                'project_id': project.id,
                'lgu_id': project.lgu_id,
                'program_id': project.program_id,
                'revision_count': len(rows),
                'template_id': template_info.get('id') if template_info else None,
                'template_file_name': template_info.get('file_name') if template_info else None,
                'placeholder_override_keys': sorted((placeholder_overrides or {}).keys()),
                'stored_object_id': stored_object.id,
            },
            file_path=stored_object.legacy_file_path or stored_object.object_key,
            file_name=filename,
            file_format='docx',
            file_size_bytes=len(content),
        )
        db.session.add(report)
        db.session.commit()
        storage_tx.commit()
    audit_service.log_action(
        user_id=user_id,
        action='generate_report',
        entity_type='generated_report',
        entity_id=report.id,
        old_value=None,
        new_value={
            'id': report.id,
            'report_type': report.report_type,
            'title': report.title,
            'project_id': project.id,
            'file_name': report.file_name,
            'file_format': report.file_format,
            'revision_count': len(rows),
            'template_id': template_info.get('id') if template_info else None,
        },
    )
    return {
        'file_name': filename,
        'content': content,
        'mimetype': DOCX_MIME_TYPE,
        'row_count': len(rows),
        'stored_object_id': stored_object.id,
    }


def generate_requirement_correction_report(requirement_id, user_id=None, transition_id=None, placeholder_overrides=None, commit=True):
    if user_id is not None:
        _require_role(user_id, 'admin', 'staff', 'field_officer')

    requirement = _get_project_requirement_with_access(requirement_id, user_id=user_id)
    project = _get_project_with_access(requirement.project_id, user_id=user_id)
    lgu = db.session.get(LguProfile, project.lgu_id)
    program = db.session.get(Program, project.program_id)
    generated_at = datetime.now(timezone.utc)

    transition = None
    if transition_id:
        transition = RequirementWorkflowTransition.query.filter_by(
            id=transition_id,
            project_requirement_id=requirement.id,
        ).first()
    if not transition:
        transition = (
            RequirementWorkflowTransition.query
            .filter_by(project_requirement_id=requirement.id, action='return_for_revision')
            .order_by(RequirementWorkflowTransition.created_at.desc(), RequirementWorkflowTransition.id.desc())
            .first()
        )

    actor_id = transition.actor_id if transition else (requirement.workflow_returned_by or user_id)
    returned_by_user = db.session.get(User, actor_id) if actor_id else None
    returned_by_display = returned_by_user.full_name if returned_by_user else (f'User #{actor_id}' if actor_id else 'Reviewer')

    returned_at = transition.created_at if transition else (requirement.workflow_returned_at or generated_at)
    return_notes = (transition.reason if transition else requirement.workflow_return_reason) or 'Returned for revision.'

    deficiencies = (
        RequirementDeficiency.query
        .filter_by(project_requirement_id=requirement.id)
        .order_by(RequirementDeficiency.created_at.asc(), RequirementDeficiency.id.asc())
        .all()
    )
    serialized_deficiencies = [
        {
            'category': requirement_workflow_service.format_deficiency_category(d.category) if d.category else 'General',
            'finding_observation': d.finding_observation,
            'corrective_instruction': d.corrective_instruction,
            'resolution_state': d.resolution_state,
        }
        for d in deficiencies
    ]

    custom_fields = requirement.custom_fields or {}
    revision_history = list(custom_fields.get('revision_history') or [])

    row = {
        'number': 1,
        'requirement_name': requirement.requirement_name or '',
        'annex_reference': requirement.annex_reference or '',
        'phase_name': requirement.phase_name or '',
        'status': _format_status_label(requirement.status),
        'latest_revision_note': return_notes,
        'returned_at': _format_report_datetime(returned_at),
        'returned_by': returned_by_display,
        'revision_count': len(revision_history) or 1,
        'latest_submission_at': '',
    }
    rows = [row]

    context = _revision_report_context(project, lgu, program, rows, generated_at)
    context['report_title'] = f'Correction Report — {requirement.requirement_name}'
    context['requirement_name'] = requirement.requirement_name
    context['annex_reference'] = requirement.annex_reference or ''
    context['phase_name'] = requirement.phase_name or ''
    context['returned_by'] = returned_by_display
    context['returned_at'] = _format_report_datetime(returned_at)
    context['return_notes'] = return_notes
    context['deficiencies'] = serialized_deficiencies

    template_info = _get_revision_report_template(None)
    context = _apply_revision_report_placeholders(context, template_info, placeholder_overrides)
    content = _render_revision_report_docx(context, template_info=template_info)
    filename = f'requirement-{requirement.id}-correction-report.docx'

    with object_storage_service.storage_transaction() as storage_tx:
        stored_object = object_storage_service.store_bytes(
            content,
            source_module='correction_report',
            original_filename=filename,
            content_type=DOCX_MIME_TYPE,
            uploaded_by=user_id,
            metadata={
                'project_id': project.id,
                'requirement_id': requirement.id,
                'transition_id': transition.id if transition else None,
            },
            tracker=storage_tx,
        )

        report = GeneratedReport(
            generated_by=user_id,
            report_type='custom',
            title=f'Correction Report - {requirement.requirement_name}',
            parameters={
                'project_id': project.id,
                'project_requirement_id': requirement.id,
                'workflow_transition_id': transition.id if transition else None,
                'stored_object_id': stored_object.id,
                'deficiencies_count': len(deficiencies),
            },
            file_path=stored_object.legacy_file_path or stored_object.object_key,
            file_name=filename,
            file_format='docx',
            file_size_bytes=len(content),
        )
        db.session.add(report)
        if commit:
            db.session.commit()
            storage_tx.commit()
        else:
            db.session.flush()

    audit_service.log_action(
        user_id=user_id,
        action='generate_report',
        entity_type='generated_report',
        entity_id=report.id,
        old_value=None,
        new_value={
            'id': report.id,
            'report_type': report.report_type,
            'title': report.title,
            'project_id': project.id,
            'project_requirement_id': requirement.id,
            'file_name': report.file_name,
            'file_format': report.file_format,
        },
        commit=commit,
    )

    return {
        'report_id': report.id,
        'file_name': filename,
        'content': content,
        'mimetype': DOCX_MIME_TYPE,
        'stored_object_id': stored_object.id,
    }


def preview_revision_report(project_id, user_id=None, template_id=None, placeholder_overrides=None):
    if user_id is not None:
        _require_role(user_id, 'admin', 'staff', 'field_officer')

    project = _get_project_with_access(project_id, user_id=user_id)
    lgu = db.session.get(LguProfile, project.lgu_id)
    program = db.session.get(Program, project.program_id)
    generated_at = datetime.now(timezone.utc)
    rows = _revision_report_rows(project)
    context = _revision_report_context(project, lgu, program, rows, generated_at)
    template_info = _get_revision_report_template(template_id)
    context = _apply_revision_report_placeholders(context, template_info, placeholder_overrides)
    content = _render_revision_report_docx(context, template_info=template_info, template_id=template_id)
    return {
        'template': _serialize_template_preview(template_info),
        'variables': document_template_service.build_template_variables(template_info, context, placeholder_overrides),
        'preview_text': document_template_service.preview_docx_text(content),
        'preview_document': document_template_service.preview_docx_document(content),
        'row_count': len(rows),
    }


def _revision_report_rows(project):
    serialized_rows = [
        _serialize_revision_requirement(requirement)
        for requirement in sorted(
            (
                item
                for item in ProjectRequirement.query.filter_by(project_id=project.id).all()
                if (item.custom_fields or {}).get('revision_history') or item.workflow_status == 'returned_for_revision'
            ),
            key=lambda item: ((item.custom_fields or {}).get('latest_revision') or {}).get('returned_at') or '',
            reverse=True,
        )
    ]
    returned_by_ids = {
        int((row.get('latest_revision') or {}).get('returned_by'))
        for row in serialized_rows
        if str((row.get('latest_revision') or {}).get('returned_by') or '').isdigit()
    }
    users = {
        user.id: user.full_name
        for user in User.query.filter(User.id.in_(returned_by_ids)).all()
    } if returned_by_ids else {}

    rows = []
    for index, row in enumerate(serialized_rows, start=1):
        latest_revision = row.get('latest_revision') or {}
        returned_by = latest_revision.get('returned_by')
        rows.append({
            'number': index,
            'requirement_name': row.get('requirement_name') or '',
            'annex_reference': row.get('annex_reference') or '',
            'phase_name': row.get('phase_name') or '',
            'status': _format_status_label(row.get('status')),
            'latest_revision_note': latest_revision.get('notes') or '',
            'returned_at': _format_report_datetime(latest_revision.get('returned_at')),
            'returned_by': users.get(int(returned_by), f'User #{returned_by}') if str(returned_by or '').isdigit() else '',
            'revision_count': row.get('revision_count') or 0,
            'latest_submission_at': _format_report_datetime(row.get('latest_submission_at')),
        })
    return rows


def _revision_report_context(project, lgu, program, rows, generated_at):
    return {
        'report_title': 'Returned for Revision Correction Report',
        'generated_at_display': _format_report_datetime(generated_at),
        'date': _format_report_datetime(generated_at),
        'project_id': project.id,
        'project_name': project.project_name,
        'source_project_code': project.source_project_code or '',
        'lgu_name': lgu.name if lgu else '',
        'program_name': program.name if program else '',
        'fiscal_year': project.fiscal_year,
        'revision_count': len(rows),
        'revision_rows': rows,
        'revision_rows_text': '\n'.join(_format_revision_report_row_text(row) for row in rows) or 'No requirements have been returned for revision.',
    }


def _get_revision_report_template(template_id=None):
    try:
        return settings_service.get_report_template('project_revision_report', template_id=template_id)
    except settings_service.SettingsError as exc:
        raise ProjectError(exc.message, exc.status_code)


def _apply_revision_report_placeholders(context, template_info=None, placeholder_overrides=None):
    merged = dict(context)
    editable_defaults = {
        'director_name': '',
        'director_title': '',
        'office_name': '',
        'prepared_by_name': '',
        'prepared_by_title': '',
        'signatory_name': '',
        'signatory_title': '',
    }
    merged.update(editable_defaults)
    editable_keys = set(editable_defaults) | {
        item.get('key')
        for item in (template_info or {}).get('placeholder_definitions', [])
        if item.get('key')
    }
    defaults = (template_info or {}).get('placeholder_defaults') or {}
    for source in (defaults, placeholder_overrides or {}):
        if not isinstance(source, dict):
            continue
        for key, value in source.items():
            cleaned_key = str(key).strip()
            if cleaned_key in editable_keys:
                merged[cleaned_key] = '' if value is None else str(value)
    if merged.get('director_name') and not merged.get('signatory_name'):
        merged['signatory_name'] = merged['director_name']
    if merged.get('director_title') and not merged.get('signatory_title'):
        merged['signatory_title'] = merged['director_title']
    return merged


def _render_revision_report_docx(context, template_info=None, template_id=None):
    target_id = template_id or (template_info.get('id') if template_info else None)
    resolved = document_template_service.resolve_document_template(
        'project_revision_report',
        template_id=target_id,
        allow_general_fallback=True,
    )
    # Tier 1 & Tier 2: Specific template
    if resolved and resolved.get('tier') in ('configured', 'seeded'):
        template_path = resolved.get('path')
        if template_path:
            try:
                return document_template_service.render_docx_template(template_path, context)
            except Exception as exc:
                logger.warning(
                    'Project revision report specific template rendering failed; falling back: %s',
                    exc,
                )

    # Tier 3: Approved General Template Fallback
    base_template_path = None
    if resolved and resolved.get('tier') == 'general_fallback':
        base_template_path = resolved.get('path')
    else:
        base_template_path = document_template_service.get_general_template_path()

    if base_template_path:
        try:
            return _render_revision_report_fallback_docx(context, base_template_path=base_template_path)
        except Exception as exc:
            logger.warning(
                'Rendering revision report on General Template failed; using raw blank fallback: %s',
                exc,
            )

    # Tier 4: Emergency raw programmatic fallback
    return _render_revision_report_fallback_docx(context, base_template_path=None)


def _serialize_template_preview(template_info):
    if not template_info:
        general_path = document_template_service.get_general_template_path()
        return {
            'id': None,
            'label': 'DILG Laguna General Letterhead Template' if general_path else 'Built-in DOCX layout',
            'file_name': 'General Template.docx' if general_path else None,
            'placeholder_definitions': [],
        }
    return {
        'id': template_info.get('id'),
        'label': template_info.get('label'),
        'file_name': template_info.get('file_name'),
        'placeholder_definitions': template_info.get('placeholder_definitions') or [],
    }


def _render_revision_report_fallback_docx(context, base_template_path=None):
    Document = _import_docx_document()
    if base_template_path and os.path.exists(base_template_path):
        doc = Document(str(base_template_path))
        # Remove the initial empty paragraph in General Template so title starts at top
        if doc.paragraphs and not doc.paragraphs[0].text.strip():
            first_p = doc.paragraphs[0]._p
            if first_p.getparent() is not None:
                first_p.getparent().remove(first_p)
    else:
        doc = Document()

    doc.add_heading(context['report_title'], level=1)
    doc.add_paragraph(f'Generated: {context["generated_at_display"]}')
    doc.add_paragraph(f'Project: {context["project_name"]}')
    doc.add_paragraph(f'LGU: {context["lgu_name"]}')
    doc.add_paragraph(f'Program: {context["program_name"]}')
    doc.add_paragraph(f'Fiscal Year: {context["fiscal_year"]}')
    if context.get('source_project_code'):
        doc.add_paragraph(f'Source Project Code: {context["source_project_code"]}')
    doc.add_paragraph(f'Returned Requirements: {context["revision_count"]}')

    rows = context.get('revision_rows') or []
    if rows:
        table = doc.add_table(rows=1, cols=9)
        try:
            table.style = 'Table Grid'
        except Exception:
            pass
        _apply_docx_table_borders(table)
        headers = [
            '#',
            'Requirement',
            'Annex',
            'Phase',
            'Status',
            'Revision Notes',
            'Returned At',
            'Returned By',
            'Revisions',
        ]
        for index, label in enumerate(headers):
            table.cell(0, index).text = label
        for item in rows:
            cells = table.add_row().cells
            cells[0].text = str(item['number'])
            cells[1].text = item['requirement_name']
            cells[2].text = item['annex_reference']
            cells[3].text = item['phase_name']
            cells[4].text = item['status']
            cells[5].text = item['latest_revision_note']
            cells[6].text = item['returned_at']
            cells[7].text = item['returned_by']
            cells[8].text = str(item['revision_count'])
    else:
        doc.add_paragraph('No requirements have been returned for revision.')

    deficiencies = context.get('deficiencies') or []
    if deficiencies:
        doc.add_paragraph('')
        doc.add_heading('Specific Findings & Corrective Instructions', level=2)
        d_table = doc.add_table(rows=1, cols=4)
        try:
            d_table.style = 'Table Grid'
        except Exception:
            pass
        _apply_docx_table_borders(d_table)
        d_headers = ['#', 'Category', 'Finding / Observation', 'Corrective Instruction']
        for index, label in enumerate(d_headers):
            d_table.cell(0, index).text = label
        for index, item in enumerate(deficiencies, 1):
            d_cells = d_table.add_row().cells
            d_cells[0].text = str(index)
            d_cells[1].text = str(item.get('category') or '')
            d_cells[2].text = str(item.get('finding_observation') or item.get('finding') or '')
            d_cells[3].text = str(item.get('corrective_instruction') or '')

    doc.add_paragraph('')
    doc.add_paragraph('Prepared by DILG Operations Workspace.')
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


ANGLE_PLACEHOLDER_ALIASES = {
    "Director's Name": 'director_name',
    'Director Name': 'director_name',
    'Regional Director': 'director_name',
    "Director's Title": 'director_title',
    'Director Title': 'director_title',
    'Office Name': 'office_name',
    'Prepared By': 'prepared_by_name',
    'Prepared By Name': 'prepared_by_name',
    'Prepared By Title': 'prepared_by_title',
    'Signatory Name': 'signatory_name',
    'Signatory Title': 'signatory_title',
}


def _replace_angle_placeholders_in_docx(content, context):
    source = BytesIO(content)
    target = BytesIO()
    try:
        with zipfile.ZipFile(source, 'r') as zin, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith('word/') and item.filename.endswith('.xml'):
                    text = data.decode('utf-8')
                    text = re.sub(r'&lt;([^<>]{1,120})&gt;', lambda match: _angle_placeholder_value(match, context), text)
                    data = text.encode('utf-8')
                zout.writestr(item, data)
    except zipfile.BadZipFile:
        return content
    target.seek(0)
    return target.read()


def _angle_placeholder_value(match, context):
    label = html.unescape(match.group(1)).strip()
    key = ANGLE_PLACEHOLDER_ALIASES.get(label) or _slug_placeholder_label(label)
    if key not in context:
        return match.group(0)
    return html.escape(str(context.get(key) or ''), quote=False)


def _slug_placeholder_label(label):
    cleaned = re.sub(r"['’]", '', label.lower())
    cleaned = re.sub(r'[^a-z0-9]+', '_', cleaned).strip('_')
    return cleaned


def _apply_docx_table_borders(table):
    try:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except Exception:
        return

    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement('w:tblPr')
        table._tbl.insert(0, tbl_pr)
    existing = tbl_pr.find(qn('w:tblBorders'))
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement('w:tblBorders')
    for edge_name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement(f'w:{edge_name}')
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '6')
        edge.set(qn('w:space'), '0')
        edge.set(qn('w:color'), '000000')
        borders.append(edge)
    tbl_pr.append(borders)


def _format_revision_report_row_text(row):
    return (
        f"{row['number']}. {row['requirement_name']} | {row['phase_name']} | "
        f"{row['status']} | {row['latest_revision_note']}"
    )


def _format_report_datetime(value):
    if not value:
        return ''
    if isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.replace('Z', '+00:00'))
        except ValueError:
            return value
    else:
        parsed = value
    return parsed.strftime('%B %d, %Y %I:%M %p')


def _format_status_label(value):
    if not value:
        return ''
    return _format_requirement_status_label(value) or str(value).replace('_', ' ').strip().title()


def _import_docxtemplate():
    try:
        from docxtpl import DocxTemplate
    except Exception as exc:
        raise ProjectError(f'docxtpl is required for DOCX report rendering: {exc}', 500)
    return DocxTemplate


def _import_docx_document():
    try:
        from docx import Document
    except Exception as exc:
        raise ProjectError(f'python-docx is required for DOCX report rendering: {exc}', 500)
    return Document


def initialize_project_requirements(project_id, user_id=None):
    project = _get_project_with_access(project_id, user_id=user_id)
    try:
        access_service.assert_project_movement_allowed(project)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)
    existing_count = ProjectRequirement.query.filter_by(project_id=project.id).count()
    if existing_count > 0:
        return existing_count

    requirement_count = _create_project_requirement_snapshots(project)
    try:
        if user_id is not None and requirement_count > 0:
            audit_service.log_action(
                user_id=user_id,
                action='create',
                entity_type='project_requirement',
                entity_id=project.id,
                old_value=None,
                new_value={'initialized_count': requirement_count},
                commit=False,
            )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    return requirement_count


def _create_project_requirement_snapshots(project):
    program = db.session.get(Program, project.program_id)
    if not program:
        raise ProjectError('Associated program not found', 404)

    effective_kind = resolve_effective_project_kind(project)
    phases = ProgramPhase.query.filter_by(program_id=program.id).order_by(ProgramPhase.phase_order).all()
    requirement_count = 0
    sort_order = 0
    for phase in phases:
        requirements_query = Requirement.query.filter_by(phase_id=phase.id)
        if project.program_cycle_id is not None:
            requirements_query = requirements_query.filter(
                db.or_(Requirement.program_cycle_id == project.program_cycle_id, Requirement.program_cycle_id.is_(None))
            )
        else:
            requirements_query = requirements_query.filter(Requirement.program_cycle_id.is_(None))
        requirements = requirements_query.order_by(Requirement.sort_order).all()
        for requirement in requirements:
            if hasattr(requirement, 'is_applicable_to') and not requirement.is_applicable_to(effective_kind):
                continue
            db.session.add(ProjectRequirement(
                project_id=project.id,
                requirement_id=requirement.id,
                requirement_name=requirement.name,
                requirement_type=requirement.requirement_type,
                document_category=requirement.document_category,
                annex_reference=requirement.annex_reference,
                is_mandatory=requirement.is_mandatory,
                has_sla=requirement.has_sla,
                sla_days=requirement.sla_days,
                phase_name=phase.name,
                status='pending',
                workflow_status='draft' if requirement.document_category else None,
                sort_order=sort_order,
            ))
            requirement_count += 1
            sort_order += 1

    return requirement_count


def reconcile_project_requirements(project_id, user_id=None, commit=True, retire_unmatched=False):
    """Reconcile requirement snapshots against the project's effective kind.

    Non-destructive:
    - Adds missing applicable requirements.
    - Deactivates (is_active=False) inapplicable requirements.
    - Historical submissions and files are preserved.
    """
    project = db.session.get(Project, project_id)
    if not project:
        raise ProjectError('Project not found', 404)

    program = db.session.get(Program, project.program_id)
    if not program:
        raise ProjectError('Associated program not found', 404)

    effective_kind = resolve_effective_project_kind(project)
    phases = ProgramPhase.query.filter_by(program_id=program.id).order_by(ProgramPhase.phase_order).all()

    existing_snapshots = ProjectRequirement.query.filter_by(project_id=project.id).all()
    snapshot_by_req_id = {s.requirement_id: s for s in existing_snapshots if s.requirement_id}

    def match_key(name, phase_name, requirement_type, document_category):
        normalize = lambda value: re.sub(r'\W+', ' ', (value or '').casefold()).strip()
        return (normalize(name), normalize(phase_name), requirement_type, document_category)

    claimed_snapshot_ids = set()
    added_count = 0
    deactivated_count = 0
    sort_order = max([s.sort_order for s in existing_snapshots] or [0]) + 1

    # 1. Identify applicable requirements across phases
    target_rows = []
    for phase in phases:
        requirements_query = Requirement.query.filter_by(phase_id=phase.id)
        if project.program_cycle_id is not None:
            requirements_query = requirements_query.filter(
                db.or_(Requirement.program_cycle_id == project.program_cycle_id, Requirement.program_cycle_id.is_(None))
            )
        else:
            requirements_query = requirements_query.filter(Requirement.program_cycle_id.is_(None))
        requirements = requirements_query.order_by(Requirement.sort_order).all()

        target_rows.extend((phase, req) for req in requirements if req.is_applicable_to(effective_kind))
    applicable_req_ids = {req.id for _, req in target_rows}
    target_keys = [match_key(req.name, phase.name, req.requirement_type, req.document_category)
                   for phase, req in target_rows]
    for phase, req in target_rows:
        if req.id not in snapshot_by_req_id:
            key = match_key(req.name, phase.name, req.requirement_type, req.document_category)
            candidates = [snap for snap in existing_snapshots
                          if snap.id not in claimed_snapshot_ids
                          and snap.requirement_id not in applicable_req_ids
                          and match_key(snap.requirement_name, snap.phase_name,
                                        snap.requirement_type, snap.document_category) == key]
            if len(candidates) == 1 and target_keys.count(key) == 1:
                snap = candidates[0]
                was_inactive = not snap.is_active
                history = list((snap.custom_fields or {}).get('requirement_reassignments') or [])
                history.append({'previous_requirement_id': snap.requirement_id,
                                'requirement_id': req.id,
                                'changed_at': datetime.now(timezone.utc).isoformat()})
                snap.custom_fields = {**(snap.custom_fields or {}), 'requirement_reassignments': history}
                snap.requirement_id = req.id
                snap.requirement = req
                snapshot_by_req_id[req.id] = snap
                snap.is_active = True
                snap.is_mandatory = req.is_mandatory
                snap.has_sla = req.has_sla
                snap.sla_days = req.sla_days
                snap.annex_reference = req.annex_reference
                snap.requirement_name = req.name
                if was_inactive and snap.status == 'waived':
                    snap.status = 'pending'
                    snap.waived_reason = None
                if not req.has_sla:
                    from app.services import sla_service
                    sla_service.resolve_deadlines_for_requirement(
                        snap.id, reason='Requirement moved to a checklist without an SLA',
                        user_id=user_id, commit=False,
                    )
                if snap.workflow_status is None and req.document_category:
                    snap.workflow_status = 'draft'
        if req.id in snapshot_by_req_id:
            claimed_snapshot_ids.add(snapshot_by_req_id[req.id].id)
        if req.id not in snapshot_by_req_id:
            # Add missing applicable requirement
            db.session.add(ProjectRequirement(
                project_id=project.id,
                requirement_id=req.id,
                requirement_name=req.name,
                requirement_type=req.requirement_type,
                document_category=req.document_category,
                annex_reference=req.annex_reference,
                is_mandatory=req.is_mandatory,
                has_sla=req.has_sla,
                sla_days=req.sla_days,
                phase_name=phase.name,
                status='pending',
                workflow_status='draft' if req.document_category else None,
                sort_order=sort_order,
                is_active=True,
            ))
            sort_order += 1
            added_count += 1
        else:
            # Re-activate if it was previously deactivated
            existing_snap = snapshot_by_req_id[req.id]
            if not existing_snap.is_active:
                existing_snap.is_active = True
                existing_snap.is_mandatory = req.is_mandatory
                if existing_snap.status == 'waived':
                    existing_snap.status = 'pending'
                    existing_snap.waived_reason = None

    # 2. Deactivate inapplicable existing snapshots
    for snap in existing_snapshots:
        is_inapplicable = retire_unmatched and snap.requirement_id is None
        if snap.requirement_id and snap.requirement_id not in applicable_req_ids:
            is_inapplicable = True
        elif snap.document_category:
            if snap.document_category in ('final_swa', 'certificate_of_final_acceptance') and effective_kind not in ('infrastructure', 'infra'):
                is_inapplicable = True
            elif snap.document_category in ('signed_distribution_list', 'acknowledgement_receipt', 'beneficiary_tracking_record') and effective_kind != 'financial_assistance':
                is_inapplicable = True
            elif snap.document_category in ('apm_template', 'non_infra_monitoring_evidence') and effective_kind == 'infrastructure':
                is_inapplicable = True

        if is_inapplicable and snap.is_active:
            has_submissions = snap.submissions.count() > 0
            snap.is_active = False
            snap.is_mandatory = False
            if not has_submissions:
                snap.status = 'waived'
                snap.waived_reason = f"Not applicable to {effective_kind} projects"
            else:
                snap.waived_reason = f"Historical submission preserved; requirement not applicable to {effective_kind} projects"
            from app.services import sla_service
            sla_service.resolve_deadlines_for_requirement(
                snap.id, reason=snap.waived_reason, user_id=user_id, commit=False,
            )
            deactivated_count += 1

    try:
        if user_id is not None:
            audit_service.log_action(
                user_id=user_id,
                action='reconcile_requirements',
                entity_type='project',
                entity_id=project.id,
                new_value={
                    'effective_kind': effective_kind,
                    'added_count': added_count,
                    'deactivated_count': deactivated_count,
                },
                commit=False,
            )
        if commit:
            db.session.commit()
        else:
            db.session.flush()
    except Exception:
        db.session.rollback()
        raise

    return {
        'project_id': project.id,
        'effective_kind': effective_kind,
        'added_count': added_count,
        'deactivated_count': deactivated_count,
    }


def change_project_kind(project_id, target_kind, reason=None, user_id=None):
    """Administratively change a project's project_kind with reconciliation and audit trail."""
    project = db.session.get(Project, project_id)
    if not project:
        raise ProjectError('Project not found', 404)

    program = db.session.get(Program, project.program_id)
    if not program:
        raise ProjectError('Associated program not found', 404)

    target_cleaned = _clean_source_value(target_kind)
    if not target_cleaned or target_cleaned not in VALID_PROJECT_KINDS:
        raise ProjectError(f"Invalid target project_kind. Must be one of: {', '.join(VALID_PROJECT_KINDS)}", 422)

    validated_kind = validate_project_kind_for_program_type(program.program_type, target_cleaned)
    old_kind = project.project_kind
    project.project_kind = validated_kind
    db.session.commit()

    reconciliation = reconcile_project_requirements(project.id, user_id=user_id)

    if user_id is not None:
        audit_service.log_action(
            user_id=user_id,
            action='change_project_kind',
            entity_type='project',
            entity_id=project.id,
            old_value={'project_kind': old_kind},
            new_value={'project_kind': validated_kind, 'reason': reason, 'reconciliation': reconciliation},
        )

    return project



def list_project_requirements(project_id, status=None, requirement_type=None, is_mandatory=None, applicability=None, user_id=None):
    _get_project_with_access(project_id, user_id=user_id)
    query = ProjectRequirement.query.filter_by(project_id=project_id)
    if status:
        query = query.filter_by(status=status)
    if requirement_type:
        query = query.filter_by(requirement_type=requirement_type)
    if is_mandatory is not None:
        query = query.filter_by(is_mandatory=is_mandatory)
    if applicability:
        query = query.filter_by(applicability=applicability)
    return [_serialize_project_requirement(requirement) for requirement in query.order_by(ProjectRequirement.sort_order).all()]


def get_project_requirement(requirement_id, user_id=None):
    return _get_project_requirement_with_access(requirement_id, user_id=user_id)


def update_project_requirement(requirement_id, data, user_id=None, allow_reviewed_reversal=False,
                               enforce_workbook_locks=True, sync_workbook=True):
    if user_id is not None:
        _require_role(user_id, 'admin', 'staff')
    requirement = _get_project_requirement_with_access(requirement_id, user_id=user_id)
    try:
        access_service.assert_project_requirement_movement_allowed(requirement)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)
    return _apply_project_requirement_update(
        requirement,
        data,
        user_id=user_id,
        allow_reviewed_reversal=allow_reviewed_reversal,
        enforce_workbook_locks=enforce_workbook_locks,
        sync_workbook=sync_workbook,
    )


def update_project_requirement_applicability(
    requirement_id,
    applicability,
    reason=None,
    project_id=None,
    user_id=None,
    commit=True,
):
    """
    Update a project requirement's applicability ('applicable' or 'not_applicable').
    Authorized for Provincial Admin ('admin') and Provincial Staff ('staff').
    """
    if user_id is not None:
        _require_role(user_id, 'admin', 'staff')

    requirement = _get_project_requirement_with_access(requirement_id, user_id=user_id)

    if project_id is not None and requirement.project_id != project_id:
        raise ProjectError('Requirement does not belong to the specified project.', 404)

    requirement = ProjectRequirement.query.filter_by(id=requirement.id).with_for_update().populate_existing().one()

    try:
        access_service.assert_project_requirement_movement_allowed(requirement)
    except access_service.AccessError as exc:
        raise ProjectError(exc.message, exc.status_code) from exc

    if not requirement.is_active:
        raise ProjectError('Cannot change applicability of an inactive historical requirement.', 409)

    if applicability not in ('applicable', 'not_applicable'):
        raise ProjectError("applicability must be either 'applicable' or 'not_applicable'.", 422)

    before = _serialize_project_requirement(requirement)

    if applicability == 'not_applicable':
        from app.services.completion_documents_service import DOCUMENT_TYPES
        if requirement.document_category in DOCUMENT_TYPES or requirement.completion_document_record:
            raise ProjectError(
                'Completion Matrix applicability is managed through the Completion controls. '
                'Fixed closeout requirements cannot be exempted through this action.',
                422,
            )

        # Protect in-flight and completed review states
        in_flight_statuses = {
            'submitted',
            'under_verification',
            'staff_approved',
            'pending_provincial_approval',
            'resubmitted',
        }
        if requirement.workflow_status in in_flight_statuses:
            raise ProjectError(
                f"Cannot mark requirement Not Applicable while document review is in flight (current workflow status: '{requirement.workflow_status}').",
                409,
            )
        if requirement.workflow_status == 'approved':
            raise ProjectError(
                'Cannot mark requirement Not Applicable after it has been approved.',
                409,
            )
        if requirement.workflow_status == 'returned_for_revision':
            raise ProjectError(
                'Cannot mark requirement Not Applicable while returned for revision. Resolve or cancel the revision first.',
                409,
            )
        if requirement.status in REVIEWED_REQ_STATUSES or requirement.status == 'under_review':
            raise ProjectError(
                f"Cannot mark requirement Not Applicable when status is '{requirement.status}'.",
                409,
            )

        open_deficiencies = RequirementDeficiency.query.filter_by(
            project_requirement_id=requirement.id,
            resolution_state='open',
        ).count()
        if open_deficiencies > 0:
            raise ProjectError(
                f'Cannot mark requirement Not Applicable with {open_deficiencies} open deficiency item(s).',
                409,
            )

        if Submission.query.filter_by(project_requirement_id=requirement.id).first() is not None:
            raise ProjectError(
                'Cannot mark requirement Not Applicable when submissions exist. Resolve the recorded evidence first.',
                409,
            )

        trimmed_reason = (reason or '').strip()
        if not trimmed_reason:
            raise ProjectError('A reason is required when marking a requirement as Not Applicable.', 422)

        import re
        words = re.findall(r'\b[a-z]+\b', trimmed_reason.lower())
        denylist = {'n/a', 'na', 'none', 'nil', 'tbd', 'test', 'x', 'nothing', 'blank'}
        valid_words = [
            w for w in words
            if w not in denylist and len(w) >= 3 and any(v in w for v in 'aeiouy')
        ]
        if len(valid_words) < 3:
            raise ProjectError(
                'Reason must contain at least 3 meaningful words (not placeholders or keyboard mashing).',
                422,
            )

    try:
        if applicability == 'not_applicable':
            now = datetime.now(timezone.utc)
            requirement.applicability = 'not_applicable'
            requirement.applicability_reason = trimmed_reason
            requirement.applicability_set_by = user_id
            requirement.applicability_set_at = now

            # Resolve active SLA deadlines
            from app.services import sla_service
            sla_service.resolve_deadlines_for_requirement(
                requirement.id,
                reason=f'Resolved automatically: requirement marked Not Applicable ({trimmed_reason})',
                user_id=user_id,
                commit=False,
            )
        else:
            # Restoring 'applicable'
            requirement.applicability = 'applicable'
            requirement.applicability_reason = None
            requirement.applicability_set_by = None
            requirement.applicability_set_at = None

        if user_id is not None:
            audit_service.log_action(
                user_id=user_id,
                action='update_requirement_applicability',
                entity_type='project_requirement',
                entity_id=requirement.id,
                old_value={
                    'project_id': requirement.project_id,
                    'project_requirement_id': requirement.id,
                    'applicability': before.get('applicability'),
                    'applicability_reason': before.get('applicability_reason'),
                    'applicability_set_by': before.get('applicability_set_by'),
                    'applicability_set_at': before.get('applicability_set_at'),
                },
                new_value={
                    'project_id': requirement.project_id,
                    'project_requirement_id': requirement.id,
                    'applicability': requirement.applicability,
                    'applicability_reason': requirement.applicability_reason,
                    'applicability_set_by': requirement.applicability_set_by,
                    'applicability_set_at': requirement.applicability_set_at.isoformat() if requirement.applicability_set_at else None,
                    'operation': 'set_not_applicable' if requirement.applicability == 'not_applicable' else 'restore_applicable',
                    'decision_reason': requirement.applicability_reason if requirement.applicability == 'not_applicable' else before.get('applicability_reason'),
                },
                commit=False,
            )

        if commit:
            db.session.commit()
    except Exception:
        if commit:
            db.session.rollback()
        raise

    return requirement


def return_requirement_for_revision(requirement_id, review_notes, user_id=None):
    if user_id is not None:
        _require_role(user_id, 'admin', 'staff')
    requirement = _get_project_requirement_with_access(requirement_id, user_id=user_id)
    try:
        access_service.assert_project_requirement_movement_allowed(requirement)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)
    notes = str(review_notes or '').strip()
    if not notes:
        raise ProjectError('Revision notes are required when returning a requirement for correction.', 422)

    try:
        submission_service.assert_documentary_evidence(requirement.id, 'under_review')
    except submission_service.SubmissionError as exc:
        raise ProjectError(exc.message, exc.status_code) from exc

    # If this requirement belongs to the shared document workflow (document_category set),
    # delegate to requirement_workflow_service so that:
    # 1. A canonical RequirementWorkflowTransition record is created in the transition ledger.
    # 2. A RequirementDeficiency record is created (using 'other_documented_defect' for free-form notes).
    # 3. State machine guards (RETURNABLE_FROM) are enforced.
    # 4. Correction reports and workflow notifications are dispatched.
    if requirement.document_category:
        from app.services import requirement_workflow_service
        try:
            # If currently draft (or None), submit first per P2-02 intake triage contract
            if requirement.workflow_status in (None, 'draft'):
                requirement_workflow_service.submit_requirement(
                    requirement.id,
                    user_id=user_id,
                    note='Auto-submitted for review prior to revision return.',
                )
            return requirement_workflow_service.return_requirement(
                requirement.id,
                user_id=user_id,
                reason=notes,
            )
        except requirement_workflow_service.WorkflowError as exc:
            raise ProjectError(exc.message, exc.status_code) from exc

    before = _serialize_project_requirement(requirement)
    revision_entry = {
        'notes': notes,
        'returned_at': datetime.now(timezone.utc).isoformat(),
        'returned_by': user_id,
    }
    custom_fields = dict(requirement.custom_fields or {})
    revision_history = list(custom_fields.get('revision_history') or [])
    revision_history.append(revision_entry)
    custom_fields['latest_revision'] = revision_entry
    custom_fields['revision_history'] = revision_history

    requirement.custom_fields = custom_fields
    requirement.status = 'under_review'
    requirement.reviewed_at = None
    requirement.workflow_status = 'returned_for_revision'
    requirement.workflow_return_reason = notes
    requirement.workflow_returned_by = user_id
    requirement.workflow_returned_at = datetime.now(timezone.utc)
    if user_id is not None:
        requirement.last_action_by = user_id

    try:
        if user_id is not None:
            audit_service.log_action(
                user_id=user_id,
                action='return_for_revision',
                entity_type='project_requirement',
                entity_id=requirement.id,
                old_value=before,
                new_value=_serialize_project_requirement(requirement),
                commit=False,
            )
        from app.services import workbook_sync_service
        workbook_sync_service.sync_requirement_status(requirement, user_id=user_id, commit=False)
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    if user_id is not None:
        # A return is an event even when the compliance status stays Under Review.
        try:
            from app.services import notification_service
            notification_service.notify_workflow_transition(
                project_requirement_id=requirement.id,
                event='returned',
                actor_user_id=user_id,
            )
        except Exception as exc:
            # Log but don't fail the operation if notification fails
            current_app.logger.warning(
                'Failed to create field officer notification for requirement %s: %s',
                requirement.id,
                exc,
            )
    
    return requirement


def create_change_status_request(requirement_id, requested_status, justification, user_id, auto_approve=False):
    _require_role(user_id, 'staff', 'admin')
    requirement = _get_project_requirement_with_access(requirement_id, user_id=user_id)
    try:
        access_service.assert_project_requirement_movement_allowed(requirement)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)

    if requested_status not in VALID_REQ_STATUSES:
        raise ProjectError(f'Invalid status. Must be one of: {", ".join(VALID_REQ_STATUSES)}', 422)
    if requirement.status not in REVIEWED_REQ_STATUSES:
        raise ProjectError('Change requests are only required for reversing a reviewed requirement state.', 422)
    if requested_status == requirement.status:
        raise ProjectError('Requested status must be different from the current status.', 422)

    existing_request = ChangeStatusRequest.query.filter_by(
        project_requirement_id=requirement.id,
        status='pending',
    ).first()
    if existing_request:
        if auto_approve:
            _require_role(user_id, 'admin')
            return review_change_status_request(
                existing_request.id,
                decision='approved',
                review_notes=justification or 'Approved by administrator.',
                reviewer_id=user_id,
            )
        raise ProjectError('A pending status change request already exists for this requirement.', 409)

    change_request = ChangeStatusRequest(
        project_requirement_id=requirement.id,
        requested_by=user_id,
        current_status=requirement.status,
        requested_status=requested_status,
        justification=justification,
        status='pending',
    )
    db.session.add(change_request)
    db.session.flush()

    try:
        if user_id is not None:
            audit_service.log_action(
                user_id=user_id,
                action='create',
                entity_type='change_status_request',
                entity_id=change_request.id,
                old_value=None,
                new_value=_serialize_change_status_request(change_request),
                commit=False,
            )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    if auto_approve:
        _require_role(user_id, 'admin')
        return review_change_status_request(
            change_request.id,
            decision='approved',
            review_notes=justification or 'Approved by administrator.',
            reviewer_id=user_id,
        )

    return change_request


def list_change_status_requests(status=None, project_id=None, fiscal_year=None, user_id=None):
    query = ChangeStatusRequest.query.join(
        ProjectRequirement, ChangeStatusRequest.project_requirement_id == ProjectRequirement.id
    ).join(
        Project, ProjectRequirement.project_id == Project.id
    ).join(
        LguProfile, Project.lgu_id == LguProfile.id
    )

    if user_id is not None:
        role = access_service.get_user_role(user_id)
        scope = access_service.get_access_scope(user_id)
        if not scope['unrestricted']:
            if scope['lgu_ids']:
                query = query.filter(LguProfile.id.in_(scope['lgu_ids']))
            elif scope['cluster_ids']:
                query = query.filter(LguProfile.cluster_id.in_(scope['cluster_ids']))
            else:
                return []
        if role != 'admin':
            query = query.filter(ChangeStatusRequest.requested_by == user_id)

    if status:
        query = query.filter(ChangeStatusRequest.status == status)
    if project_id is not None:
        query = query.filter(Project.id == project_id)
    query = apply_project_fiscal_year(query, fiscal_year)

    return [_serialize_change_status_request(item) for item in query.order_by(ChangeStatusRequest.created_at.desc()).all()]


def review_change_status_request(request_id, decision, review_notes=None, reviewer_id=None):
    if reviewer_id is not None:
        _require_role(reviewer_id, 'admin')

    change_request = db.session.get(ChangeStatusRequest, request_id)
    if not change_request:
        raise ProjectError('Change status request not found', 404)
    try:
        access_service.assert_project_requirement_movement_allowed(change_request.project_requirement_id)
    except access_service.AccessError as exc:
        raise _translate_access_error(exc)
    if change_request.status != 'pending':
        raise ProjectError('This change status request has already been reviewed.', 409)
    if decision not in ('approved', 'rejected'):
        raise ProjectError('Decision must be either approved or rejected.', 422)

    before = _serialize_change_status_request(change_request)
    requirement_audit_before = None
    requirement = None
    if decision == 'approved':
        update_payload = {'status': change_request.requested_status}
        if change_request.requested_status == 'waived':
            update_payload['waived_reason'] = change_request.justification
        requirement = _get_project_requirement_with_access(change_request.project_requirement_id, user_id=reviewer_id)
        requirement_audit_before = _serialize_project_requirement(requirement)
        _apply_project_requirement_update(
            requirement,
            update_payload,
            user_id=reviewer_id,
            allow_reviewed_reversal=True,
            commit=False,
        )

    change_request.status = decision
    change_request.review_decision = decision
    change_request.review_notes = review_notes
    change_request.reviewed_by = reviewer_id
    change_request.reviewed_at = datetime.now(timezone.utc)

    try:
        if reviewer_id is not None:
            audit_service.log_action(
                user_id=reviewer_id,
                action='review',
                entity_type='change_status_request',
                entity_id=change_request.id,
                old_value=before,
                new_value=_serialize_change_status_request(change_request),
                commit=False,
            )
            if decision == 'approved' and requirement is not None:
                audit_service.log_action(
                    user_id=reviewer_id,
                    action='status_change',
                    entity_type='project_requirement',
                    entity_id=requirement.id,
                    old_value=requirement_audit_before,
                    new_value=_serialize_project_requirement(requirement),
                    commit=False,
                )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return change_request


def get_project_detail(project_id, user_id=None):
    project = _get_project_with_access(project_id, user_id=user_id)
    result = _serialize_project(project)
    from app.services import workbook_sync_service
    result['workbook_field_locks'] = workbook_sync_service.project_field_lock_summary(project.id)

    lgu = db.session.get(LguProfile, project.lgu_id)
    if lgu:
        result['lgu'] = {
            'id': lgu.id,
            'name': lgu.name,
            'municipality': lgu.municipality,
            'province': lgu.province,
            'cluster_id': lgu.cluster_id,
            'cluster_name': lgu.cluster.name if lgu.cluster else None,
        }

    program = db.session.get(Program, project.program_id)
    if program:
        result['program'] = {
            'id': program.id,
            'name': program.name,
            'code': program.code,
            'program_type': program.program_type,
            'is_active': program.is_active,
        }

    cycle = db.session.get(ProgramCycle, project.program_cycle_id) if project.program_cycle_id else None
    result['program_cycle'] = {
        'id': cycle.id,
        'name': cycle.name,
        'fiscal_year': cycle.fiscal_year,
        'cycle_code': cycle.cycle_code,
    } if cycle else None

    all_requirements = ProjectRequirement.query.filter_by(project_id=project.id).all()
    requirements = [item for item in all_requirements if item.is_active]
    applicable_requirements = [item for item in requirements if getattr(item, 'applicability', 'applicable') == 'applicable']
    not_applicable_requirements = [item for item in requirements if getattr(item, 'applicability', 'applicable') == 'not_applicable']
    result['requirement_summary'] = {
        'total': len(requirements),
        'applicable': len(applicable_requirements),
        'pending': sum(1 for item in applicable_requirements if item.status == 'pending'),
        'under_review': sum(1 for item in applicable_requirements if item.status == 'under_review'),
        'compliant': sum(1 for item in applicable_requirements if item.status == 'compliant'),
        'non_compliant': sum(1 for item in applicable_requirements if item.status == 'non_compliant'),
        'waived': sum(1 for item in applicable_requirements if item.status == 'waived'),
        'not_applicable': len(not_applicable_requirements),
        'mandatory': sum(1 for item in applicable_requirements if item.is_mandatory),
    }
    result['pending_change_requests'] = ChangeStatusRequest.query.join(
        ProjectRequirement, ChangeStatusRequest.project_requirement_id == ProjectRequirement.id
    ).filter(
        ProjectRequirement.project_id == project.id,
        ProjectRequirement.is_active == True,
        ChangeStatusRequest.status == 'pending',
        ).count()

    active_returned = [
        item for item in requirements
        if item.workflow_status == 'returned_for_revision'
    ]
    historical_returned = [
        item for item in all_requirements
        if (item.custom_fields or {}).get('revision_history') or item.workflow_status == 'returned_for_revision'
    ]

    result['revision_requirements'] = [
        _serialize_revision_requirement(requirement)
        for requirement in sorted(
            active_returned,
            key=lambda item: ((item.custom_fields or {}).get('latest_revision') or {}).get('returned_at') or '',
            reverse=True,
        )
    ]
    result['historical_revision_requirements'] = [
        _serialize_revision_requirement(requirement)
        for requirement in sorted(
            historical_returned,
            key=lambda item: ((item.custom_fields or {}).get('latest_revision') or {}).get('returned_at') or '',
            reverse=True,
        )
    ]

    result['phase_history'] = [
        _serialize_phase_log(log)
        for log in ProjectPhaseLog.query.filter_by(project_id=project.id).order_by(ProjectPhaseLog.transitioned_at).all()
    ]

    current_phase_label = PHASE_NAME_BY_KEY.get(project.current_phase)
    if current_phase_label:
        blocking_requirements = ProjectRequirement.query.filter(
            ProjectRequirement.project_id == project.id,
            ProjectRequirement.phase_name == current_phase_label,
            ProjectRequirement.is_mandatory == True,
            ProjectRequirement.is_active == True,
            ProjectRequirement.status != 'compliant',
        ).order_by(ProjectRequirement.sort_order.asc()).all()
        effective_kind = resolve_effective_project_kind(project) if project else 'infrastructure'
        applicable_blocking = []
        for req in blocking_requirements:
            if getattr(req, 'applicability', 'applicable') == 'not_applicable':
                continue
            if req.requirement and hasattr(req.requirement, 'is_applicable_to'):
                if not req.requirement.is_applicable_to(effective_kind):
                    continue
            if _is_valid_waived_requirement(req):
                continue
            applicable_blocking.append(req)

        result['phase_gate'] = {
            'current_phase_label': current_phase_label,
            'can_advance': len(applicable_blocking) == 0,
            'blocking_requirements': [
                {
                    'id': requirement.id,
                    'requirement_name': requirement.requirement_name,
                    'status': requirement.status,
                }
                for requirement in applicable_blocking
            ],
        }

    return result


def _apply_project_requirement_update(requirement, data, user_id=None,
                                      allow_reviewed_reversal=False, commit=True,
                                      enforce_workbook_locks=True, sync_workbook=True):
    before = _serialize_project_requirement(requirement)
    old_status = requirement.status  # Track old status for notification

    if 'applicability' in data:
        raise ProjectError(
            'applicability cannot be modified directly through generic update. Use the dedicated applicability action.',
            422,
        )

    if getattr(requirement, 'applicability', 'applicable') == 'not_applicable' and 'status' in data:
        raise ProjectError(
            'Cannot update status of a requirement marked Not Applicable. Restore applicability first.',
            409,
        )

    if 'workflow_status' in data:
        if requirement.document_category:
            raise ProjectError(
                'Requirements with a managed document workflow cannot have workflow_status modified directly through generic update. Use the authoritative documentary workflow actions.',
                422,
            )
        raise ProjectError(
            'workflow_status cannot be modified directly through generic update. Use the authoritative documentary workflow actions.',
            422,
        )

    if requirement.document_category:
        if 'status' in data and not allow_reviewed_reversal:
            raise ProjectError(
                'Requirements with a managed document workflow cannot have status modified directly through generic update. Use the authoritative documentary workflow actions.',
                422,
            )

    target_status = data.get('status', requirement.status)
    if target_status not in VALID_REQ_STATUSES:
        raise ProjectError(f'Invalid status. Must be one of: {", ".join(VALID_REQ_STATUSES)}', 422)

    if requirement.status in REVIEWED_REQ_STATUSES and target_status != requirement.status and not allow_reviewed_reversal:
        raise ProjectError(
            'Reviewed requirement statuses can only be reversed through an approved change status request.',
            409,
        )

    if 'status' in data and target_status == 'compliant':
        open_count = RequirementDeficiency.query.filter_by(
            project_requirement_id=requirement.id,
            resolution_state='open',
        ).count()
        if open_count > 0:
            raise ProjectError(
                f'Cannot approve requirement: {open_count} open deficiency item(s) must be resolved first.',
                422,
            )

        if user_id is not None:
            role = access_service.get_user_role(user_id)
            if role != 'admin':
                raise ProjectError(
                    'Staff cannot directly approve requirement compliance. Final approval requires Provincial Admin.',
                    403,
                )

    if 'status' in data and target_status == 'non_compliant' and target_status != requirement.status:
        assert_requirement_can_be_non_compliant(requirement.id)

    if 'status' in data and target_status != requirement.status:
        try:
            submission_service.assert_documentary_evidence(requirement.id, target_status)
        except submission_service.SubmissionError as exc:
            raise ProjectError(exc.message, exc.status_code) from exc

    status_changed = 'status' in data and target_status != requirement.status
    if enforce_workbook_locks and status_changed:
        try:
            from app.services import workbook_sync_service
            workbook_sync_service.assert_requirement_status_unlocked(requirement)
        except workbook_sync_service.WorkbookSyncError as exc:
            raise ProjectError(exc.message, exc.status_code) from exc

    if 'status' in data:
        requirement.status = target_status
        if requirement.document_category:
            if target_status == 'compliant':
                requirement.workflow_status = 'approved'
                requirement.workflow_approved_at = datetime.now(timezone.utc)
                requirement.workflow_return_reason = None
            elif target_status == 'waived':
                requirement.workflow_status = 'approved'
                requirement.workflow_return_reason = None
            elif target_status == 'under_review':
                if not requirement.workflow_status or requirement.workflow_status == 'draft':
                    requirement.workflow_status = 'submitted'
            elif target_status == 'pending':
                requirement.workflow_status = 'draft'
        else:
            if target_status in ('compliant', 'waived'):
                requirement.workflow_status = None
                requirement.workflow_return_reason = None
    if 'waived_reason' in data:
        if target_status != 'waived':
            raise ProjectError('Waived reason can only be set when the requirement status is waived.', 422)
        requirement.waived_reason = data['waived_reason']
    elif target_status != 'waived':
        requirement.waived_reason = None

    if target_status == 'waived' and status_changed:
        open_deficiencies = RequirementDeficiency.query.filter_by(
            project_requirement_id=requirement.id,
            resolution_state='open',
        ).all()
        now_dt = datetime.now(timezone.utc)
        effective_reason = requirement.waived_reason or (data.get('waived_reason') if isinstance(data, dict) else None)
        waiver_note = (
            f"Waived automatically upon requirement waiver: {effective_reason}"
            if effective_reason
            else "Waived automatically upon requirement waiver."
        )
        for d in open_deficiencies:
            d.resolution_state = 'waived'
            d.resolution_notes = waiver_note
            d.resolved_by_user_id = user_id
            d.resolved_at = now_dt
            if user_id is not None:
                audit_service.log_action(
                    user_id=user_id,
                    action='deficiency_waived',
                    entity_type='requirement_deficiency',
                    entity_id=d.id,
                    old_value={'resolution_state': 'open'},
                    new_value={
                        'id': d.id,
                        'project_requirement_id': requirement.id,
                        'resolution_state': 'waived',
                        'resolution_notes': waiver_note,
                        'resolved_by_user_id': user_id,
                        'resolved_at': now_dt.isoformat(),
                    },
                    commit=False,
                )

    if 'is_active' in data:
        requirement.is_active = data['is_active']
    if 'custom_fields' in data:
        requirement.custom_fields = data['custom_fields']

    if user_id is not None:
        requirement.last_action_by = user_id
    if target_status in REVIEWED_REQ_STATUSES:
        requirement.reviewed_at = datetime.now(timezone.utc)
    elif 'status' in data:
        requirement.reviewed_at = None

    if sync_workbook and status_changed:
        try:
            from app.services import workbook_sync_service
            workbook_sync_service.sync_requirement_status(requirement, user_id=user_id, commit=False)
        except workbook_sync_service.WorkbookSyncError as exc:
            raise ProjectError(exc.message, exc.status_code) from exc

    if requirement.status == 'compliant' or _is_valid_waived_requirement(requirement):
        from app.services import sla_service
        sync_reason = (
            'Resolved automatically upon requirement compliance.'
            if requirement.status == 'compliant'
            else f'Resolved automatically upon valid requirement waiver: {requirement.waived_reason}'
        )
        sla_service.resolve_deadlines_for_requirement(
            requirement.id,
            reason=sync_reason,
            user_id=user_id,
            commit=False,
        )

    if commit:
        try:
            if user_id is not None:
                audit_service.log_action(
                    user_id=user_id,
                    action='status_change' if 'status' in data else 'update',
                    entity_type='project_requirement',
                    entity_id=requirement.id,
                    old_value=before,
                    new_value=_serialize_project_requirement(requirement),
                    commit=False,
                )
            db.session.commit()
        except Exception:
            db.session.rollback()
            raise
        
        # Notify field officer when staff changes document status
        if status_changed and user_id is not None:
            try:
                from app.services import notification_service
                notification_service.notify_field_officer_on_status_change(
                    project_requirement_id=requirement.id,
                    old_status=old_status,
                    new_status=target_status,
                    reviewer_user_id=user_id,
                )
            except Exception as exc:
                # Log but don't fail the status update if notification fails
                current_app.logger.warning(
                    'Failed to create field officer notification for requirement %s: %s',
                    requirement.id,
                    exc,
                )
    else:
        db.session.flush()

    return requirement


def _requirement_has_overdue_deadline(requirement_id):
    return requirement_has_overdue_deadline(requirement_id)


def requirement_has_overdue_deadline(requirement_id):
    return (
        SlaDeadline.query
        .filter(
            SlaDeadline.project_requirement_id == requirement_id,
            SlaDeadline.is_resolved == False,
            SlaDeadline.deadline_type.in_(FORMAL_COMPLIANCE_DEADLINE_TYPES),
            SlaDeadline.deadline_date < datetime.now(timezone.utc),
        )
        .first()
        is not None
    )


def assert_requirement_can_be_non_compliant(requirement_id, status_code=409):
    if not requirement_has_overdue_deadline(requirement_id):
        raise ProjectError(NON_COMPLIANT_REQUIRES_OVERDUE_DEADLINE_MESSAGE, status_code)


def _serialize_requirement_active_deadline(requirement_id):
    deadline = (
        SlaDeadline.query
        .filter(
            SlaDeadline.project_requirement_id == requirement_id,
            SlaDeadline.is_resolved == False,
        )
        .order_by(SlaDeadline.deadline_date.asc(), SlaDeadline.id.asc())
        .first()
    )
    if not deadline:
        return None

    now = datetime.now(timezone.utc)
    deadline_at = deadline.deadline_date
    if deadline_at and deadline_at.tzinfo is None:
        deadline_at = deadline_at.replace(tzinfo=timezone.utc)
    elif deadline_at:
        deadline_at = deadline_at.astimezone(timezone.utc)

    return {
        'id': deadline.id,
        'deadline_type': deadline.deadline_type,
        'trigger_date': deadline.trigger_date.isoformat() if deadline.trigger_date else None,
        'deadline_date': deadline.deadline_date.isoformat() if deadline.deadline_date else None,
        'sla_working_days': deadline.sla_working_days,
        'is_overdue': bool(deadline_at and deadline_at < now),
        'escalation_triggered': deadline.escalation_triggered,
    }


def _is_valid_waived_requirement(requirement):
    if requirement.status != 'waived':
        return False
    if not requirement.waived_reason or not str(requirement.waived_reason).strip():
        return False
    if requirement.document_category and requirement.workflow_status not in ('approved', None):
        return False
    return True


is_valid_waived_requirement = _is_valid_waived_requirement


def _assert_phase_gate(project_id, current_phase):
    phase_label = PHASE_NAME_BY_KEY.get(current_phase)
    if not phase_label:
        return

    project = db.session.get(Project, project_id)
    effective_kind = resolve_effective_project_kind(project) if project else 'infrastructure'

    blocking_requirements = ProjectRequirement.query.filter(
        ProjectRequirement.project_id == project_id,
        ProjectRequirement.phase_name == phase_label,
        ProjectRequirement.is_mandatory == True,
        ProjectRequirement.is_active == True,
        ProjectRequirement.status != 'compliant',
    ).order_by(ProjectRequirement.sort_order.asc()).all()
    if not blocking_requirements:
        return

    applicable_blocking = []
    for item in blocking_requirements:
        if getattr(item, 'applicability', 'applicable') == 'not_applicable':
            continue
        if item.requirement and hasattr(item.requirement, 'is_applicable_to'):
            if not item.requirement.is_applicable_to(effective_kind):
                continue
        if _is_valid_waived_requirement(item):
            continue
        applicable_blocking.append(item)

    if not applicable_blocking:
        return

    names = ', '.join(item.requirement_name for item in applicable_blocking[:3])
    if len(applicable_blocking) > 3:
        names = f'{names}, and {len(applicable_blocking) - 3} more'
    raise ProjectError(
        f'Cannot advance project from {phase_label} until all mandatory requirements are compliant. Outstanding: {names}.',
        422,
    )


def _serialize_project(project):
    cycle = db.session.get(ProgramCycle, project.program_cycle_id) if project.program_cycle_id else None
    program = project.program
    lgu = project.lgu
    return {
        'id': project.id,
        'lgu_id': project.lgu_id,
        'lgu_name': lgu.name if lgu else None,
        'program_id': project.program_id,
        'program_name': program.name if program else None,
        'program_type': program.program_type if program else None,
        'program_is_active': bool(program.is_active) if program is not None else True,
        'program_cycle_id': project.program_cycle_id,
        'program_cycle_name': cycle.name if cycle else None,
        'project_name': project.project_name,
        'fiscal_year': project.fiscal_year,
        'source_project_code': project.source_project_code,
        'source_row_key': project.source_row_key,
        'source_sheet_name': project.source_sheet_name,
        'current_phase': project.current_phase,
        'overall_status': project.overall_status,
        'budget_amount': float(project.budget_amount) if project.budget_amount else None,
        'target_start_date': project.target_start_date.isoformat() if project.target_start_date else None,
        'target_end_date': project.target_end_date.isoformat() if project.target_end_date else None,
        'actual_start_date': project.actual_start_date.isoformat() if project.actual_start_date else None,
        'actual_end_date': project.actual_end_date.isoformat() if project.actual_end_date else None,
        'physical_completion_pct': float(project.physical_completion_pct) if project.physical_completion_pct else 0,
        'project_kind': project.project_kind,
        'monitoring_applicability': project.monitoring_applicability,
        'qaar_category': project.qaar_category,
        'subaybayan_project_code': project.subaybayan_project_code,
        'subaybayan_url': project.subaybayan_url,
        'nadai_received_at': project.nadai_received_at.isoformat() if project.nadai_received_at else None,
        'subaybayan_enrolled_at': project.subaybayan_enrolled_at.isoformat() if project.subaybayan_enrolled_at else None,
        'subaybayan_enrollment_status': project.subaybayan_enrollment_status,
        'notes': project.notes,
        'custom_fields': project.custom_fields or {},
        'created_at': project.created_at.isoformat() if project.created_at else None,
        'updated_at': project.updated_at.isoformat() if project.updated_at else None,
    }


def _sanitize_custom_fields_for_field_officer(custom_fields):
    if not isinstance(custom_fields, dict):
        return {}
    sanitized = {}
    if 'latest_revision' in custom_fields and isinstance(custom_fields['latest_revision'], dict):
        rev = dict(custom_fields['latest_revision'])
        rev.pop('stored_object_id', None)
        sanitized['latest_revision'] = rev
    if 'revision_history' in custom_fields and isinstance(custom_fields['revision_history'], list):
        clean_history = []
        for item in custom_fields['revision_history']:
            if isinstance(item, dict):
                clean_item = dict(item)
                clean_item.pop('stored_object_id', None)
                clean_history.append(clean_item)
            else:
                clean_history.append(item)
        sanitized['revision_history'] = clean_history
    return sanitized


def _serialize_project_requirement(requirement, user_role=None):
    actor = db.session.get(User, requirement.last_action_by) if requirement.last_action_by else None
    active_deadline = _serialize_requirement_active_deadline(requirement.id)
    latest_submission = (
        Submission.query
        .filter_by(project_requirement_id=requirement.id)
        .order_by(Submission.created_at.desc(), Submission.id.desc())
        .first()
    )
    evidence_exists = (
        db.session.query(DocumentFile.id)
        .join(Submission, DocumentFile.submission_id == Submission.id)
        .filter(Submission.project_requirement_id == requirement.id)
        .first()
        is not None
    )
    custom_fields = requirement.custom_fields or {}
    if user_role == 'field_officer':
        custom_fields = _sanitize_custom_fields_for_field_officer(custom_fields)
    is_returned = requirement.workflow_status == 'returned_for_revision'
    latest_rev = custom_fields.get('latest_revision')
    current_return = None
    if is_returned:
        current_return = {
            'notes': requirement.workflow_return_reason or (latest_rev.get('notes') if latest_rev else ''),
            'returned_at': requirement.workflow_returned_at.isoformat() if requirement.workflow_returned_at else (latest_rev.get('returned_at') if latest_rev else None),
            'returned_by': requirement.workflow_returned_by or (latest_rev.get('returned_by') if latest_rev else None),
            'report_id': latest_rev.get('report_id') if latest_rev else None,
            'report_file_name': latest_rev.get('report_file_name') if latest_rev else None,
            'has_report': bool(latest_rev and (latest_rev.get('report_id') or latest_rev.get('stored_object_id'))),
            'deficiencies_count': latest_rev.get('deficiencies_count', 0) if latest_rev else 0,
        }
    return {
        'id': requirement.id,
        'project_id': requirement.project_id,
        'requirement_id': requirement.requirement_id,
        'requirement_name': requirement.requirement_name,
        'requirement_type': requirement.requirement_type,
        'document_category': requirement.document_category,
        'annex_reference': requirement.annex_reference,
        'is_mandatory': requirement.is_mandatory,
        'has_sla': requirement.has_sla,
        'sla_days': requirement.sla_days,
        'active_deadline': active_deadline,
        'has_active_deadline': active_deadline is not None,
        'phase_name': requirement.phase_name,
        'status': requirement.status,
        'workflow_status': requirement.workflow_status,
        'workflow_submitted_by': requirement.workflow_submitted_by,
        'workflow_submitted_at': requirement.workflow_submitted_at.isoformat() if requirement.workflow_submitted_at else None,
        'workflow_staff_reviewer_id': requirement.workflow_staff_reviewer_id,
        'workflow_staff_approved_at': requirement.workflow_staff_approved_at.isoformat() if requirement.workflow_staff_approved_at else None,
        'workflow_provincial_reviewer_id': requirement.workflow_provincial_reviewer_id,
        'workflow_approved_at': requirement.workflow_approved_at.isoformat() if requirement.workflow_approved_at else None,
        'workflow_return_reason': requirement.workflow_return_reason,
        'workflow_returned_by': requirement.workflow_returned_by,
        'workflow_returned_at': requirement.workflow_returned_at.isoformat() if requirement.workflow_returned_at else None,
        'is_returned': is_returned,
        'current_return': current_return,
        'return_history': custom_fields.get('revision_history', []),
        'latest_revision': latest_rev,
        'sort_order': requirement.sort_order,
        'is_active': requirement.is_active,
        'program_is_active': bool(requirement.project.program.is_active) if requirement.project and requirement.project.program else True,
        'waived_reason': requirement.waived_reason,
        'applicability': getattr(requirement, 'applicability', 'applicable') or 'applicable',
        'applicability_reason': requirement.applicability_reason,
        'applicability_set_by': requirement.applicability_set_by,
        'applicability_set_by_name': (
            requirement.applicability_setter.full_name
            if getattr(requirement, 'applicability_setter', None)
            else (
                db.session.get(User, requirement.applicability_set_by).full_name
                if requirement.applicability_set_by
                else None
            )
        ),
        'applicability_set_at': requirement.applicability_set_at.isoformat() if requirement.applicability_set_at else None,
        'custom_fields': custom_fields,
        'last_action_by': requirement.last_action_by,
        'last_action_by_name': actor.full_name if actor else None,
        'reviewed_at': requirement.reviewed_at.isoformat() if requirement.reviewed_at else None,
        'has_documentary_evidence': evidence_exists,
        'submission_count': requirement.submissions.count(),
        'current_version_number': latest_submission.current_version if latest_submission else 1,
        'latest_submission_at': latest_submission.created_at.isoformat() if latest_submission and latest_submission.created_at else None,
        'created_at': requirement.created_at.isoformat() if requirement.created_at else None,
        'updated_at': requirement.updated_at.isoformat() if requirement.updated_at else None,
    }


def _serialize_change_status_request(change_request):
    requirement = db.session.get(ProjectRequirement, change_request.project_requirement_id)
    project = db.session.get(Project, requirement.project_id) if requirement else None
    requester = db.session.get(User, change_request.requested_by)
    reviewer = db.session.get(User, change_request.reviewed_by) if change_request.reviewed_by else None
    return {
        'id': change_request.id,
        'project_requirement_id': change_request.project_requirement_id,
        'requested_by': change_request.requested_by,
        'requested_by_name': requester.full_name if requester else None,
        'current_status': change_request.current_status,
        'requested_status': change_request.requested_status,
        'justification': change_request.justification,
        'status': change_request.status,
        'review_decision': change_request.review_decision,
        'review_notes': change_request.review_notes,
        'reviewed_by': change_request.reviewed_by,
        'reviewed_by_name': reviewer.full_name if reviewer else None,
        'reviewed_at': change_request.reviewed_at.isoformat() if change_request.reviewed_at else None,
        'created_at': change_request.created_at.isoformat() if change_request.created_at else None,
        'updated_at': change_request.updated_at.isoformat() if change_request.updated_at else None,
        'requirement': {
            'id': requirement.id,
            'requirement_name': requirement.requirement_name,
            'project_id': requirement.project_id,
            'project_name': project.project_name if project else None,
        } if requirement else None,
    }


def _serialize_revision_requirement(requirement):
    custom_fields = requirement.custom_fields or {}
    revision_history = list(custom_fields.get('revision_history') or [])
    latest_revision = custom_fields.get('latest_revision') or (revision_history[-1] if revision_history else None)
    if not latest_revision and requirement.workflow_status == 'returned_for_revision':
        latest_revision = {
            'notes': requirement.workflow_return_reason or '',
            'returned_at': requirement.workflow_returned_at.isoformat() if requirement.workflow_returned_at else None,
            'returned_by': requirement.workflow_returned_by,
        }
        revision_history = [latest_revision]
    latest_submission = (
        Submission.query
        .filter_by(project_requirement_id=requirement.id)
        .order_by(Submission.created_at.desc(), Submission.id.desc())
        .first()
    )
    is_active_return = requirement.workflow_status == 'returned_for_revision'
    return {
        'id': requirement.id,
        'project_id': requirement.project_id,
        'requirement_name': requirement.requirement_name,
        'annex_reference': requirement.annex_reference,
        'phase_name': requirement.phase_name,
        'status': requirement.status,
        'workflow_status': requirement.workflow_status,
        'is_active_return': is_active_return,
        'latest_revision': latest_revision,
        'current_return': latest_revision if is_active_return else None,
        'revision_history': revision_history,
        'revision_count': len(revision_history),
        'latest_submission_at': latest_submission.created_at.isoformat() if latest_submission and latest_submission.created_at else None,
        'current_version_number': latest_submission.current_version if latest_submission else 1,
        'has_report': bool(latest_revision and (latest_revision.get('report_id') or latest_revision.get('stored_object_id'))),
    }


def _serialize_phase_log(log):
    return {
        'id': log.id,
        'project_id': log.project_id,
        'from_phase': log.from_phase,
        'to_phase': log.to_phase,
        'transitioned_by': log.transitioned_by,
        'notes': log.notes,
        'transitioned_at': log.transitioned_at.isoformat() if log.transitioned_at else None,
    }


def _normalize_project_create_values(project_name, fiscal_year, budget_amount=None,
                                     target_start_date=None, target_end_date=None):
    errors = {}
    normalized = {
        'project_name': _normalize_project_name(project_name, errors),
        'fiscal_year': _normalize_fiscal_year(fiscal_year, errors),
        'budget_amount': _normalize_decimal(
            budget_amount,
            errors,
            field='budget_amount',
            label='Budget amount',
            minimum=Decimal('0'),
        ),
        'target_start_date': _normalize_date_field(target_start_date, errors, 'target_start_date', 'Target start date'),
        'target_end_date': _normalize_date_field(target_end_date, errors, 'target_end_date', 'Target end date'),
    }
    _validate_date_order(
        normalized['target_start_date'],
        normalized['target_end_date'],
        errors,
        end_field='target_end_date',
        start_label='target start date',
        end_label='Target end date',
    )
    _raise_project_validation_error(errors)
    return normalized


def _normalize_project_update_values(project, data):
    errors = {}
    normalized = {}

    if 'project_name' in data and data['project_name']:
        normalized['project_name'] = _normalize_project_name(data['project_name'], errors)
    if 'fiscal_year' in data:
        normalized['fiscal_year'] = _normalize_fiscal_year(data['fiscal_year'], errors)
    if 'budget_amount' in data:
        normalized['budget_amount'] = _normalize_decimal(
            data['budget_amount'],
            errors,
            field='budget_amount',
            label='Budget amount',
            minimum=Decimal('0'),
        )
    if 'physical_completion_pct' in data:
        normalized['physical_completion_pct'] = _normalize_decimal(
            data['physical_completion_pct'],
            errors,
            field='physical_completion_pct',
            label='Physical completion',
            minimum=Decimal('0'),
            maximum=Decimal('100'),
        )

    target_start = project.target_start_date
    target_end = project.target_end_date
    if 'target_start_date' in data:
        target_start = _normalize_date_field(data['target_start_date'], errors, 'target_start_date', 'Target start date')
        normalized['target_start_date'] = target_start
    if 'target_end_date' in data:
        target_end = _normalize_date_field(data['target_end_date'], errors, 'target_end_date', 'Target end date')
        normalized['target_end_date'] = target_end
    _validate_date_order(
        target_start,
        target_end,
        errors,
        end_field='target_end_date',
        start_label='target start date',
        end_label='Target end date',
    )

    actual_start = project.actual_start_date
    actual_end = project.actual_end_date
    if 'actual_start_date' in data:
        actual_start = _normalize_date_field(data['actual_start_date'], errors, 'actual_start_date', 'Actual start date')
        normalized['actual_start_date'] = actual_start
    if 'actual_end_date' in data:
        actual_end = _normalize_date_field(data['actual_end_date'], errors, 'actual_end_date', 'Actual end date')
        normalized['actual_end_date'] = actual_end
    _validate_date_order(
        actual_start,
        actual_end,
        errors,
        end_field='actual_end_date',
        start_label='actual start date',
        end_label='Actual end date',
    )

    _raise_project_validation_error(errors)
    return normalized


def _normalize_project_name(value, errors):
    if not isinstance(value, str) or not value.strip():
        errors['project_name'] = 'Project name is required'
        return value
    normalized = value.strip()
    if len(normalized) > 300:
        errors['project_name'] = 'Project name must be 300 characters or less'
    return normalized


def _normalize_fiscal_year(value, errors):
    field = 'fiscal_year'
    normalized = None
    if isinstance(value, bool):
        errors[field] = f'Fiscal year must be between {MIN_FISCAL_YEAR} and {MAX_FISCAL_YEAR}'
        return value
    if isinstance(value, int):
        normalized = value
    elif isinstance(value, str) and value.strip().isdigit():
        normalized = int(value.strip())
    elif isinstance(value, float) and value.is_integer():
        normalized = int(value)
    else:
        errors[field] = f'Fiscal year must be between {MIN_FISCAL_YEAR} and {MAX_FISCAL_YEAR}'
        return value

    if normalized < MIN_FISCAL_YEAR or normalized > MAX_FISCAL_YEAR:
        errors[field] = f'Fiscal year must be between {MIN_FISCAL_YEAR} and {MAX_FISCAL_YEAR}'
    return normalized


def _normalize_decimal(value, errors, field, label, minimum=None, maximum=None):
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    try:
        normalized = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        errors[field] = f'{label} must be a valid number'
        return value

    if not normalized.is_finite():
        errors[field] = f'{label} must be a valid number'
    elif minimum is not None and normalized < minimum:
        if minimum == 0:
            errors[field] = f'{label} must be non-negative'
        else:
            errors[field] = f'{label} must be at least {minimum}'
    elif maximum is not None and normalized > maximum:
        errors[field] = f'{label} must be between {minimum} and {maximum}'
    return normalized


def _normalize_date_field(value, errors, field, label):
    try:
        return _parse_date(value)
    except ProjectError as exc:
        errors[field] = f'{label} must use YYYY-MM-DD'
        return value


def _validate_date_order(start_date, end_date, errors, end_field, start_label, end_label):
    if isinstance(start_date, date) and isinstance(end_date, date) and end_date < start_date:
        errors[end_field] = f'{end_label} cannot be earlier than {start_label}'


def _raise_project_validation_error(errors):
    if errors:
        first_message = next(iter(errors.values()))
        raise ProjectError(first_message, 422, errors=errors)


def _parse_date(date_value):
    if date_value is None or date_value == '':
        return None
    if isinstance(date_value, datetime):
        return date_value.date()
    if isinstance(date_value, date):
        return date_value
    if isinstance(date_value, str):
        try:
            return date.fromisoformat(date_value)
        except ValueError:
            raise ProjectError(f'Invalid date format: {date_value}. Use YYYY-MM-DD.', 422)
    raise ProjectError(f'Invalid date format: {date_value}. Use YYYY-MM-DD.', 422)


def _project_identity_error(source_project_code=None):
    if _clean_source_value(source_project_code):
        message = 'A project with this source project code already exists. Project codes must be unique across all projects.'
        return ProjectError(message, 409, errors={'source_project_code': message})
    message = 'A project with this name already exists for this LGU and fiscal year'
    return ProjectError(message, 409, errors={'project_name': message})


def _assert_project_identity_available(lgu_id, program_id, fiscal_year, project_name,
                                       program_cycle_id=None, source_project_code=None,
                                       exclude_project_id=None):
    query = Project.query
    if exclude_project_id is not None:
        query = query.filter(Project.id != exclude_project_id)

    cleaned_code = _clean_source_value(source_project_code)
    if cleaned_code:
        if len(cleaned_code) > 120:
            raise ProjectError('Project code must be 120 characters or less', 422,
                               errors={'source_project_code': 'Project code must be 120 characters or less'})
        existing = query.filter(
            func.lower(func.trim(Project.source_project_code)) == cleaned_code.lower(),
        ).first()
        if existing:
            raise _project_identity_error(cleaned_code)
        return

    normalized_name = project_name.strip().lower() if project_name else ''
    if normalized_name:
        existing_name = query.filter(
            Project.lgu_id == lgu_id,
            Project.fiscal_year == fiscal_year,
            func.lower(func.trim(Project.project_name)) == normalized_name,
        ).first()
        if existing_name:
            raise ProjectError('A project with this name already exists for this LGU and fiscal year', 409)


