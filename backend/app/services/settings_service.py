import base64
import hashlib
from pathlib import Path
from datetime import datetime, timezone
import json
import mimetypes
import os
import re
import shutil
import uuid
import zipfile

from cryptography.fernet import Fernet, InvalidToken
from flask import current_app
from werkzeug.utils import secure_filename

from app.extensions import db
from app.models import SystemConfig
from app.services import audit_service, document_template_service, object_storage_service
from app.services.upload_validation import ArchiveSecurityError, resolve_upload_mime_type, validate_ooxml_container


VALID_VALUE_TYPES = ('string', 'integer', 'boolean', 'json', 'secret')
ENCRYPTED_VALUE_PREFIX = 'fernet:v1:'
BRANDING_ASSET_TYPES = {'logo', 'favicon'}
DOCUMENT_TEMPLATE_TYPES = {
    'project_revision_report': {
        'label': 'Returned for Revision Report',
        'description': 'DOCX template used when generating project correction reports from returned documentary requirements.',
        'config_prefix': 'report_template_project_revision_report',
        'category': 'reports',
        'allowed_extension': 'docx',
        'mime_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    },
    'annex_e': {
        'label': 'Annex E Notice of Deficiencies',
        'description': 'DOCX template used for notice-of-deficiency drafts and downloadable notice documents.',
        'config_prefix': 'document_template_annex_e',
        'category': 'notices',
        'allowed_extension': 'docx',
        'mime_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    },
    'annex_m': {
        'label': 'Annex M Notice of Non-Compliance No. 1',
        'description': 'DOCX template used for the first notice of non-compliance.',
        'config_prefix': 'document_template_annex_m',
        'category': 'notices',
        'allowed_extension': 'docx',
        'mime_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    },
    'annex_n': {
        'label': 'Annex N Notice of Non-Compliance No. 2',
        'description': 'DOCX template used for the second notice of non-compliance.',
        'config_prefix': 'document_template_annex_n',
        'category': 'notices',
        'allowed_extension': 'docx',
        'mime_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    },
    'annex_o': {
        'label': 'Annex O Notice of Non-Compliance No. 3',
        'description': 'DOCX template used for the third notice of non-compliance and final escalation package.',
        'config_prefix': 'document_template_annex_o',
        'category': 'notices',
        'allowed_extension': 'docx',
        'mime_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    },
    'general_document': {
        'label': 'General Letter or Document',
        'description': 'Reusable DOCX layout for office letters, memoranda, notices, or other manually prepared document outputs.',
        'config_prefix': 'document_template_general_document',
        'category': 'documents',
        'allowed_extension': 'docx',
        'mime_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    },
}
REPORT_TEMPLATE_TYPES = DOCUMENT_TEMPLATE_TYPES
NOTICE_DOCUMENT_TEMPLATE_TYPES = ('annex_e', 'annex_m', 'annex_n', 'annex_o')
REVISION_REPORT_TEMPLATE_TYPES = ('project_revision_report',)
BRANDING_TEXT_FIELDS = {
    'app_name': {
        'key': 'app_name',
        'label': 'Application Name',
        'max_length': 80,
        'required': True,
    },
    'assistant_name': {
        'key': 'assistant_name',
        'label': 'Assistant Name',
        'max_length': 80,
        'required': False,
    },
    'app_tagline': {
        'key': 'app_tagline',
        'label': 'Application Tagline',
        'max_length': 180,
        'required': False,
    },
    'logo_alt_text': {
        'key': 'branding_logo_alt_text',
        'label': 'Logo Alternative Text',
        'max_length': 160,
        'required': False,
    },
}
OPERATIONAL_DEFAULTS = (
    {
        'key': 'default_intake_source',
        'label': 'Default Intake Source',
        'value_type': 'string',
        'default_value': 'dms_ingestion',
        'description': 'Fallback source when staff records a documentary intake without an explicit channel.',
        'allowed_values': ['dms_ingestion', 'direct_upload', 'field_upload', 'api_push'],
    },
    {
        'key': 'default_notice_channel',
        'label': 'Default Notice Channel',
        'value_type': 'string',
        'default_value': 'dms',
        'description': 'Preferred dispatch channel for reminder and notice preparation.',
        'allowed_values': ['email', 'dms', 'intranet', 'manual_print'],
    },
    {
        'key': 'default_monitoring_status',
        'label': 'Default Monitoring Status',
        'value_type': 'string',
        'default_value': 'draft',
        'description': 'Initial state for newly recorded QAAR, Aide Memoire, or site-validation records.',
        'allowed_values': ['draft', 'submitted', 'reviewed', 'approved', 'flagged'],
    },
    {
        'key': 'default_monitoring_transmittal_status',
        'label': 'Default Monitoring Transmittal',
        'value_type': 'string',
        'default_value': 'not_required',
        'description': 'Initial transmittal state for implementation-monitoring findings.',
        'allowed_values': ['not_required', 'pending', 'transmitted', 'acknowledged'],
    },
    {
        'key': 'regional_handoff_destination_label',
        'label': 'Regional Handoff Destination Label',
        'value_type': 'string',
        'default_value': 'Regional Google Drive',
        'description': 'Operator-facing name shown for the configured regional reference destination.',
        'allowed_values': [],
    },
    {
        'key': 'guideline_library_retention_years',
        'label': 'Guideline Retention Years',
        'value_type': 'integer',
        'default_value': '5',
        'description': 'Planning default for how long historical issuances remain prominent in lookup screens.',
        'allowed_values': [],
    },
)


class SettingsError(Exception):
    def __init__(self, message, status_code=400):
        self.message = message
        self.status_code = status_code
        super().__init__(message)


UNREADABLE_GOOGLE_DRIVE_TOKEN_MESSAGE = (
    'Stored Google Drive account credentials cannot be decrypted; reconnect Google Drive from Settings.'
)
GOOGLE_DRIVE_SERVICE_ACCOUNT_REQUIRED_FIELDS = {'client_email', 'private_key', 'token_uri'}


RISK_THRESHOLD_MODERATE_KEY = 'monitoring.risk_threshold_moderate_pct'
RISK_THRESHOLD_HIGH_KEY = 'monitoring.risk_threshold_high_pct'
DEFAULT_RISK_THRESHOLD_MODERATE_PCT = 5
DEFAULT_RISK_THRESHOLD_HIGH_PCT = 15


def get_risk_thresholds():
    """Return the configurable Monthly Progress Report slippage risk thresholds.

    Values represent the minimum negative slippage percentage
    (target_to_date_pct - actual_to_date_pct) required to classify a report
    as Moderate Risk (minor) or High Risk (severe). Reuses the existing
    SystemConfig key/value pattern instead of introducing a new settings
    model.
    """
    moderate_raw = _config_by_key(RISK_THRESHOLD_MODERATE_KEY)
    high_raw = _config_by_key(RISK_THRESHOLD_HIGH_KEY)
    try:
        moderate_pct = int(moderate_raw) if moderate_raw not in (None, '') else DEFAULT_RISK_THRESHOLD_MODERATE_PCT
    except (TypeError, ValueError):
        moderate_pct = DEFAULT_RISK_THRESHOLD_MODERATE_PCT
    try:
        high_pct = int(high_raw) if high_raw not in (None, '') else DEFAULT_RISK_THRESHOLD_HIGH_PCT
    except (TypeError, ValueError):
        high_pct = DEFAULT_RISK_THRESHOLD_HIGH_PCT
    return {
        'moderate_pct': moderate_pct,
        'high_pct': high_pct,
    }


def update_risk_thresholds(data, user_id):
    before = get_risk_thresholds()
    moderate_pct = data.get('moderate_pct', before['moderate_pct'])
    high_pct = data.get('high_pct', before['high_pct'])
    try:
        moderate_pct = int(moderate_pct)
        high_pct = int(high_pct)
    except (TypeError, ValueError) as exc:
        raise SettingsError('Risk thresholds must be integers.', 422) from exc
    if moderate_pct < 0 or high_pct < 0:
        raise SettingsError('Risk thresholds must be non-negative.', 422)
    if high_pct <= moderate_pct:
        raise SettingsError('The High Risk threshold must be greater than the Moderate Risk threshold.', 422)

    _upsert_config(
        key=RISK_THRESHOLD_MODERATE_KEY,
        value=str(moderate_pct),
        value_type='integer',
        description=(
            'Minimum negative slippage percentage (Target to Date - Actual to Date) '
            'at which a Monthly Progress Report is classified as Moderate Risk (minor).'
        ),
        user_id=user_id,
    )
    _upsert_config(
        key=RISK_THRESHOLD_HIGH_KEY,
        value=str(high_pct),
        value_type='integer',
        description=(
            'Minimum negative slippage percentage (Target to Date - Actual to Date) '
            'at which a Monthly Progress Report is classified as High Risk (severe) '
            'and may auto-trigger a Management Action Report.'
        ),
        user_id=user_id,
    )
    db.session.flush()

    audit_service.log_action(
        user_id=int(user_id),
        action='update',
        entity_type='system_config',
        entity_id=0,
        old_value=before,
        new_value=get_risk_thresholds(),
    )
    db.session.commit()
    return get_risk_thresholds()


def get_readiness_snapshot(object_storage_check=None):
    storage_check = object_storage_check if object_storage_check is not None else _object_storage_check()
    checks = [
        _database_check(),
        _jwt_check(),
        _smtp_check(),
        _ai_provider_check(),
        _document_intelligence_check(),
        _google_drive_check(),
        storage_check,
        _uploads_check(),
        _guideline_uploads_check(),
        _celery_check(),
    ]
    summary = {
        'total': len(checks),
        'ready': sum(1 for item in checks if item['status'] == 'ready'),
        'warning': sum(1 for item in checks if item['status'] == 'warning'),
        'disabled': sum(1 for item in checks if item['status'] == 'disabled'),
    }
    return {
        'summary': summary,
        'checks': checks,
    }


def get_google_drive_handoff_readiness():
    return _google_drive_check()


def get_branding():
    try:
        configs = _branding_config_map()
    except Exception:
        configs = {}
    app_name = _config_value(configs, 'app_name', _default_app_name())
    app_tagline = _config_value(configs, 'app_tagline', _default_app_tagline())
    logo_alt_text = _config_value(configs, 'branding_logo_alt_text', f'{app_name} logo')
    assistant_name = _config_value(configs, 'assistant_name', _default_assistant_name(app_name))
    return {
        'app_name': app_name,
        'app_tagline': app_tagline,
        'logo_alt_text': logo_alt_text or f'{app_name} logo',
        'assistant_name': assistant_name,
        'logo_url': _asset_url('logo', configs),
        'favicon_url': _asset_url('favicon', configs),
    }


def get_app_name():
    return get_branding()['app_name']


def get_app_tagline():
    return get_branding()['app_tagline']


def get_assistant_name():
    return get_branding()['assistant_name']


def format_dynamic_branding(text, app_name=None, assistant_name=None):
    if not text or not isinstance(text, str):
        return text
    resolved_app = app_name or get_app_name()
    resolved_assistant = assistant_name or get_assistant_name()
    return text.replace('{app_name}', resolved_app).replace('{assistant_name}', resolved_assistant)



def update_branding(data, user_id):
    updates = {}
    for field, definition in BRANDING_TEXT_FIELDS.items():
        if field not in data:
            continue
        value = str(data.get(field) or '').strip()
        if definition['required'] and not value:
            raise SettingsError(f"{definition['label']} is required.", 422)
        if len(value) > definition['max_length']:
            raise SettingsError(f"{definition['label']} must be {definition['max_length']} characters or less.", 422)
        updates[definition['key']] = value

    if not updates:
        return get_branding()

    before = get_branding()
    for key, value in updates.items():
        _upsert_config(
            key=key,
            value=value,
            value_type='string',
            description=_branding_description(key),
            user_id=user_id,
        )
    db.session.flush()

    audit_service.log_action(
        user_id=int(user_id),
        action='update',
        entity_type='system_config',
        entity_id=0,
        old_value=before,
        new_value=get_branding(),
    )
    db.session.commit()
    return get_branding()


def store_brand_asset(asset_type, upload, user_id):
    if asset_type not in BRANDING_ASSET_TYPES:
        raise SettingsError('Branding asset type must be logo or favicon.', 404)
    if upload is None or not getattr(upload, 'filename', '').strip():
        raise SettingsError('An image file is required.', 422)

    original_filename = secure_filename(upload.filename)
    if not original_filename:
        raise SettingsError('Uploaded image file name is invalid.', 422)

    extension = Path(original_filename).suffix.lower().lstrip('.')
    if extension == 'svg':
        raise SettingsError('SVG branding images are not supported for security reasons.', 422)

    allowed_extensions = set(current_app.config.get('BRANDING_ALLOWED_EXTENSIONS') or ())
    if not extension or extension not in allowed_extensions:
        allowed = ', '.join(sorted(allowed_extensions))
        raise SettingsError(f'Branding images must use one of these file types: {allowed}.', 422)

    mime_type = _validate_brand_asset_mime_type(original_filename, upload)
    _validate_brand_asset_size(upload)

    upload_dir = _branding_upload_root()
    stored_filename = f'{asset_type}_{uuid.uuid4().hex}.{extension}'
    destination_path = upload_dir / stored_filename
    upload.save(destination_path)
    try:
        _validate_brand_image_signature(destination_path, extension)
    except SettingsError:
        destination_path.unlink(missing_ok=True)
        raise

    before = get_branding()
    _upsert_config(
        key=f'branding_{asset_type}_file_path',
        value=str(destination_path),
        value_type='string',
        description=f'Stored {asset_type} image path for application branding.',
        user_id=user_id,
    )
    _upsert_config(
        key=f'branding_{asset_type}_mime_type',
        value=mime_type,
        value_type='string',
        description=f'MIME type for the stored application {asset_type}.',
        user_id=user_id,
    )
    _upsert_config(
        key=f'branding_{asset_type}_file_name',
        value=original_filename,
        value_type='string',
        description=f'Original file name for the stored application {asset_type}.',
        user_id=user_id,
    )
    db.session.flush()

    audit_service.log_action(
        user_id=int(user_id),
        action='upload',
        entity_type='system_config',
        entity_id=0,
        old_value=before,
        new_value=get_branding(),
    )
    db.session.commit()
    return get_branding()


def get_brand_asset(asset_type):
    if asset_type not in BRANDING_ASSET_TYPES:
        raise SettingsError('Branding asset type must be logo or favicon.', 404)

    config = SystemConfig.query.filter_by(key=f'branding_{asset_type}_file_path').first()
    if not config or not config.value:
        raise SettingsError('Branding image has not been uploaded.', 404)

    path = _resolve_under_root(config.value, _branding_upload_root())
    if not path or not path.exists() or not path.is_file():
        raise SettingsError('Branding image is not available on disk.', 404)

    mime_type = _config_by_key(f'branding_{asset_type}_mime_type')
    file_name = _config_by_key(f'branding_{asset_type}_file_name')
    return path, mime_type or (mimetypes.guess_type(path.name)[0] or 'application/octet-stream'), file_name or path.name


def list_document_templates():
    configs = {item.key: item for item in SystemConfig.query.all()}
    catalogs = {
        template_type: _load_document_template_catalog(template_type, configs=configs)
        for template_type in DOCUMENT_TEMPLATE_TYPES
    }
    default_sources = _document_template_default_sources(catalogs)
    templates = []
    for template_type, definition in DOCUMENT_TEMPLATE_TYPES.items():
        prefix = definition['config_prefix']
        path_config = configs.get(f'{prefix}_file_path')
        file_name_config = configs.get(f'{prefix}_file_name')
        updated_config = configs.get(f'{prefix}_file_path')
        path = Path(path_config.value) if path_config and path_config.value else None
        catalog = catalogs.get(template_type) or []
        serialized_templates = [
            _serialize_document_template_entry(
                template_type,
                entry,
                default_sources=default_sources,
            )
            for entry in catalog
        ]
        placeholder_definitions = _document_template_placeholders(template_type, serialized_templates)
        templates.append({
            'template_type': template_type,
            'label': definition['label'],
            'description': definition['description'],
            'category': definition.get('category') or 'documents',
            'configured': any(item['available'] for item in serialized_templates) or bool(_is_safe_existing_file(path, _report_template_upload_root()) if path else False),
            'file_name': file_name_config.value if file_name_config else None,
            'updated_at': updated_config.updated_at.isoformat() if updated_config and updated_config.updated_at else None,
            'placeholder_fields': [item['key'] for item in placeholder_definitions],
            'placeholder_definitions': placeholder_definitions,
            'default_template_id': _default_document_template_id(serialized_templates),
            'templates': serialized_templates,
        })
    return templates


def list_report_templates():
    return [
        item for item in list_document_templates()
        if item.get('template_type') == 'project_revision_report'
    ]


def list_notice_document_templates():
    return _workflow_template_groups(NOTICE_DOCUMENT_TEMPLATE_TYPES)


def list_revision_report_templates():
    return _workflow_template_groups(REVISION_REPORT_TEMPLATE_TYPES)


def store_document_template(template_type, upload, user_id, label=None, placeholder_defaults=None, is_default=True):
    definition = DOCUMENT_TEMPLATE_TYPES.get(template_type)
    if not definition:
        raise SettingsError('Document template type is not supported.', 404)
    parsed_is_default = _parse_template_default_flag(is_default)
    if upload is None or not getattr(upload, 'filename', '').strip():
        raise SettingsError('A DOCX template file is required.', 422)

    original_filename = secure_filename(upload.filename)
    if not original_filename:
        raise SettingsError('Uploaded template file name is invalid.', 422)

    extension = Path(original_filename).suffix.lower().lstrip('.')
    if extension != definition['allowed_extension']:
        raise SettingsError('Document templates must be uploaded as DOCX files.', 422)

    _validate_report_template_size(upload)
    upload_dir = _report_template_upload_root()
    stored_filename = f'{template_type}_{uuid.uuid4().hex}.{extension}'
    destination_path = upload_dir / stored_filename
    upload.save(destination_path)

    try:
        _validate_docx_file(destination_path)
        placeholder_definitions = document_template_service.extract_docx_placeholders(destination_path)
    except SettingsError:
        destination_path.unlink(missing_ok=True)
        raise
    except document_template_service.DocumentTemplateError as exc:
        destination_path.unlink(missing_ok=True)
        raise SettingsError(exc.message, exc.status_code) from exc

    before = list_document_templates()
    default_sources = _document_template_default_sources()
    catalog = _load_document_template_catalog(template_type)
    parsed_placeholder_defaults = _parse_placeholder_defaults(placeholder_defaults)
    merged_placeholder_defaults = _merge_unambiguous_placeholder_defaults(
        placeholder_definitions,
        parsed_placeholder_defaults,
        default_sources,
    )
    entry = {
        'id': uuid.uuid4().hex,
        'label': _clean_template_label(label) or original_filename,
        'file_name': original_filename,
        'file_path': str(destination_path),
        'mime_type': definition['mime_type'],
        'uploaded_at': datetime.now(timezone.utc).isoformat(),
        'uploaded_by': int(user_id),
        'placeholder_defaults': merged_placeholder_defaults,
        'placeholder_definitions': placeholder_definitions,
        'placeholder_fields': [item['key'] for item in placeholder_definitions],
        'is_default': parsed_is_default,
    }
    if entry['is_default'] or not catalog:
        for item in catalog:
            item['is_default'] = False
        entry['is_default'] = True
    catalog.append(entry)
    _sync_current_document_template_config(template_type, catalog, user_id)
    _upsert_config(
        key=f'{definition["config_prefix"]}_templates',
        value=json.dumps(catalog, sort_keys=True),
        value_type='json',
        description=f'Available DOCX document template options for {definition["label"]}.',
        user_id=user_id,
    )
    db.session.flush()

    audit_service.log_action(
        user_id=int(user_id),
        action='upload',
        entity_type='system_config',
        entity_id=0,
        old_value=before,
        new_value=list_document_templates(),
    )
    db.session.commit()
    return list_document_templates()


def store_report_template(template_type, upload, user_id, label=None, placeholder_defaults=None, is_default=True):
    store_document_template(template_type, upload, user_id, label, placeholder_defaults, is_default)
    return list_report_templates()


def preview_document_template_placeholders(template_type, upload):
    """Scan an uploaded DOCX for placeholders without persisting it."""
    definition = DOCUMENT_TEMPLATE_TYPES.get(template_type)
    if not definition:
        raise SettingsError('Document template type is not supported.', 404)
    if upload is None or not getattr(upload, 'filename', '').strip():
        raise SettingsError('A DOCX template file is required.', 422)

    original_filename = secure_filename(upload.filename)
    if not original_filename:
        raise SettingsError('Uploaded template file name is invalid.', 422)

    extension = Path(original_filename).suffix.lower().lstrip('.')
    if extension != definition['allowed_extension']:
        raise SettingsError('Document templates must be uploaded as DOCX files.', 422)

    _validate_report_template_size(upload)
    upload_dir = _report_template_upload_root()
    preview_path = upload_dir / f'preview_{uuid.uuid4().hex}.{extension}'

    try:
        upload.save(preview_path)
        _validate_docx_file(preview_path)
        placeholder_definitions = document_template_service.extract_docx_placeholders(preview_path)
    except SettingsError:
        raise
    except document_template_service.DocumentTemplateError as exc:
        raise SettingsError(exc.message, exc.status_code) from exc
    finally:
        preview_path.unlink(missing_ok=True)

    default_sources = _document_template_default_sources()
    placeholder_defaults = _merge_unambiguous_placeholder_defaults(
        placeholder_definitions,
        {},
        default_sources,
    )
    return {
        'file_name': original_filename,
        'mime_type': definition['mime_type'],
        'placeholder_defaults': placeholder_defaults,
        'placeholder_default_suggestions': _placeholder_default_suggestions(
            placeholder_definitions,
            default_sources,
        ),
        'placeholder_fields': [item['key'] for item in placeholder_definitions],
        'placeholder_definitions': placeholder_definitions,
    }


def update_document_template(template_type, template_id, data, user_id):
    definition = DOCUMENT_TEMPLATE_TYPES.get(template_type)
    if not definition:
        raise SettingsError('Document template type is not supported.', 404)
    catalog = _load_document_template_catalog(template_type)
    entry = next((item for item in catalog if item.get('id') == template_id), None)
    if entry is None:
        raise SettingsError('Document template not found.', 404)

    before = list_document_templates()
    if 'label' in data:
        entry['label'] = _clean_template_label(data.get('label')) or entry.get('file_name') or definition['label']
    if 'placeholder_defaults' in data:
        entry['placeholder_defaults'] = _parse_placeholder_defaults(data.get('placeholder_defaults'))
    if 'is_default' in data:
        wants_default = _parse_template_default_flag(data.get('is_default'))
        if wants_default:
            for item in catalog:
                item['is_default'] = item.get('id') == template_id
            _sync_current_document_template_config(template_type, catalog, user_id)
        elif entry.get('is_default') and len(catalog) > 1:
            entry['is_default'] = False
            replacement_default = _first_available_template_entry(
                [item for item in catalog if item.get('id') != template_id],
            ) or next((item for item in catalog if item.get('id') != template_id), None)
            for item in catalog:
                item['is_default'] = item is replacement_default
            _sync_current_document_template_config(template_type, catalog, user_id)

    _upsert_config(
        key=f'{definition["config_prefix"]}_templates',
        value=json.dumps(catalog, sort_keys=True),
        value_type='json',
        description=f'Available DOCX document template options for {definition["label"]}.',
        user_id=user_id,
    )
    db.session.flush()
    audit_service.log_action(
        user_id=int(user_id),
        action='update',
        entity_type='system_config',
        entity_id=0,
        old_value=before,
        new_value=list_document_templates(),
    )
    db.session.commit()
    return list_document_templates()


def update_report_template(template_type, template_id, data, user_id):
    update_document_template(template_type, template_id, data, user_id)
    return list_report_templates()


def delete_document_template(template_type, template_id, user_id):
    definition = DOCUMENT_TEMPLATE_TYPES.get(template_type)
    if not definition:
        raise SettingsError('Document template type is not supported.', 404)
    catalog = _load_document_template_catalog(template_type)
    entry = next((item for item in catalog if item.get('id') == template_id), None)
    if entry is None:
        raise SettingsError('Document template not found.', 404)

    before = list_document_templates()
    remaining = [item for item in catalog if item.get('id') != template_id]
    if remaining and not any(item.get('is_default') for item in remaining):
        replacement_default = _first_available_template_entry(remaining) or remaining[0]
        for item in remaining:
            item['is_default'] = item is replacement_default

    _upsert_config(
        key=f'{definition["config_prefix"]}_templates',
        value=json.dumps(remaining, sort_keys=True),
        value_type='json',
        description=f'Available DOCX document template options for {definition["label"]}.',
        user_id=user_id,
    )
    _sync_current_document_template_config(template_type, remaining, user_id)

    removed_path = _resolve_under_root(entry.get('file_path'), _report_template_upload_root())
    if removed_path and not _path_is_referenced(removed_path, remaining):
        try:
            removed_path.unlink(missing_ok=True)
        except OSError:
            current_app.logger.warning('Unable to delete document template file: %s', removed_path)

    db.session.flush()
    audit_service.log_action(
        user_id=int(user_id),
        action='delete',
        entity_type='system_config',
        entity_id=0,
        old_value=before,
        new_value=list_document_templates(),
    )
    db.session.commit()
    return list_document_templates()


def delete_report_template(template_type, template_id, user_id):
    delete_document_template(template_type, template_id, user_id)
    return list_report_templates()


def get_document_template(template_type, template_id=None):
    definition = DOCUMENT_TEMPLATE_TYPES.get(template_type)
    if not definition:
        return None
    catalog = _load_document_template_catalog(template_type)
    selected = None
    if template_id:
        selected = next((item for item in catalog if item.get('id') == template_id), None)
        if selected is None:
            raise SettingsError('Document template not found.', 404)
    elif catalog:
        selected = next((item for item in catalog if item.get('is_default')), None) or catalog[-1]

    if selected:
        candidate = _resolve_under_root(selected.get('file_path'), _report_template_upload_root())
        if not candidate or not candidate.exists() or not candidate.is_file():
            if template_id:
                raise SettingsError('Document template file is not available on disk.', 404)
            return None
        return {
            'id': selected.get('id'),
            'label': selected.get('label') or selected.get('file_name') or definition['label'],
            'file_name': selected.get('file_name') or candidate.name,
            'path': candidate,
            'mime_type': selected.get('mime_type') or definition['mime_type'],
            'placeholder_defaults': selected.get('placeholder_defaults') or {},
            'placeholder_definitions': _entry_placeholder_definitions(selected, candidate),
            'placeholder_fields': selected.get('placeholder_fields') or [],
        }

    path = _config_by_key(f'{definition["config_prefix"]}_file_path')
    if not path:
        return None
    candidate = _resolve_under_root(path, _report_template_upload_root())
    if not candidate:
        return None
    if not candidate.exists() or not candidate.is_file():
        return None
    return {
        'id': None,
        'label': definition['label'],
        'file_name': _config_by_key(f'{definition["config_prefix"]}_file_name') or candidate.name,
        'path': candidate,
        'mime_type': definition['mime_type'],
        'placeholder_defaults': {},
        'placeholder_definitions': _entry_placeholder_definitions({}, candidate),
        'placeholder_fields': [],
    }


def get_report_template(template_type, template_id=None):
    return get_document_template(template_type, template_id)


def get_document_template_file(template_type):
    template = get_document_template(template_type)
    return template['path'] if template else None


def get_report_template_file(template_type):
    return get_document_template_file(template_type)


def list_configs():
    return [
        _serialize_config(item)
        for item in SystemConfig.query.order_by(SystemConfig.key.asc()).all()
        if not _is_protected_config_key(item.key)
    ]


def get_defaults_catalog():
    configs = {item.key: item for item in SystemConfig.query.all()}
    catalog = []
    for default in OPERATIONAL_DEFAULTS:
        config = configs.get(default['key'])
        serialized = _serialize_config(config) if config else None
        catalog.append({
            **default,
            'configured': config is not None and bool(config.value),
            'current_value': serialized['value'] if serialized else None,
            'has_value': serialized['has_value'] if serialized else False,
            'config_id': config.id if config else None,
        })
    return catalog


def create_config(data, user_id):
    _ensure_public_config_key(data['key'])
    if SystemConfig.query.filter_by(key=data['key']).first():
        raise SettingsError('Configuration key already exists', 409)
    value_type = data.get('value_type', 'string')
    _validate_value_type(value_type)
    _validate_config_value(data['key'], data.get('value'), value_type)
    should_encrypt = value_type == 'secret' or bool(data.get('is_encrypted', False))

    config = SystemConfig(
        key=data['key'],
        value=_store_config_value(data.get('value'), should_encrypt),
        value_type=value_type,
        description=data.get('description'),
        is_encrypted=bool(should_encrypt and data.get('value') not in (None, '')),
        updated_by=int(user_id),
    )
    db.session.add(config)
    db.session.commit()

    audit_service.log_action(
        user_id=int(user_id),
        action='create',
        entity_type='system_config',
        entity_id=config.id,
        old_value=None,
        new_value=_serialize_config(config),
    )
    return config


def update_config(config_id, data, user_id):
    config = db.session.get(SystemConfig, config_id)
    if not config:
        raise SettingsError('System config not found', 404)
    _ensure_public_config_key(config.key)

    before = _serialize_config(config)

    next_key = data.get('key', config.key)
    if next_key != config.key:
        _ensure_public_config_key(next_key)
        existing = SystemConfig.query.filter(SystemConfig.key == next_key, SystemConfig.id != config.id).first()
        if existing:
            raise SettingsError('Configuration key already exists', 409)
    next_value_type = data.get('value_type', config.value_type)
    _validate_value_type(next_value_type)
    next_value = data['value'] if 'value' in data else _read_config_value(config)
    _validate_config_value(next_key, next_value, next_value_type)
    next_is_encrypted = (
        next_value_type == 'secret'
        or bool(data.get('is_encrypted', config.is_encrypted))
    )

    config.key = next_key
    config.value = _store_config_value(next_value, next_is_encrypted)
    config.value_type = next_value_type
    if 'description' in data:
        config.description = data['description']
    config.is_encrypted = bool(next_is_encrypted and next_value not in (None, ''))
    config.updated_by = int(user_id)

    db.session.commit()

    audit_service.log_action(
        user_id=int(user_id),
        action='update',
        entity_type='system_config',
        entity_id=config.id,
        old_value=before,
        new_value=_serialize_config(config),
    )
    return config


def _validate_value_type(value_type):
    if value_type not in VALID_VALUE_TYPES:
        raise SettingsError(
            f'value_type must be one of: {", ".join(VALID_VALUE_TYPES)}',
            422,
        )


def _validate_config_value(key, value, value_type):
    if value in (None, '') or value_type == 'secret':
        return

    if value_type == 'integer':
        try:
            int(value)
        except (TypeError, ValueError) as exc:
            raise SettingsError('Configuration value must be an integer.', 422) from exc
    elif value_type == 'boolean':
        if str(value).lower() not in {'true', 'false', '1', '0', 'yes', 'no', 'on', 'off'}:
            raise SettingsError('Configuration value must be boolean-like.', 422)
    elif value_type == 'json':
        try:
            json.loads(value)
        except (TypeError, ValueError) as exc:
            raise SettingsError('Configuration value must be valid JSON.', 422) from exc

    default = _default_for_key(key)
    if default and default.get('allowed_values') and value not in default['allowed_values']:
        allowed = ', '.join(default['allowed_values'])
        raise SettingsError(f'{key} must be one of: {allowed}', 422)


def _default_for_key(key):
    return next((item for item in OPERATIONAL_DEFAULTS if item['key'] == key), None)


def _serialize_config(config):
    is_secret = config.value_type == 'secret' or bool(config.is_encrypted)
    return {
        'id': config.id,
        'key': config.key,
        'value': None if is_secret else config.value,
        'has_value': bool(config.value),
        'value_type': config.value_type,
        'description': config.description,
        'is_encrypted': bool(config.is_encrypted),
        'updated_by': config.updated_by,
        'created_at': config.created_at.isoformat() if config.created_at else None,
        'updated_at': config.updated_at.isoformat() if config.updated_at else None,
    }


def _read_config_value(config):
    if not config:
        return None
    if config.is_encrypted:
        return _decrypt_secret_value(config.value)
    return config.value


def _store_config_value(value, encrypt=False):
    if value in (None, ''):
        return value
    return _encrypt_secret_value(value) if encrypt else value


def _is_encrypted_value(value):
    return isinstance(value, str) and value.startswith(ENCRYPTED_VALUE_PREFIX)


def _encrypt_secret_value(value):
    if value in (None, '') or _is_encrypted_value(value):
        return value
    token = _secret_cipher().encrypt(str(value).encode('utf-8')).decode('utf-8')
    return f'{ENCRYPTED_VALUE_PREFIX}{token}'


def _decrypt_secret_value(value):
    if value in (None, '') or not _is_encrypted_value(value):
        return value
    token = value[len(ENCRYPTED_VALUE_PREFIX):]
    try:
        return _secret_cipher().decrypt(token.encode('utf-8')).decode('utf-8')
    except InvalidToken as exc:
        raise SettingsError(
            'Encrypted configuration could not be decrypted. Check SYSTEM_CONFIG_SECRET_KEY.',
            503,
        ) from exc


def _secret_cipher():
    secret = current_app.config.get('SYSTEM_CONFIG_SECRET_KEY')
    if not secret:
        raise SettingsError('SYSTEM_CONFIG_SECRET_KEY is required for encrypted settings.', 503)
    key = base64.urlsafe_b64encode(hashlib.sha256(str(secret).encode('utf-8')).digest())
    return Fernet(key)


def _is_decryption_error(exc):
    return (
        isinstance(exc, SettingsError)
        and exc.status_code == 503
        and 'could not be decrypted' in str(exc.message)
    )


def _branding_config_keys():
    keys = {definition['key'] for definition in BRANDING_TEXT_FIELDS.values()}
    for asset_type in BRANDING_ASSET_TYPES:
        keys.update({
            f'branding_{asset_type}_file_path',
            f'branding_{asset_type}_mime_type',
            f'branding_{asset_type}_file_name',
        })
    return keys


def _branding_config_map():
    keys = _branding_config_keys()
    return {
        item.key: item
        for item in SystemConfig.query.filter(SystemConfig.key.in_(keys)).all()
    }


def _config_value(configs, key, default=None):
    config = configs.get(key)
    if config is None or config.value in (None, ''):
        return default
    return _read_config_value(config)


def _config_by_key(key):
    config = SystemConfig.query.filter_by(key=key).first()
    return _read_config_value(config) if config and config.value not in (None, '') else None


def _ensure_public_config_key(key):
    if _is_protected_config_key(key):
        raise SettingsError('This configuration key is managed by a dedicated upload or branding workflow.', 422)


def _is_protected_config_key(key):
    return key in _protected_config_keys()


def _protected_config_keys():
    keys = set()
    keys.update({
        'google_drive_oauth_token',
        'google_drive_oauth_state',
        'google_drive_root_folder_id',
    })
    for asset_type in BRANDING_ASSET_TYPES:
        keys.update({
            f'branding_{asset_type}_file_path',
            f'branding_{asset_type}_mime_type',
            f'branding_{asset_type}_file_name',
        })
    keys.update({
        'GEMINI_API_KEY',
        'GROQ_API_KEY',
        'OPENROUTER_API_KEY',
        'GEMINI_MODEL',
        'GROQ_MODEL',
        'OPENROUTER_MODEL',
        'AI_PROVIDER_ORDER',
        'AI_REQUEST_TIMEOUT_SECONDS',
        'AI_TEMPERATURE',
        'AI_MAX_OUTPUT_TOKENS',
        'AI_MAX_PROMPT_CHARS',
        'OBJECT_STORAGE_PROVIDER',
        'R2_BUCKET_NAME',
        'R2_ENDPOINT_URL',
        'R2_ACCESS_KEY_ID',
        'R2_SECRET_ACCESS_KEY',
        'R2_REGION',
        'OBJECT_STORAGE_BUCKET_PREFIX',
        'SMTP_HOST',
        'SMTP_PORT',
        'SMTP_USERNAME',
        'SMTP_PASSWORD',
        'SMTP_USE_TLS',
        'MAIL_FROM',
        'GOOGLE_DRIVE_SHARED_DRIVE_ID',
    })
    for definition in DOCUMENT_TEMPLATE_TYPES.values():
        prefix = definition['config_prefix']
        keys.update({
            f'{prefix}_file_path',
            f'{prefix}_file_name',
            f'{prefix}_mime_type',
            f'{prefix}_templates',
        })
    return keys


def _resolve_under_root(path_value, root):
    if not path_value:
        return None
    try:
        root_path = Path(root).resolve()
        candidate = Path(path_value).resolve()
        candidate.relative_to(root_path)
        return candidate
    except (OSError, RuntimeError, ValueError):
        return None


def _is_safe_existing_file(path, root):
    if path is None:
        return False
    candidate = _resolve_under_root(str(path), root)
    return bool(candidate and candidate.exists() and candidate.is_file())


def _default_app_name():
    try:
        if current_app and current_app.config.get('APP_NAME'):
            return str(current_app.config['APP_NAME']).strip()
    except Exception:
        pass
    env_name = os.environ.get('APP_NAME')
    if env_name and env_name.strip():
        return env_name.strip()
    return 'DILG Operations Workspace'


def _default_app_tagline():
    try:
        if current_app and current_app.config.get('APP_TAGLINE'):
            return str(current_app.config['APP_TAGLINE']).strip()
    except Exception:
        pass
    env_tagline = os.environ.get('APP_TAGLINE')
    if env_tagline and env_tagline.strip():
        return env_tagline.strip()
    return 'DILG Laguna PDMU Operations System'


def _default_assistant_name(app_name=None):
    try:
        if current_app and current_app.config.get('ASSISTANT_NAME'):
            return str(current_app.config['ASSISTANT_NAME']).strip()
    except Exception:
        pass
    env_name = os.environ.get('ASSISTANT_NAME')
    if env_name and env_name.strip():
        return env_name.strip()
    base_app = app_name or _default_app_name()
    return f'{base_app} Assistant'


def _asset_url(asset_type, configs):
    config = configs.get(f'branding_{asset_type}_file_path')
    if config is None or not config.value:
        return None
    version = config.updated_at.strftime('%Y%m%d%H%M%S') if config.updated_at else ''
    suffix = f'?v={version}' if version else ''
    return f'/api/settings/branding/{asset_type}{suffix}'


def _upsert_config(*, key, value, value_type, description, user_id):
    config = SystemConfig.query.filter_by(key=key).first()
    if config is None:
        config = SystemConfig(
            key=key,
            value_type=value_type,
            description=description,
            updated_by=int(user_id),
        )
        db.session.add(config)
    config.value = value
    config.value_type = value_type
    config.description = description
    config.is_encrypted = False
    config.updated_by = int(user_id)
    return config


def _branding_description(key):
    descriptions = {
        'app_name': 'Public application name shown across the frontend and assistant surfaces.',
        'app_tagline': 'Short public description shown in the browser title and landing page.',
        'branding_logo_alt_text': 'Accessible alternative text for the uploaded application logo.',
    }
    return descriptions.get(key, 'Application branding setting.')


def _branding_upload_root():
    configured = current_app.config.get('BRANDING_UPLOAD_DIR')
    root = Path(configured) if configured else Path(current_app.instance_path) / 'uploads' / 'branding'
    root.mkdir(parents=True, exist_ok=True)
    return root


def _report_template_upload_root():
    configured = current_app.config.get('REPORT_TEMPLATE_UPLOAD_DIR')
    root = Path(configured) if configured else Path(current_app.instance_path) / 'uploads' / 'report_templates'
    root.mkdir(parents=True, exist_ok=True)
    return root


def _load_document_template_catalog(template_type, configs=None):
    definition = DOCUMENT_TEMPLATE_TYPES.get(template_type)
    if not definition:
        return []
    key = f'{definition["config_prefix"]}_templates'
    raw_value = None
    if configs is not None and key in configs:
        raw_value = configs[key].value
    elif configs is None:
        raw_value = _config_by_key(key)

    catalog = []
    if raw_value:
        try:
            parsed = json.loads(raw_value)
            if isinstance(parsed, list):
                catalog = [item for item in parsed if isinstance(item, dict)]
        except (TypeError, ValueError):
            catalog = []

    if catalog:
        return catalog

    legacy = _legacy_document_template_entry(template_type, configs=configs)
    return [legacy] if legacy else []


def _load_report_template_catalog(template_type, configs=None):
    return _load_document_template_catalog(template_type, configs=configs)


def _legacy_document_template_entry(template_type, configs=None):
    definition = DOCUMENT_TEMPLATE_TYPES.get(template_type)
    if not definition:
        return None
    prefix = definition['config_prefix']

    def get_config_value(key):
        if configs is not None:
            config = configs.get(key)
            return config.value if config and config.value not in (None, '') else None
        return _config_by_key(key)

    path = get_config_value(f'{prefix}_file_path')
    if not path:
        return None
    return {
        'id': 'legacy-current',
        'label': definition['label'],
        'file_name': get_config_value(f'{prefix}_file_name') or Path(path).name,
        'file_path': path,
        'mime_type': get_config_value(f'{prefix}_mime_type') or definition['mime_type'],
        'uploaded_at': None,
        'uploaded_by': None,
        'placeholder_defaults': {},
        'placeholder_definitions': [],
        'placeholder_fields': [],
        'is_default': True,
    }


def _legacy_report_template_entry(template_type, configs=None):
    return _legacy_document_template_entry(template_type, configs=configs)


def _serialize_document_template_entry(template_type, entry, default_sources=None):
    definition = DOCUMENT_TEMPLATE_TYPES.get(template_type) or {}
    candidate = _resolve_under_root(entry.get('file_path'), _report_template_upload_root())
    available = bool(candidate and candidate.exists() and candidate.is_file())
    placeholder_definitions = _entry_placeholder_definitions(entry, candidate if available else None)
    placeholder_suggestions = _placeholder_default_suggestions(
        placeholder_definitions,
        default_sources or [],
        exclude_template_type=template_type,
        exclude_template_id=entry.get('id'),
    )
    return {
        'id': entry.get('id'),
        'label': entry.get('label') or entry.get('file_name') or 'DOCX template',
        'file_name': entry.get('file_name'),
        'mime_type': entry.get('mime_type') or definition.get('mime_type'),
        'uploaded_at': entry.get('uploaded_at'),
        'uploaded_by': entry.get('uploaded_by'),
        'placeholder_defaults': entry.get('placeholder_defaults') or {},
        'placeholder_default_suggestions': placeholder_suggestions,
        'placeholder_fields': [item['key'] for item in placeholder_definitions],
        'placeholder_definitions': placeholder_definitions,
        'is_default': bool(entry.get('is_default')),
        'available': available,
    }


def _serialize_report_template_entry(template_type, entry):
    return _serialize_document_template_entry(template_type, entry)


def _workflow_template_groups(template_types):
    allowed_types = set(template_types or ())
    return [
        _workflow_template_group(group)
        for group in list_document_templates()
        if group.get('template_type') in allowed_types
    ]


def _workflow_template_group(group):
    return {
        'template_type': group.get('template_type'),
        'label': group.get('label'),
        'description': group.get('description'),
        'category': group.get('category'),
        'configured': group.get('configured'),
        'placeholder_fields': group.get('placeholder_fields') or [],
        'placeholder_definitions': group.get('placeholder_definitions') or [],
        'default_template_id': group.get('default_template_id'),
        'templates': [
            _workflow_template_entry(template)
            for template in group.get('templates') or []
        ],
    }


def _workflow_template_entry(entry):
    return {
        'id': entry.get('id'),
        'label': entry.get('label'),
        'file_name': entry.get('file_name'),
        'placeholder_defaults': entry.get('placeholder_defaults') or {},
        'placeholder_fields': entry.get('placeholder_fields') or [],
        'placeholder_definitions': entry.get('placeholder_definitions') or [],
        'is_default': bool(entry.get('is_default')),
        'available': bool(entry.get('available')),
    }


def _sync_current_document_template_config(template_type, catalog, user_id):
    definition = DOCUMENT_TEMPLATE_TYPES[template_type]
    current = next((item for item in catalog if item.get('is_default')), None)
    if current is None and catalog:
        current = catalog[-1]
    _upsert_config(
        key=f'{definition["config_prefix"]}_file_path',
        value=current.get('file_path') if current else '',
        value_type='string',
        description=f'Stored DOCX document template path for {definition["label"]}.',
        user_id=user_id,
    )
    _upsert_config(
        key=f'{definition["config_prefix"]}_file_name',
        value=current.get('file_name') if current else '',
        value_type='string',
        description=f'Original DOCX document template file name for {definition["label"]}.',
        user_id=user_id,
    )
    _upsert_config(
        key=f'{definition["config_prefix"]}_mime_type',
        value=current.get('mime_type') if current else definition['mime_type'],
        value_type='string',
        description=f'MIME type for {definition["label"]} document template.',
        user_id=user_id,
    )


def _first_available_template_entry(entries):
    root = _report_template_upload_root()
    return next(
        (item for item in entries if _is_safe_existing_file(Path(str(item.get('file_path') or '')), root)),
        None,
    )


def _path_is_referenced(path, entries):
    try:
        resolved_path = Path(path).resolve()
    except (OSError, RuntimeError):
        return False
    for entry in entries:
        candidate = _resolve_under_root(entry.get('file_path'), _report_template_upload_root())
        if candidate and candidate.resolve() == resolved_path:
            return True
    return False


def _default_document_template_id(entries):
    default = next((item for item in entries if item.get('is_default') and item.get('available')), None)
    if default:
        return default.get('id')
    first_available = next((item for item in entries if item.get('available')), None)
    return first_available.get('id') if first_available else None


def _default_report_template_id(entries):
    return _default_document_template_id(entries)


def _clean_template_label(value):
    if value in (None, ''):
        return None
    return str(value).strip()[:120] or None


def _parse_placeholder_defaults(value):
    if value in (None, ''):
        return {}
    if isinstance(value, dict):
        parsed = value
    elif isinstance(value, str):
        try:
            parsed = json.loads(value)
        except (TypeError, ValueError) as exc:
            raise SettingsError('placeholder_defaults must be a valid JSON object.', 422) from exc
    else:
        raise SettingsError('placeholder_defaults must be an object.', 422)
    if not isinstance(parsed, dict):
        raise SettingsError('placeholder_defaults must be an object.', 422)
    return {
        str(key).strip(): '' if val is None else str(val)
        for key, val in parsed.items()
        if str(key).strip()
    }


def _parse_template_default_flag(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, int) and not isinstance(value, bool):
        if value in (0, 1):
            return value != 0
        raise SettingsError('Default flag must be true or false.', 422)
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in {'true', '1', 'yes', 'on'}:
            return True
        if normalized in {'false', '0', 'no', 'off', ''}:
            return False
    raise SettingsError('Default flag must be true or false.', 422)


def _merge_unambiguous_placeholder_defaults(placeholder_definitions, explicit_defaults, default_sources):
    merged = dict(explicit_defaults or {})
    suggestions = _placeholder_default_suggestions(
        placeholder_definitions,
        default_sources,
    )
    for definition in placeholder_definitions or []:
        key = definition.get('key') if isinstance(definition, dict) else None
        if not key or _has_placeholder_default_value(merged.get(key)):
            continue
        key_suggestions = suggestions.get(key) or []
        if len(key_suggestions) == 1:
            merged[key] = key_suggestions[0]['value']
    return merged


def _placeholder_default_suggestions(placeholder_definitions, default_sources, exclude_template_type=None, exclude_template_id=None):
    wanted_keys = {}
    for definition in placeholder_definitions or []:
        key = definition.get('key') if isinstance(definition, dict) else None
        lookup_key = _placeholder_default_lookup_key(key)
        if key and lookup_key and lookup_key not in wanted_keys:
            wanted_keys[lookup_key] = key
    if not wanted_keys:
        return {}

    suggestions = {key: {} for key in wanted_keys.values()}
    for source in default_sources or []:
        if (
            exclude_template_type
            and source.get('template_type') == exclude_template_type
            and source.get('template_id') == exclude_template_id
        ):
            continue
        defaults = source.get('placeholder_defaults') or {}
        if not isinstance(defaults, dict):
            continue
        for source_key, raw_value in defaults.items():
            if not _has_placeholder_default_value(raw_value):
                continue
            lookup_key = _placeholder_default_lookup_key(source_key)
            target_key = wanted_keys.get(lookup_key)
            if not target_key:
                continue
            value = _clean_placeholder_default_value(raw_value)
            by_value = suggestions[target_key]
            if value not in by_value:
                by_value[value] = {
                    'value': value,
                    'count': 0,
                    'source_templates': [],
                }
            by_value[value]['count'] += 1
            source_template = {
                'template_type': source.get('template_type'),
                'template_type_label': source.get('template_type_label'),
                'template_id': source.get('template_id'),
                'template_label': source.get('template_label'),
                'file_name': source.get('file_name'),
                'placeholder_key': source_key,
            }
            if source_template not in by_value[value]['source_templates']:
                by_value[value]['source_templates'].append(source_template)

    return {
        key: sorted(
            values.values(),
            key=lambda item: (-item['count'], item['value'].casefold()),
        )
        for key, values in suggestions.items()
        if values
    }


def _document_template_default_sources(catalogs=None):
    catalogs = catalogs or {
        template_type: _load_document_template_catalog(template_type)
        for template_type in DOCUMENT_TEMPLATE_TYPES
    }
    sources = []
    for template_type, catalog in catalogs.items():
        definition = DOCUMENT_TEMPLATE_TYPES.get(template_type) or {}
        for entry in catalog or []:
            if not isinstance(entry, dict):
                continue
            sources.append({
                'template_type': template_type,
                'template_type_label': definition.get('label') or template_type,
                'template_id': entry.get('id'),
                'template_label': entry.get('label') or entry.get('file_name') or definition.get('label') or 'DOCX template',
                'file_name': entry.get('file_name'),
                'placeholder_defaults': entry.get('placeholder_defaults') or {},
            })
    return sources


def _placeholder_default_lookup_key(value):
    if value in (None, ''):
        return ''
    return str(value).strip().casefold()


def _has_placeholder_default_value(value):
    return bool(_clean_placeholder_default_value(value))


def _clean_placeholder_default_value(value):
    if value is None:
        return ''
    return str(value).strip()


def _validate_report_template_size(upload):
    max_size = int(current_app.config.get('MAX_SUBMISSION_FILE_BYTES') or (25 * 1024 * 1024))
    size = _measure_upload_size(upload)
    if size > max_size:
        raise SettingsError(f'Document templates must be {max_size // (1024 * 1024)} MB or smaller.', 422)


def _validate_docx_file(path):
    try:
        validate_ooxml_container(path, expected_kind='docx')
        document_template_service._validate_docx_template_safety(path)
    except (ArchiveSecurityError, document_template_service.DocumentTemplateError) as exc:
        raise SettingsError(getattr(exc, 'message', str(exc)), 422) from exc
    except (zipfile.BadZipFile, OSError) as exc:
        raise SettingsError('Uploaded document template must be a valid DOCX file.', 422) from exc


def _document_template_placeholders(template_type, serialized_templates=None):
    serialized_templates = serialized_templates or []
    default_template = next((item for item in serialized_templates if item.get('is_default') and item.get('placeholder_definitions')), None)
    first_template = next((item for item in serialized_templates if item.get('placeholder_definitions')), None)
    if default_template or first_template:
        return (default_template or first_template).get('placeholder_definitions') or []

    if template_type == 'project_revision_report':
        return [_placeholder_definition(key) for key in [
            'report_title',
            'generated_at_display',
            'date',
            'project_name',
            'source_project_code',
            'lgu_name',
            'program_name',
            'fiscal_year',
            'revision_count',
            'revision_rows',
            'revision_rows_text',
            'director_name',
            'director_title',
            'office_name',
            'prepared_by_name',
            'prepared_by_title',
            'signatory_name',
            'signatory_title',
        ]]
    if template_type in {'annex_e', 'annex_m', 'annex_n', 'annex_o', 'general_document'}:
        return [_placeholder_definition(key) for key in [
            'date',
            'notice_title',
            'notice_label',
            'notice_date_display',
            'subject_line',
            'lce_name_upper',
            'lce_position',
            'lgu_display',
            'lce_surname',
            'program_title',
            'program_name',
            'fiscal_year',
            'project_name',
            'project_names_text',
            'project_rows_text',
            'findings_text',
            'recommendations_text',
            'timelines_text',
            'budget_allocation_text',
            'compliance_period_text',
            'default_body_text',
            'custom_body_text',
            'content',
            'signatory_name',
            'signatory_title',
            'region_name',
            'regional_office_address',
            'regional_office_website',
        ]]
    return []


def _report_template_placeholders(template_type):
    return [item['key'] for item in _document_template_placeholders(template_type)]


def _entry_placeholder_definitions(entry, candidate=None):
    definitions = entry.get('placeholder_definitions') if isinstance(entry, dict) else None
    if definitions:
        normalized = []
        for item in definitions:
            definition = _normalize_placeholder_definition(item)
            if definition:
                normalized.append(definition)
        return normalized
    fields = entry.get('placeholder_fields') if isinstance(entry, dict) else None
    if fields:
        return [_placeholder_definition(field) for field in fields]
    if candidate:
        try:
            return document_template_service.extract_docx_placeholders(candidate)
        except document_template_service.DocumentTemplateError:
            return []
    return []


def _normalize_placeholder_definition(item):
    if isinstance(item, str):
        return _placeholder_definition(item)
    if not isinstance(item, dict):
        return None
    key = str(item.get('key') or document_template_service.placeholder_key(item.get('label') or '')).strip()
    if not key:
        return None
    return {
        'key': key,
        'label': str(item.get('label') or document_template_service.label_for_key(key)),
        'syntax': item.get('syntax') if isinstance(item.get('syntax'), list) else [],
    }


def _placeholder_definition(key):
    normalized_key = document_template_service.placeholder_key(key)
    return {
        'key': normalized_key,
        'label': document_template_service.label_for_key(normalized_key),
        'syntax': [],
    }


def _validate_brand_asset_mime_type(filename, upload):
    allowed_mime_types = set(current_app.config.get('BRANDING_ALLOWED_MIME_TYPES') or ())
    allowed = ', '.join(sorted(allowed_mime_types))
    return resolve_upload_mime_type(
        filename,
        upload,
        allowed_mime_types,
        lambda message: SettingsError(
            f'Branding image MIME type must be one of: {allowed}.' if 'MIME type' in message else message,
            422,
        ),
        label='Branding image',
    )


def _validate_brand_asset_size(upload):
    max_size = int(current_app.config.get('MAX_BRANDING_FILE_BYTES') or 0)
    if not max_size:
        return
    size = _measure_upload_size(upload)
    if size > max_size:
        raise SettingsError(f'Branding images must be {max_size // 1024} KB or smaller.', 422)


def _measure_upload_size(upload):
    content_length = getattr(upload, 'content_length', None)
    if content_length is not None and int(content_length) > 0:
        return int(content_length)

    stream = getattr(upload, 'stream', None)
    if stream is None:
        return 0

    position = stream.tell()
    stream.seek(0, os.SEEK_END)
    size_bytes = stream.tell()
    stream.seek(position)
    return int(size_bytes)


def _validate_brand_image_signature(path, extension):
    try:
        header = path.read_bytes()[:16]
    except OSError as exc:
        raise SettingsError('Unable to read branding image file.', 422) from exc

    if len(header) < 4:
        raise SettingsError('Branding image file is truncated or empty.', 422)

    if extension == 'png' and not header.startswith(b'\x89PNG\r\n\x1a\n'):
        raise SettingsError('Branding image content does not match the PNG file type.', 422)
    elif extension in ('jpg', 'jpeg') and not header.startswith(b'\xff\xd8\xff'):
        raise SettingsError('Branding image content does not match the JPEG file type.', 422)
    elif extension == 'webp' and (not header.startswith(b'RIFF') or header[8:12] != b'WEBP'):
        raise SettingsError('Branding image content does not match the WEBP file type.', 422)


def _validate_svg_file(path):
    raise SettingsError('SVG branding images are not supported for security reasons.', 422)


def _database_check():
    uri = current_app.config.get('SQLALCHEMY_DATABASE_URI')
    if uri:
        return _check('database', 'ready', 'Database connection string is configured.')
    return _check('database', 'warning', 'Database connection string is missing.')


def _jwt_check():
    secret = current_app.config.get('JWT_SECRET_KEY')
    if secret and len(secret) >= 16:
        return _check('jwt', 'ready', 'JWT signing secret is configured.')
    return _check('jwt', 'warning', 'JWT signing secret is missing or too short.')


def _smtp_check():
    from app.services.integration_config_service import resolve_email_runtime_value
    smtp_host = resolve_email_runtime_value('SMTP_HOST')
    if smtp_host:
        return _check('smtp', 'ready', f'SMTP delivery is configured for {smtp_host}.')
    return _check('smtp', 'warning', 'SMTP delivery is not configured. Temporary passwords require manual handoff.')


def _ai_provider_check():
    configured = []
    from app.services.system_admin_service import resolve_ai_runtime_value
    gemini_key = resolve_ai_runtime_value('GEMINI_API_KEY')
    gemini_model = resolve_ai_runtime_value('GEMINI_MODEL') or 'gemini-3.5-flash'
    if _has_real_secret(gemini_key):
        configured.append(f"Gemini ({gemini_model})")

    groq_key = resolve_ai_runtime_value('GROQ_API_KEY')
    groq_model = resolve_ai_runtime_value('GROQ_MODEL') or 'llama-3.1-8b-instant'
    if _has_real_secret(groq_key):
        configured.append(f"Groq ({groq_model})")

    openrouter_key = resolve_ai_runtime_value('OPENROUTER_API_KEY')
    openrouter_model = resolve_ai_runtime_value('OPENROUTER_MODEL') or 'openai/gpt-4o-mini'
    if _has_real_secret(openrouter_key):
        configured.append(f"OpenRouter ({openrouter_model})")

    raw_order = resolve_ai_runtime_value('AI_PROVIDER_ORDER') or ('gemini', 'groq', 'openrouter')
    if isinstance(raw_order, str):
        order_list = [item.strip() for item in raw_order.split(',') if item.strip()]
    elif isinstance(raw_order, (list, tuple)):
        order_list = [str(item).strip() for item in raw_order if str(item).strip()]
    else:
        order_list = ['gemini', 'groq', 'openrouter']

    if configured:
        order_str = ', '.join(order_list)
        return _check('ai_assist', 'ready', f"{get_assistant_name()} providers configured: {', '.join(configured)}. Failover order: {order_str}.")
    return _check('ai_assist', 'warning', f'{get_assistant_name()} has no configured provider API keys.')


def _document_intelligence_check():
    warnings = []
    try:
        import pypdf  # noqa: F401
    except Exception:
        warnings.append('PDF text extraction dependency is unavailable')

    try:
        import openpyxl  # noqa: F401
    except Exception:
        warnings.append('Excel text extraction dependency is unavailable')

    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
    except Exception:
        warnings.append('image OCR Python dependencies are unavailable')

    if not shutil.which('tesseract'):
        warnings.append('Tesseract OCR binary is unavailable')

    try:
        from app.services.document_intelligence_service import is_provider_refinement_enabled
        refinement_active = is_provider_refinement_enabled()
    except Exception:
        refinement_active = bool(current_app.config.get('AI_DOCUMENT_CLASSIFICATION_ENABLED'))

    ai_status = (
        'provider refinement enabled'
        if refinement_active
        else 'provider refinement disabled by default'
    )
    max_pages = current_app.config.get('DOCUMENT_OCR_MAX_PAGES') or 3

    if warnings:
        return _check(
            'document_intelligence',
            'warning',
            f"Document intake scanning is partially available ({'; '.join(warnings)}; {ai_status}; max pages: {max_pages}).",
        )
    return _check(
        'document_intelligence',
        'ready',
        f'Document intake scanning and OCR are available ({ai_status}; max pages: {max_pages}).',
    )


def _google_drive_check():
    if not current_app.config.get('GOOGLE_DRIVE_ENABLED'):
        return _check('google_drive', 'disabled', 'Google Drive handoff is disabled in this deployment.')

    root_folder_id = _config_by_key('google_drive_root_folder_id') or current_app.config.get('GOOGLE_DRIVE_ROOT_FOLDER_ID')
    service_account = current_app.config.get('GOOGLE_DRIVE_SERVICE_ACCOUNT_FILE') or current_app.config.get('GOOGLE_DRIVE_SERVICE_ACCOUNT_JSON')
    try:
        oauth_token = _config_by_key('google_drive_oauth_token')
    except SettingsError as exc:
        if not _is_decryption_error(exc):
            raise
        oauth_token = None
        return _check('google_drive', 'warning', UNREADABLE_GOOGLE_DRIVE_TOKEN_MESSAGE)
    oauth_client = bool(current_app.config.get('GOOGLE_OAUTH_CLIENT_ID') and current_app.config.get('GOOGLE_OAUTH_CLIENT_SECRET'))
    if root_folder_id and oauth_token and oauth_client:
        return _check('google_drive', 'ready', 'Google Drive handoff is configured with OAuth access.')
    if root_folder_id and service_account and not oauth_token:
        service_account_warning = _google_drive_service_account_warning()
        if service_account_warning:
            return _check('google_drive', 'warning', service_account_warning)
        mode = 'service account'
        return _check('google_drive', 'ready', f'Google Drive handoff is configured with {mode} access.')
    missing = []
    if not service_account and not oauth_token:
        missing.append('Drive account connection')
    if (oauth_token or not service_account) and not oauth_client:
        missing.append('Google OAuth client')
    if not root_folder_id:
        missing.append('root folder id')
    return _check(
        'google_drive',
        'warning',
        f"Google Drive handoff is enabled but missing: {', '.join(missing)}.",
    )


def _google_drive_service_account_warning():
    raw_info = current_app.config.get('GOOGLE_DRIVE_SERVICE_ACCOUNT_JSON')
    file_path = current_app.config.get('GOOGLE_DRIVE_SERVICE_ACCOUNT_FILE')

    if raw_info:
        try:
            info = json.loads(raw_info)
        except json.JSONDecodeError:
            return 'Google Drive service account JSON is not valid JSON.'
        return _google_drive_service_account_shape_warning(info, 'JSON')

    if file_path:
        credential_path = Path(file_path)
        if not credential_path.exists():
            return 'Google Drive service account file does not exist.'
        try:
            info = json.loads(credential_path.read_text(encoding='utf-8'))
        except json.JSONDecodeError:
            return 'Google Drive service account file is not valid JSON.'
        except OSError:
            return 'Google Drive service account file cannot be read.'
        return _google_drive_service_account_shape_warning(info, 'file')

    return None


def _google_drive_service_account_shape_warning(info, source_label):
    if not isinstance(info, dict):
        return f'Google Drive service account {source_label} must be a JSON object.'
    missing = sorted(
        field for field in GOOGLE_DRIVE_SERVICE_ACCOUNT_REQUIRED_FIELDS
        if not str(info.get(field) or '').strip()
    )
    if missing:
        return f"Google Drive service account {source_label} is missing required field(s): {', '.join(missing)}."
    return None


def _object_storage_check():
    ready, message = object_storage_service.provider_ready()
    return _check('object_storage', 'ready' if ready else 'warning', message)


def _uploads_check():
    upload_dir = current_app.config.get('SUBMISSION_UPLOAD_DIR')
    if not upload_dir:
        return _check('uploads', 'warning', 'Submission upload directory is not configured.')

    path = Path(upload_dir)
    if path.exists() or path.parent.exists():
        return _check('uploads', 'ready', f'Submission upload storage is pointed at {path}.')
    return _check('uploads', 'warning', 'Submission upload directory is configured but the path is not available on disk.')


def _guideline_uploads_check():
    upload_dir = current_app.config.get('GUIDELINE_UPLOAD_DIR')
    if not upload_dir:
        base_upload_dir = current_app.config.get('SUBMISSION_UPLOAD_DIR')
        upload_dir = Path(base_upload_dir) / 'guidelines' if base_upload_dir else Path(current_app.instance_path) / 'uploads' / 'guidelines'
    path = Path(upload_dir)
    if path.exists() or path.parent.exists():
        return _check('guideline_library', 'ready', f'Guideline upload storage is pointed at {path}.')
    return _check('guideline_library', 'warning', 'Guideline upload storage is configured but the path is not available on disk.')


def _celery_check():
    broker = current_app.config.get('CELERY_BROKER_URL')
    backend = current_app.config.get('CELERY_RESULT_BACKEND')
    if broker and backend:
        return _check('celery', 'ready', 'Celery broker and result backend are configured.')
    return _check('celery', 'warning', 'Celery broker or result backend configuration is incomplete.')


def _has_real_secret(value):
    normalized = str(value or '').strip().lower()
    if not normalized:
        return False
    return normalized not in {'placeholder', 'changeme', 'change-me', 'none'} and not normalized.startswith('your-')


def _check(key, status, message):
    return {
        'key': key,
        'status': status,
        'message': message,
    }
