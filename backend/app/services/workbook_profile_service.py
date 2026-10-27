import csv
import hashlib
import json
import re
from pathlib import Path

from flask import current_app

from app.services.import_field_catalog import normalize_label, match_canonical_field


class WorkbookProfileError(Exception):
    def __init__(self, message, status_code=400):
        self.message = message
        self.status_code = status_code


PROJECT_CODE_RE = re.compile(r'\b([A-Z]{2,}|FA|G3P|SAFPB|FALGU)[A-Z0-9-]{6,}\b', re.IGNORECASE)


def profile_workbook(file_path, file_type=None):
    path = Path(file_path)
    detected_type = (file_type or path.suffix.lower().lstrip('.')).lower()
    if detected_type == 'csv':
        sheets = [_profile_csv(path)]
    elif detected_type == 'xlsx':
        sheets = _profile_xlsx(path)
    elif detected_type == 'xls':
        raise WorkbookProfileError('Legacy .xls workbooks are not supported; save as .xlsx or CSV first.', 422)
    else:
        raise WorkbookProfileError('Unsupported import file type.', 422)

    profile = {
        'file_type': detected_type,
        'sheets': sheets,
    }
    profile = _json_safe(profile)
    profile['schema_signature'] = build_schema_signature(profile)
    return profile


def extract_rows(file_path, profile, sheet_name=None):
    path = Path(file_path)
    file_type = profile.get('file_type') or path.suffix.lower().lstrip('.')
    sheets = profile.get('sheets', [])
    if sheet_name:
        sheets = [sheet for sheet in sheets if sheet.get('name') == sheet_name]
    if not sheets:
        return []
    sheet_names = {sheet.get('name') for sheet in sheets if sheet.get('name')}

    if file_type == 'csv':
        matrix_by_sheet = {'CSV Import': _read_csv_matrix(path)}
    else:
        matrix_by_sheet = _read_xlsx_matrices(path, sheet_names=sheet_names)

    extracted = []
    for sheet in sheets:
        matrix = matrix_by_sheet.get(sheet.get('name'))
        if matrix is None:
            continue
        columns = sheet.get('columns') or []
        for row_number in range(sheet.get('data_start_row') or 1, len(matrix) + 1):
            row = matrix[row_number - 1]
            values = {}
            for column in columns:
                col_index = column['index'] - 1
                value = _clean(row[col_index]) if col_index < len(row) else None
                values[column['column']] = _json_safe(value)
            if any(value is not None for value in values.values()):
                extracted.append({
                    'sheet_name': sheet.get('name'),
                    'sheet_type': sheet.get('sheet_type'),
                    'row_number': row_number,
                    'values': values,
                })
    return extracted


def build_schema_signature(profile):
    parts = []
    for sheet in profile.get('sheets', []):
        headers = sorted(
            column.get('normalized_header') or column.get('header') or ''
            for column in sheet.get('columns', [])
            if column.get('normalized_header') or column.get('header')
        )
        parts.append({
            'sheet': normalize_label(sheet.get('name')),
            'type': sheet.get('sheet_type'),
            'headers': headers,
        })
    payload = json.dumps(parts, sort_keys=True, default=str)
    return hashlib.sha256(payload.encode('utf-8')).hexdigest()


def _profile_csv(path):
    rows = _read_csv_matrix(path)
    return _profile_matrix('CSV Import', rows, merged_ranges=[])


def _profile_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise WorkbookProfileError('openpyxl is required to import Excel workbooks.', 500) from exc

    workbook = load_workbook(path, data_only=False, read_only=False)
    values_workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            values_worksheet = values_workbook[worksheet.title]
            formula_rows = _worksheet_rows(worksheet)
            value_rows = _worksheet_rows(values_worksheet)
            rows = _best_profile_rows(value_rows, formula_rows)
            sheets.append(_profile_matrix(
                worksheet.title,
                rows,
                merged_ranges=[str(item) for item in worksheet.merged_cells.ranges],
                formula_rows=formula_rows,
            ))
        return sheets
    finally:
        workbook.close()
        values_workbook.close()


def _read_csv_matrix(path):
    rows = []
    with path.open('r', encoding='utf-8-sig', newline='') as handle:
        reader = csv.reader(handle)
        for row_number, row in enumerate(reader, start=1):
            if row_number > _max_profile_rows():
                raise WorkbookProfileError(f'Workbook profiling is limited to {_max_profile_rows()} rows.', 422)
            if len(row) > _max_profile_columns():
                raise WorkbookProfileError(f'Workbook profiling is limited to {_max_profile_columns()} columns.', 422)
            rows.append(row)
    return rows


def _read_xlsx_matrices(path, sheet_names=None):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise WorkbookProfileError('openpyxl is required to import Excel workbooks.', 500) from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    matrices = {}
    try:
        for worksheet in workbook.worksheets:
            if sheet_names and worksheet.title not in sheet_names:
                continue
            matrices[worksheet.title] = _worksheet_rows(worksheet)
        return matrices
    finally:
        workbook.close()


def _worksheet_rows(worksheet):
    if worksheet.max_row and worksheet.max_row > _max_profile_rows():
        raise WorkbookProfileError(f'Workbook profiling is limited to {_max_profile_rows()} rows.', 422)
    if worksheet.max_column and worksheet.max_column > _max_profile_columns():
        raise WorkbookProfileError(f'Workbook profiling is limited to {_max_profile_columns()} columns.', 422)
    max_row = min(worksheet.max_row or 1, _max_profile_rows())
    max_col = min(worksheet.max_column or 1, _max_profile_columns())
    rows = []
    for row in worksheet.iter_rows(max_row=max_row, max_col=max_col):
        rows.append([cell.value for cell in row])
    return rows


def _max_profile_rows():
    return max(int(current_app.config.get('MAX_IMPORT_PROFILE_ROWS') or 5000), 1)


def _max_profile_columns():
    return max(int(current_app.config.get('MAX_IMPORT_PROFILE_COLUMNS') or 200), 1)


def _best_profile_rows(value_rows, formula_rows):
    if any(_looks_like_project_code(value) for row in value_rows[:100] for value in row[:10]):
        return value_rows
    value_count = sum(1 for row in value_rows for value in row if _clean(value) is not None)
    formula_count = sum(1 for row in formula_rows for value in row if _clean(value) is not None)
    if value_count >= max(5, int(formula_count * 0.6)):
        return value_rows
    return formula_rows


def _profile_matrix(sheet_name, rows, merged_ranges=None, formula_rows=None):
    trimmed_rows = _trim_rows(rows)
    max_row = len(trimmed_rows)
    max_column = max((len(row) for row in trimmed_rows), default=0)
    normalized_rows = [row + [None] * (max_column - len(row)) for row in trimmed_rows]
    header_candidates = _header_candidates(normalized_rows)
    data_start = _detect_data_start(normalized_rows, header_candidates)
    header_rows = _select_header_rows(normalized_rows, data_start, header_candidates)
    columns = _build_columns(normalized_rows, header_rows, data_start)
    sheet_type = _classify_sheet(sheet_name, columns)

    return {
        'name': sheet_name,
        'sheet_type': sheet_type,
        'max_row': max_row,
        'max_column': max_column,
        'merged_ranges': merged_ranges or [],
        'formula_count': _formula_count(formula_rows or normalized_rows),
        'header_candidates': header_candidates[:5],
        'data_start_row': data_start,
        'header_rows': header_rows,
        'columns': columns,
        'sample_rows': _sample_rows(normalized_rows, columns, data_start),
    }


def _trim_rows(rows):
    last_row = 0
    last_col = 0
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row, start=1):
            if _clean(value) is not None:
                last_row = max(last_row, row_index)
                last_col = max(last_col, col_index)
    return [list(row[:last_col]) for row in rows[:last_row]]


def _header_candidates(rows):
    candidates = []
    for index, row in enumerate(rows[:100], start=1):
        nonempty = [_clean(value) for value in row if _clean(value) is not None]
        if len(nonempty) < 3:
            continue
        text_count = sum(1 for value in nonempty if isinstance(value, str))
        formula_count = sum(1 for value in nonempty if isinstance(value, str) and value.startswith('='))
        numeric_count = sum(1 for value in nonempty if isinstance(value, (int, float)))
        score = text_count * 1.5 + len(nonempty) * 0.3 - formula_count * 2 - numeric_count * 0.2 - index * 0.02
        candidates.append({
            'row': index,
            'score': round(score, 3),
            'nonempty_count': len(nonempty),
            'text_count': text_count,
            'formula_count': formula_count,
        })
    return sorted(candidates, key=lambda item: item['score'], reverse=True)


def _detect_data_start(rows, header_candidates):
    for index, row in enumerate(rows, start=1):
        if index <= 1:
            continue
        nonempty = [_clean(value) for value in row if _clean(value) is not None]
        if len(nonempty) < 4:
            continue
        first_values = row[: min(10, len(row))]
        if any(_looks_like_project_code(value) for value in first_values):
            return index

    for index, row in enumerate(rows, start=1):
        if index <= 1:
            continue
        nonempty = [_clean(value) for value in row if _clean(value) is not None]
        if len(nonempty) < 4:
            continue
        first_values = row[: min(10, len(row))]
        text_values = [str(value).strip().lower() for value in first_values if isinstance(value, str)]
        numeric_count = sum(1 for value in row if isinstance(value, (int, float)))
        if numeric_count >= 2 and any(word in ' '.join(text_values) for word in ['city', 'municipality', 'province', 'laguna', 'cavite', 'batangas', 'quezon', 'rizal']):
            return index
    if header_candidates:
        return header_candidates[0]['row'] + 1
    return 2 if len(rows) > 1 else 1


def _select_header_rows(rows, data_start, header_candidates):
    candidate_rows = set()
    for candidate in header_candidates:
        row_number = candidate['row']
        if row_number < data_start and candidate['text_count'] >= 2:
            candidate_rows.add(row_number)
    for row_number in range(max(1, data_start - 4), data_start):
        row = rows[row_number - 1]
        if sum(1 for value in row if isinstance(_clean(value), str)) >= 2:
            candidate_rows.add(row_number)
    return sorted(candidate_rows)


def _build_columns(rows, header_rows, data_start):
    max_column = max((len(row) for row in rows), default=0)
    columns = []
    for col_index in range(max_column):
        header_path = []
        for row_number in header_rows:
            value = _clean(rows[row_number - 1][col_index])
            if value is not None and not _is_numeric_marker(value):
                header_path.append(str(value))
        sample_values = []
        for row in rows[data_start - 1:data_start + 9]:
            if col_index < len(row):
                value = _clean(row[col_index])
                if value is not None:
                    sample_values.append(_json_safe(value))
        if not header_path and not sample_values:
            continue
        header = ' / '.join(dict.fromkeys(header_path))
        columns.append({
            'column': _column_letter(col_index + 1),
            'index': col_index + 1,
            'header_path': header_path,
            'header': header,
            'normalized_header': normalize_label(header),
            'sample_values': sample_values[:5],
        })
    return columns


def _sample_rows(rows, columns, data_start):
    samples = []
    for row_number in range(data_start, min(len(rows), data_start + 5) + 1):
        row = rows[row_number - 1]
        values = {}
        for column in columns:
            col_index = column['index'] - 1
            values[column['column']] = _json_safe(_clean(row[col_index]) if col_index < len(row) else None)
        if any(value is not None for value in values.values()):
            samples.append({'row_number': row_number, 'values': values})
    return samples


def _classify_sheet(sheet_name, columns):
    text = normalize_label(sheet_name + ' ' + ' '.join(column.get('header') or '' for column in columns))
    if 'summary' in text:
        return 'summary'
    if 'project reporting' in text or any(month in text.split() for month in ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']):
        return 'project_reporting'
    if 'physical accomplishment' in text or 'disbursement' in text or 'obligation' in text or 'project monitoring' in text:
        return 'implementation_monitoring'
    if 'requirement' in text or 'annex' in text or 'ded' in text:
        return 'pre_implementation_requirements'
    if 'project title' in text and ('allocation' in text or 'lgu' in text):
        return 'project_master'
    recognized = {match['target_field'] for column in columns
                  if (match := match_canonical_field(column.get('header_path') or [column.get('header')]))}
    if ('project.project_name' in recognized
            and recognized.intersection({'project.budget_amount', 'lgu.name', 'lgu.municipality'})):
        return 'project_master'
    if 'fund release' in text:
        return 'fund_release'
    if 'notice to implement' in text or 'nti' in text:
        return 'notice_to_implement'
    if 'congressional district' in text or 'district lgu' in text:
        return 'reference_lookup'
    return 'unknown'


def _formula_count(rows):
    return sum(1 for row in rows for value in row if isinstance(value, str) and value.startswith('='))


def _clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        stripped = re.sub(r'\s+', ' ', value).strip()
        return stripped or None
    return value


def _json_safe(value):
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    try:
        json.dumps(value)
        return value
    except TypeError:
        return str(value)


def _looks_like_project_code(value):
    if value is None:
        return False
    text = str(value).strip()
    if text.startswith('='):
        return False
    if not re.search(r'\d', text):
        return False
    if text.count('-') < 2 and not re.search(r'\b(FA|G3P|SAFPB|FALGU|LGSF)[A-Z0-9-]*\b', text, re.IGNORECASE):
        return False
    return bool(re.search(r'\b(FA|G3P|SAFPB|FALGU|LGSF)[A-Z0-9-]{4,}\b|[A-Z]{2,}[A-Z0-9]*-\d[\w-]*-\d[\w-]*', text, re.IGNORECASE))


def _is_numeric_marker(value):
    return isinstance(value, (int, float)) or (isinstance(value, str) and value.strip().isdigit())


def _column_letter(index):
    result = ''
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result
