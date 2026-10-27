import base64
import os
import re
from io import BytesIO
from datetime import datetime, timezone
import textwrap

from flask import current_app, render_template

from app.extensions import db
from app.models import CommunicationLog, EscalationNotice, LegalFinding, LguProfile, NoticeRecipient, Project, ProjectRequirement, SlaDeadline
from app.services import access_service, audit_service, document_template_service, settings_service


VALID_NOTICE_TYPES = ('annex_e', 'annex_m', 'annex_n', 'annex_o')
VALID_NOTICE_STATUSES = ('draft', 'reviewed', 'sent', 'acknowledged', 'expired')
VALID_OUTPUT_FORMATS = ('html', 'pdf', 'docx')
VALID_RELEASE_CHANNELS = ('email', 'dms', 'intranet', 'manual_print')
MANUAL_NOTICE_TYPES = ('annex_e',)
NOTICE_LABELS = {
    'annex_e': 'Annex E Notice of Deficiencies',
    'annex_m': 'Notice of Non-Compliance No. 1',
    'annex_n': 'Notice of Non-Compliance No. 2',
    'annex_o': 'Notice of Non-Compliance No. 3',
}
NOTICE_FILENAME_PREFIXES = {
    'annex_e': 'annex-e-notice-of-deficiencies',
    'annex_m': 'annex-m-notice-of-non-compliance-no-1',
    'annex_n': 'annex-n-notice-of-non-compliance-no-2',
    'annex_o': 'annex-o-notice-of-non-compliance-no-3',
}


class NoticeError(Exception):
    def __init__(self, message, status_code=400):
        self.message = message
        self.status_code = status_code
        super().__init__(message)


def create_notice(
    project_id,
    notice_type,
    generated_by=None,
    trigger_sla_id=None,
    trigger_project_requirement_id=None,
    parent_notice_id=None,
    title=None,
    final_text=None,
    status='draft',
    legal_findings=None,
    user_id=None,
    generated_at=None,
):
    if notice_type not in VALID_NOTICE_TYPES:
        raise NoticeError(f'Invalid notice_type. Must be one of: {", ".join(VALID_NOTICE_TYPES)}', 422)
    if notice_type not in MANUAL_NOTICE_TYPES:
        raise NoticeError('Manual notice creation is only allowed for Annex E notices. Use the escalation engine for later notices.', 422)
    if status == 'sent':
        raise NoticeError('Notices cannot be created in sent status. Use the release action to mark a notice released.', 422)
    if status not in ('draft', None):
        raise NoticeError("New notices must be created with status 'draft'.", 422)
    status = 'draft'

    project = _get_project_with_access(project_id, user_id=user_id)
    try:
        access_service.assert_project_movement_allowed(project)
    except access_service.AccessError as exc:
        raise NoticeError(exc.message, exc.status_code) from exc
    lgu = db.session.get(LguProfile, project.lgu_id)
    deadline = _get_deadline_with_access(trigger_sla_id, user_id=user_id) if trigger_sla_id else None
    requirement = None
    if trigger_project_requirement_id is not None:
        requirement = _get_requirement_with_access(trigger_project_requirement_id, user_id=user_id)
        if requirement.project_id != project.id:
            raise NoticeError('Project requirement does not belong to this project.', 422)
    elif deadline and deadline.project_requirement_id:
        requirement = db.session.get(ProjectRequirement, deadline.project_requirement_id)
        trigger_project_requirement_id = requirement.id if requirement else None

    if deadline and deadline.project_id != project.id:
        raise NoticeError('SLA deadline does not belong to this project.', 422)

    body_text = build_default_body_text(
        notice_type=notice_type,
        project=project,
        lgu=lgu,
        requirement=requirement,
        deadline=deadline,
        parent_notice=None,
    )
    notice = EscalationNotice(
        project_id=project.id,
        parent_notice_id=parent_notice_id,
        trigger_sla_id=deadline.id if deadline else None,
        trigger_project_requirement_id=trigger_project_requirement_id,
        generated_by=_coerce_optional_user_id(generated_by if generated_by is not None else user_id),
        notice_type=notice_type,
        notice_number=_next_notice_number(project.id, notice_type),
        title=title or NOTICE_LABELS[notice_type],
        body_text=body_text,
        final_text=final_text,
        status=status,
        generated_at=generated_at or datetime.now(timezone.utc),
    )
    db.session.add(notice)
    db.session.flush()

    recipient = _build_recipient(notice.id, lgu)
    if recipient:
        db.session.add(recipient)

    try:
        _replace_legal_findings(
            notice,
            legal_findings if legal_findings is not None else _default_legal_findings(requirement, deadline),
        )

        audit_service.log_action(
            user_id=_coerce_optional_user_id(user_id if user_id is not None else generated_by),
            action='generate',
            entity_type='escalation_notice',
            entity_id=notice.id,
            old_value=None,
            new_value=_serialize_notice_snapshot(notice),
            commit=False,
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return notice


def update_notice(notice_id, data, user_id=None):
    notice = _get_notice_with_access(notice_id, user_id=user_id)
    try:
        access_service.assert_project_movement_allowed(notice.project_id)
    except access_service.AccessError as exc:
        raise NoticeError(exc.message, exc.status_code) from exc
    before = _serialize_notice_snapshot(notice)

    if 'title' in data and data['title']:
        notice.title = data['title']
    if 'body_text' in data:
        notice.body_text = (data['body_text'] or '').strip() or None
    if 'final_text' in data:
        notice.final_text = data['final_text'] or None
    if 'status' in data and data['status']:
        if data['status'] not in VALID_NOTICE_STATUSES:
            raise NoticeError(f'Invalid status. Must be one of: {", ".join(VALID_NOTICE_STATUSES)}', 422)
        if data['status'] == 'sent':
            raise NoticeError('Use the release action to mark a notice released.', 422)
        if data['status'] == 'acknowledged' and notice.sent_at is None:
            raise NoticeError('Notice must be marked released before receipt can be recorded.', 409)
        notice.status = data['status']
        if data['status'] == 'acknowledged':
            acknowledged_at = (
                _parse_optional_datetime(data.get('acknowledged_at'))
                or notice.acknowledged_at
                or datetime.now(timezone.utc)
            )
            notice.acknowledged_at = acknowledged_at
            for recipient in notice.recipients.order_by(NoticeRecipient.id.asc()).all():
                recipient.delivery_status = 'acknowledged'
                if recipient.acknowledged_at is None or data.get('acknowledged_at'):
                    recipient.acknowledged_at = acknowledged_at

    try:
        if 'legal_findings' in data:
            _replace_legal_findings(notice, data['legal_findings'])

        audit_service.log_action(
            user_id=_coerce_optional_user_id(user_id),
            action='update',
            entity_type='escalation_notice',
            entity_id=notice.id,
            old_value=before,
            new_value=_serialize_notice_snapshot(notice),
            commit=False,
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return notice


def release_notice(notice_id, data, user_id=None):
    notice = _get_notice_with_access(notice_id, user_id=user_id)
    try:
        access_service.assert_project_movement_allowed(notice.project_id)
    except access_service.AccessError as exc:
        raise NoticeError(exc.message, exc.status_code) from exc
    if notice.status == 'acknowledged':
        raise NoticeError('This notice has already been received.', 409)
    if notice.sent_at is not None:
        raise NoticeError('This notice has already been marked released.', 409)

    channel = data.get('channel')
    if channel not in VALID_RELEASE_CHANNELS:
        raise NoticeError(f'channel must be one of: {", ".join(VALID_RELEASE_CHANNELS)}', 422)

    released_at = _parse_optional_datetime(data.get('released_at'), field_name='released_at') or datetime.now(timezone.utc)
    before_notice = _serialize_notice_snapshot(notice)
    communication = CommunicationLog(
        lgu_id=notice.project.lgu_id,
        sent_by=_coerce_optional_user_id(user_id),
        communication_type='escalation_notice',
        channel=channel,
        subject=notice.title,
        body=notice.final_text or notice.body_text,
        recipient_email=_default_release_recipient_email(notice, channel),
        recipient_name=_clean_text(data.get('recipient_name')) or _default_release_recipient_name(notice),
        recipient_destination=_clean_text(data.get('recipient_destination')),
        recipient_office=_clean_text(data.get('recipient_office')),
        reference_no=_clean_text(data.get('reference_no')),
        related_notice_id=notice.id,
        delivery_status='sent',
        sent_at=released_at,
        posted_at=released_at,
        remarks=_clean_text(data.get('remarks')),
    )
    db.session.add(communication)

    notice.status = 'sent'
    notice.sent_at = released_at
    for recipient in notice.recipients.order_by(NoticeRecipient.id.asc()).all():
        recipient.delivery_method = channel
        recipient.delivery_status = 'sent'
        recipient.sent_at = released_at

    try:
        db.session.flush()

        audit_service.log_action(
            user_id=_coerce_optional_user_id(user_id),
            action='release',
            entity_type='escalation_notice',
            entity_id=notice.id,
            old_value=before_notice,
            new_value=_serialize_notice_snapshot(notice),
            commit=False,
        )
        audit_service.log_action(
            user_id=_coerce_optional_user_id(user_id),
            action='create',
            entity_type='communication_log',
            entity_id=communication.id,
            old_value=None,
            new_value=_serialize_communication(communication),
            commit=False,
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return {
        'notice': notice,
        'communication': _serialize_communication(communication),
    }


def record_notice_receipt(notice_id, data, user_id=None):
    notice = _get_notice_with_access(notice_id, user_id=user_id)
    try:
        access_service.assert_project_movement_allowed(notice.project_id)
    except access_service.AccessError as exc:
        raise NoticeError(exc.message, exc.status_code) from exc
    if notice.sent_at is None:
        raise NoticeError('Notice must be marked released before receipt can be recorded.', 409)

    acknowledged_at = _parse_optional_datetime(data.get('acknowledged_at')) or datetime.now(timezone.utc)
    sent_at = _to_aware_datetime(notice.sent_at)
    if sent_at and acknowledged_at < sent_at:
        raise NoticeError('Receipt timestamp cannot be earlier than the release timestamp.', 422)

    before = _serialize_notice_snapshot(notice)
    notice.status = 'acknowledged'
    notice.acknowledged_at = acknowledged_at
    for recipient in notice.recipients.order_by(NoticeRecipient.id.asc()).all():
        recipient.delivery_status = 'acknowledged'
        recipient.acknowledged_at = acknowledged_at

    try:
        audit_service.log_action(
            user_id=_coerce_optional_user_id(user_id),
            action='receipt',
            entity_type='escalation_notice',
            entity_id=notice.id,
            old_value=before,
            new_value=_serialize_notice_snapshot(notice),
            commit=False,
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return notice


def render_notice_document(notice_id, output_format='pdf', user_id=None, template_id=None, placeholder_overrides=None):
    notice = _get_notice_with_access(notice_id, user_id=user_id)
    output_format = (output_format or 'pdf').lower()
    if output_format not in VALID_OUTPUT_FORMATS:
        raise NoticeError(f'Invalid format. Must be one of: {", ".join(VALID_OUTPUT_FORMATS)}', 422)

    template_info = _get_notice_document_template(notice.notice_type, template_id=template_id) if output_format == 'docx' else None
    context = _apply_notice_document_placeholders(build_notice_context(notice), template_info, placeholder_overrides)
    filename_base = _build_filename_base(notice, context)

    if output_format == 'html':
        html = _render_notice_html(notice.notice_type, context)
        return html.encode('utf-8'), 'text/html; charset=utf-8', f'{filename_base}.html'

    if output_format == 'pdf':
        html = _render_notice_html(notice.notice_type, context)
        try:
            from app.services.pdf_service import safe_weasyprint_url_fetcher
            HTML = _import_weasyprint_html()
            pdf_bytes = HTML(string=html, base_url=current_app.root_path, url_fetcher=safe_weasyprint_url_fetcher).write_pdf()
        except Exception as exc:
            current_app.logger.warning(
                'WeasyPrint PDF rendering failed for notice %s; using text fallback instead: %s',
                notice.id,
                exc,
            )
            pdf_bytes = _render_fallback_pdf(context)
        return pdf_bytes, 'application/pdf', f'{filename_base}.pdf'

    try:
        template_path = template_info.get('path') if template_info else _ensure_docx_template(notice.notice_type)
        docx_bytes = document_template_service.render_docx_template(template_path, context)
        docx_bytes = _finalize_notice_docx(docx_bytes, context)
    except Exception as exc:
        current_app.logger.warning(
            'DOCX template rendering failed for notice %s; using python-docx fallback instead: %s',
            notice.id,
            exc,
        )
        docx_bytes = _render_fallback_docx(context)
    return docx_bytes, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', f'{filename_base}.docx'


def preview_notice_document(notice_id, user_id=None, template_id=None, placeholder_overrides=None, draft_payload=None):
    notice = _get_notice_with_access(notice_id, user_id=user_id)
    template_info = _get_notice_document_template(notice.notice_type, template_id=template_id)
    context = _apply_notice_document_placeholders(
        build_notice_context(notice, draft_payload=draft_payload),
        template_info,
        placeholder_overrides,
    )
    try:
        template_path = template_info.get('path') if template_info else _ensure_docx_template(notice.notice_type)
        docx_bytes = document_template_service.render_docx_template(template_path, context)
        docx_bytes = _finalize_notice_docx(docx_bytes, context)
    except Exception as exc:
        current_app.logger.warning(
            'DOCX preview rendering failed for notice %s; using python-docx fallback instead: %s',
            notice.id,
            exc,
        )
        docx_bytes = _render_fallback_docx(context)
    return {
        'template': _serialize_template_preview(template_info),
        'variables': document_template_service.build_template_variables(template_info, context, placeholder_overrides),
        'docx_base64': base64.b64encode(docx_bytes).decode('utf-8'),
        'preview_text': document_template_service.preview_docx_text(docx_bytes),
        'preview_document': document_template_service.preview_docx_document(docx_bytes),
    }


def build_default_body_text(notice_type, project, lgu, requirement=None, deadline=None, parent_notice=None):
    project_name = project.project_name if project else 'the subject project'
    requirement_name = requirement.requirement_name if requirement else 'the outstanding compliance requirement'
    program_name = project.program.name if project and project.program else 'the identified program'
    fiscal_year = project.fiscal_year if project else 'the covered fiscal year'
    annex_e_date = parent_notice.generated_at.strftime('%B %d, %Y') if parent_notice and parent_notice.generated_at else 'the referenced Annex E notice'
    parent_date = parent_notice.generated_at.strftime('%B %d, %Y') if parent_notice and parent_notice.generated_at else 'the previous notice'

    is_non_infra = bool(project and getattr(project, 'project_kind', None) in ('goods', 'financial_assistance', 'motor_vehicles', 'vehicle'))

    if notice_type == 'annex_e':
        deadline_window = deadline.sla_working_days if deadline and deadline.sla_working_days else 15
        if is_non_infra:
            return (
                f'Based on the evaluation conducted for {project_name}, the following deficiencies, missing documentation, or non-compliance findings in the project implementation were noted in relation to '
                f'{requirement_name} under {program_name} for FY {fiscal_year}.\n\n'
                f'You are hereby given the prescribed and agreed timelines to submit the required documentary compliance, rectify identified discrepancies, or provide written explanation within '
                f'{deadline_window} working days from receipt of this Notice of Deficiencies.'
            )
        return (
            f'Based on the evaluation conducted for {project_name}, the following defects and deficiencies were noted in relation to '
            f'{requirement_name} under {program_name} for FY {fiscal_year}.\n\n'
            f'You are hereby given the prescribed and agreed timelines to address the findings and comply within '
            f'{deadline_window} working days from receipt of this Notice of Deficiencies.'
        )

    if notice_type == 'annex_m':
        if is_non_infra:
            return (
                f'This pertains to our letter containing the findings and recommendations dated {annex_e_date} in relation to the '
                f'milestone, delivery, and distribution status of the project funded under {program_name} for FY {fiscal_year}.\n\n'
                f'Please be reminded that the above-mentioned findings and corresponding recommendations regarding delivery milestones '
                f'and distribution status are for your immediate action and compliance. Continued inaction shall result in the '
                f'issuance of another notification and endorsement to the appropriate office.'
            )
        return (
            f'This pertains to our letter containing the findings and recommendations dated {annex_e_date} in relation to the '
            f'implementation of the project funded under {program_name} for FY {fiscal_year}.\n\n'
            f'Please be reminded that the above-mentioned findings and corresponding recommendations are for your immediate '
            f'action and compliance. Continued inaction shall result in the issuance of another notification and endorsement '
            f'to the appropriate office.'
        )

    if notice_type == 'annex_n':
        return (
            f'This is in reference to the Notice of Non-Compliance No. 1 dated {parent_date}.\n\n'
            f'To date, we have yet to receive reports on the actions taken relative to the findings on {project_name}. '
            f'You are hereby given fifteen (15) days from receipt of this letter to comply as required in the previous notice.'
        )

    if notice_type == 'annex_o':
        return (
            f'This is in reference to the earlier notices issued relative to the findings on {project_name}.\n\n'
            f'Please be advised that continued inaction to address the findings on this project shall result in the endorsement '
            f'of a Non-Compliance Report to the DILG Central Office for final disposition. You are given fifteen (15) days from '
            f'receipt of this letter to comply as required in the preceding notice.'
        )

    return ''


def build_notice_context(notice, draft_payload=None):
    project = db.session.get(Project, notice.project_id)
    lgu = db.session.get(LguProfile, project.lgu_id) if project else None
    requirement = (
        db.session.get(ProjectRequirement, notice.trigger_project_requirement_id)
        if notice.trigger_project_requirement_id else None
    )
    deadline = db.session.get(SlaDeadline, notice.trigger_sla_id) if notice.trigger_sla_id else None
    parent_notice = db.session.get(EscalationNotice, notice.parent_notice_id) if notice.parent_notice_id else None
    recipients = notice.recipients.order_by(NoticeRecipient.id.asc()).all()
    notice_chain = _get_notice_chain_lookup(notice)

    draft_payload = draft_payload or {}
    draft_title = draft_payload.get('title')
    draft_body = draft_payload.get('body_text')
    draft_final = draft_payload.get('final_text')
    draft_findings = draft_payload.get('legal_findings')

    lce_name = (lgu.mayor_name or lgu.contact_person or lgu.name or 'Local Chief Executive').strip() if lgu else 'Local Chief Executive'
    lce_surname = _extract_surname(lce_name)
    province = lgu.province if lgu and lgu.province else 'Laguna'
    region_name = current_app.config.get('DILG_REGION_NAME') or 'Region IV-A (CALABARZON)'
    location_line = _build_location_line(lgu, region_name=region_name)
    funding_year = (
        project.program_cycle.fiscal_year
        if project and project.program_cycle and project.program_cycle.fiscal_year
        else (project.fiscal_year if project else '')
    )
    prior_nnc_dates_display = _prior_nnc_dates_display(notice_chain, parent_notice)
    provincial_director_name = _default_provincial_director_name()
    regional_director_name = _default_regional_director_name()
    signatory_name = _default_signatory_name(notice.notice_type)
    signatory_title = _default_signatory_title(notice.notice_type, province)
    mlgoo_name = lgu.contact_person if lgu and lgu.contact_person else 'MLGOO'

    default_body = draft_body if draft_body is not None else (notice.body_text or build_default_body_text(
        notice.notice_type,
        project,
        lgu,
        requirement=requirement,
        deadline=deadline,
        parent_notice=parent_notice,
    ))
    custom_body = draft_final if draft_final is not None else (notice.final_text or '')
    if draft_findings is not None:
        findings = draft_findings
        serialized_findings = [_serialize_legal_finding(finding) for finding in findings]
    else:
        findings = _context_findings(notice, parent_notice, notice_chain)
        serialized_findings = [_serialize_legal_finding(finding) for finding in findings]
    project_rows = _build_notice_project_rows(notice, project, requirement, deadline, findings)
    project_names_text = '\n'.join(item['project_name'] for item in project_rows if item.get('project_name'))
    budget_allocation_text = '\n'.join(
        item['budget_allocation_display']
        for item in project_rows
        if item.get('budget_allocation_display')
    )
    compliance_period_text = '\n'.join(
        item['compliance_period_display']
        for item in project_rows
        if item.get('compliance_period_display')
    )
    context = {
        'notice_type': notice.notice_type,
        'notice_label': NOTICE_LABELS.get(notice.notice_type, notice.notice_type.replace('_', ' ').title()),
        'notice_title': draft_title if draft_title is not None else notice.title,
        'notice_date_display': _format_notice_date(notice.generated_at),
        'subject_line': _subject_line(notice),
        'lce_name': lce_name,
        'lce_name_upper': lce_name.upper(),
        'lce_position': _default_lce_position(lgu),
        'lgu_name': lgu.name if lgu else 'LGU',
        'lgu_display': location_line,
        'lgu_office_heading': _default_lgu_office_heading(lgu),
        'lce_surname': lce_surname,
        'program_title': project.program.name if project and project.program else 'Program',
        'program_name': project.program.name if project and project.program else 'Program',
        'program_cycle_name': project.program_cycle.name if project and project.program_cycle else '',
        'program_cycle_code': project.program_cycle.cycle_code if project and project.program_cycle else '',
        'fiscal_year': funding_year,
        'funding_year': funding_year,
        'project_name': project.project_name if project else 'Project',
        'project_names_text': project_names_text,
        'budget_allocation_text': budget_allocation_text,
        'project_rows': project_rows,
        'project_rows_text': '\n'.join(_format_project_row_text(item) for item in project_rows),
        'default_body_text': default_body,
        'default_body_paragraphs': _split_paragraphs(default_body),
        'custom_body_text': custom_body,
        'custom_body_paragraphs': _split_paragraphs(custom_body),
        'annex_e_date_display': _notice_chain_date_display(notice_chain, parent_notice, 'annex_e'),
        'annex_m_date_display': _notice_chain_date_display(notice_chain, parent_notice, 'annex_m'),
        'annex_n_date_display': _notice_chain_date_display(notice_chain, parent_notice, 'annex_n'),
        'prior_nnc_dates_display': prior_nnc_dates_display,
        'annex_m_and_n_date_display': prior_nnc_dates_display,
        'signatory_name': signatory_name,
        'signatory_title': signatory_title,
        'provincial_director_name': provincial_director_name,
        'provincial_director_title': f'Provincial Director, {province}',
        'regional_director_name': regional_director_name,
        'regional_director_title': 'Regional Director',
        'province_name': province,
        'province': province,
        'mlgoo_name': mlgoo_name,
        'recipient_lines': [_format_recipient_line(recipient) for recipient in recipients],
        'findings_rows': serialized_findings,
        'findings_text': _join_context_values(
            [item['formal_findings'] for item in serialized_findings if item['formal_findings']],
            [item['finding_summary'] for item in project_rows if item.get('finding_summary')],
        ),
        'recommendations_text': _join_context_values(
            [item['recommendations'] for item in serialized_findings if item['recommendations']],
            [item['recommendation_summary'] for item in project_rows if item.get('recommendation_summary')],
        ),
        'timelines_text': _join_context_values(
            [item['agreements_timelines'] for item in serialized_findings if item['agreements_timelines']],
            [item['compliance_period_display'] for item in project_rows if item.get('compliance_period_display')],
        ),
        'compliance_period_text': compliance_period_text,
        'remarks_text': '\n\n'.join(item['remarks'] for item in serialized_findings if item['remarks']),
        'date': _format_notice_date(notice.generated_at),
        'content': '\n\n'.join(part for part in [default_body, custom_body] if part),
        'region_name': region_name,
        'regional_office_address': current_app.config.get('DILG_REGIONAL_OFFICE_ADDRESS') or '',
        'regional_office_website': current_app.config.get('DILG_REGIONAL_OFFICE_WEBSITE') or '',
        'aide_memoire_date_display': compliance_period_text or _format_notice_date(notice.generated_at),
        'project_kind': getattr(project, 'project_kind', None),
        'is_non_infra': bool(project and getattr(project, 'project_kind', None) in ('goods', 'financial_assistance', 'motor_vehicles', 'vehicle')),
    }
    return context


def _get_notice_document_template(notice_type, template_id=None):
    try:
        return settings_service.get_document_template(notice_type, template_id=template_id)
    except settings_service.SettingsError as exc:
        raise NoticeError(exc.message, exc.status_code)


def _apply_notice_document_placeholders(context, template_info=None, placeholder_overrides=None):
    merged = dict(context)
    defaults = (template_info or {}).get('placeholder_defaults') or {}
    allowed_keys = set(merged) | {
        item.get('key')
        for item in (template_info or {}).get('placeholder_definitions', [])
        if item.get('key')
    }
    for source in (defaults, placeholder_overrides or {}):
        for key, value in _nonblank_notice_placeholder_values(source, allowed_keys).items():
            merged[key] = value
    return document_template_service.expand_context_aliases(merged)


def _nonblank_notice_placeholder_values(values, allowed_keys):
    if not isinstance(values, dict):
        return {}
    cleaned = {}
    for key, value in values.items():
        cleaned_key = str(key).strip()
        if cleaned_key not in allowed_keys or value is None:
            continue
        cleaned_value = str(value)
        if not cleaned_value.strip():
            continue
        cleaned[cleaned_key] = cleaned_value
    return cleaned


def _serialize_template_preview(template_info):
    if not template_info:
        return {
            'id': None,
            'label': 'Built-in DOCX layout',
            'file_name': None,
            'placeholder_definitions': [],
        }
    return {
        'id': template_info.get('id'),
        'label': template_info.get('label'),
        'file_name': template_info.get('file_name'),
        'placeholder_definitions': template_info.get('placeholder_definitions') or [],
    }


def _render_notice_html(notice_type, context):
    template_name = f'pdf/{notice_type}.html'
    return render_template(template_name, **context)


def _render_fallback_pdf(context):
    lines = _build_fallback_pdf_lines(context)
    content_lines = ['BT', '/F1 11 Tf', '50 770 Td', '14 TL']
    first_line = True
    for raw_line in lines:
        escaped = _escape_pdf_text(raw_line)
        if first_line:
            content_lines.append(f'({escaped}) Tj')
            first_line = False
        else:
            content_lines.append(f'T* ({escaped}) Tj')
    content_lines.append('ET')
    content_stream = '\n'.join(content_lines).encode('latin-1', errors='replace')
    return _build_pdf_document(content_stream)


def _render_fallback_docx(context):
    Document = _import_docx_document()
    doc = Document()
    notice_label = context.get('notice_label') or context.get('notice_title') or 'Notice'
    notice_date = context.get('notice_date_display') or ''
    recipient = context.get('lce_name_upper') or context.get('lce_name') or 'Local Chief Executive'
    recipient_position = context.get('lce_position') or ''
    recipient_lgu = context.get('lgu_display') or ''
    subject_line = context.get('subject_line') or context.get('notice_title') or notice_label

    doc.add_heading(notice_label, level=1)
    if notice_date:
        doc.add_paragraph(notice_date)
    doc.add_paragraph(recipient)
    if recipient_position:
        doc.add_paragraph(recipient_position)
    if recipient_lgu:
        doc.add_paragraph(recipient_lgu)
    doc.add_paragraph(f'Subject: {subject_line}')

    for paragraph in context.get('default_body_paragraphs') or []:
        doc.add_paragraph(paragraph)

    findings_rows = context.get('findings_rows') or []
    if findings_rows:
        table = doc.add_table(rows=1, cols=4)
        headers = ['Findings / Observations', 'Recommendations', 'Agreements / Timelines', 'Remarks']
        for index, label in enumerate(headers):
            table.cell(0, index).text = label
        for item in findings_rows:
            row = table.add_row().cells
            row[0].text = item.get('formal_findings') or ''
            row[1].text = item.get('recommendations') or ''
            row[2].text = item.get('agreements_timelines') or ''
            row[3].text = item.get('remarks') or ''

    project_rows_text = context.get('project_rows_text')
    if project_rows_text:
        doc.add_paragraph('Projects / Findings:')
        doc.add_paragraph(project_rows_text)

    for paragraph in context.get('custom_body_paragraphs') or []:
        doc.add_paragraph(paragraph)

    doc.add_paragraph('Very truly yours,')
    doc.add_paragraph(context.get('signatory_name') or '')
    doc.add_paragraph(context.get('signatory_title') or '')

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def _finalize_notice_docx(content, context):
    try:
        Document = _import_docx_document()
        doc = Document(BytesIO(content))
    except Exception:
        return content

    changed = False
    for paragraph in _iter_notice_docx_paragraphs(doc):
        changed = _replace_notice_date_blank(paragraph, context) or changed
        changed = _replace_notice_non_infra_docx_text(paragraph, context) or changed

    changed = _sync_notice_recipient_block(doc, context) or changed

    for table in doc.tables:
        changed = _fill_notice_project_table(table, context) or changed

    if not changed:
        return document_template_service.repair_docx_styles(content)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return document_template_service.repair_docx_styles(buffer.read())


def _iter_notice_docx_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _replace_notice_date_blank(paragraph, context):
    text = paragraph.text or ''
    if '___' not in text or 'dated' not in text.casefold():
        return False

    replacement = (
        context.get('annex_e_date_display')
        or context.get('aide_memoire_date_display')
        or context.get('notice_date_display')
        or context.get('date')
        or ''
    )
    if not replacement:
        return False

    changed = False
    for run in paragraph.runs:
        if '___' not in (run.text or ''):
            continue
        run.text = re.sub(r'_{3,}', replacement, run.text, count=1)
        changed = True

    if changed:
        return True

    updated = re.sub(
        r'(?i)(dated\s+)_+',
        lambda match: f'{match.group(1)}{replacement}',
        text,
        count=1,
    )
    if updated == text:
        return False
    paragraph.text = updated
    return True


# NOTE: Matches canonical phrasing from the DILG Omnibus Annex M template.
# If template text is edited or customized, ensure these target phrases stay in sync
# or migrate to token/regex-based replacement to avoid silent no-ops.
def _replace_notice_non_infra_docx_text(paragraph, context):
    if not context.get('is_non_infra'):
        return False
    text = paragraph.text or ''
    changed = False
    if 'in relation to the implementation of the following project' in text:
        text = text.replace(
            'in relation to the implementation of the following project',
            'in relation to the milestone, delivery, and distribution status of the following project',
        )
        changed = True
    if 'above-mentioned finding(s) with the corresponding recommendation(s)' in text and 'delivery milestones' not in text:
        text = text.replace(
            'above-mentioned finding(s) with the corresponding recommendation(s)',
            'above-mentioned finding(s) with the corresponding recommendation(s) regarding delivery milestones and distribution status',
        )
        changed = True
    if changed:
        paragraph.text = text
        return True
    return False


def _sync_notice_recipient_block(doc, context):
    paragraphs = [paragraph for paragraph in doc.paragraphs if (paragraph.text or '').strip()]
    if not paragraphs:
        return False

    subject_index = next(
        (
            index for index, paragraph in enumerate(paragraphs[:10])
            if paragraph.text.strip().casefold().startswith('subject')
        ),
        None,
    )
    if subject_index is None:
        return False

    search_limit = subject_index
    hon_index = next(
        (
            index for index, paragraph in enumerate(paragraphs[:search_limit])
            if _is_hon_recipient_line(paragraph.text)
        ),
        None,
    )
    if hon_index is None:
        return False

    changed = False
    notice_date = context.get('notice_date_display') or context.get('date') or ''
    if hon_index > 0 and notice_date and _looks_like_notice_date_line(paragraphs[hon_index - 1].text):
        changed = _set_paragraph_text(paragraphs[hon_index - 1], notice_date) or changed

    lce_name = context.get('lce_name_upper') or context.get('lce_name') or ''
    if lce_name:
        changed = _set_paragraph_text(paragraphs[hon_index], f'HON. {lce_name}') or changed

    if subject_index is not None and hon_index + 1 < subject_index:
        lce_position = context.get('lce_position') or ''
        if lce_position:
            changed = _set_paragraph_text(paragraphs[hon_index + 1], lce_position) or changed

    if subject_index is not None and hon_index + 2 < subject_index:
        lgu_display = context.get('lgu_display') or ''
        if lgu_display:
            changed = _set_paragraph_text(paragraphs[hon_index + 2], lgu_display) or changed

    return changed


def _is_hon_recipient_line(value):
    return bool(re.match(r'^\s*HON\.\s+', str(value or ''), flags=re.IGNORECASE))


def _looks_like_notice_date_line(value):
    text = str(value or '').strip()
    if not text:
        return False
    if '<date>' in text.casefold() or '{{ date' in text.casefold():
        return True
    return bool(re.match(
        r'^(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+\d{4}$',
        text,
        flags=re.IGNORECASE,
    ))


def _set_paragraph_text(paragraph, value):
    value = '' if value is None else str(value)
    if (paragraph.text or '') == value:
        return False
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(value)
    return True


def _fill_notice_project_table(table, context):
    rows = context.get('project_rows') or []
    if not rows or not table.rows:
        return False
    table_kind = _notice_project_table_kind(table)
    if not table_kind:
        return False

    table_element = table._tbl
    for existing_row in list(table.rows)[1:]:
        table_element.remove(existing_row._tr)

    for item in rows:
        cells = table.add_row().cells
        cells[0].text = item.get('project_name') or ''
        if table_kind == 'allocation':
            cells[1].text = item.get('budget_allocation_display') or ''
        else:
            cells[1].text = item.get('finding_summary') or ''
            cells[2].text = item.get('recommendation_summary') or ''
            cells[3].text = item.get('compliance_period_display') or ''
    return True


def _notice_project_table_kind(table):
    if len(table.rows[0].cells) < 2:
        return False
    headers = [_canonical_table_label(cell.text) for cell in table.rows[0].cells]
    if (
        'project' in headers[0]
        and any(label in headers[1] for label in ('allocation', 'budget', 'amount', 'nadai'))
    ):
        return 'allocation'
    if len(headers) < 4:
        return False
    if (
        'project' in headers[0]
        and 'finding' in headers[1]
        and 'recommendation' in headers[2]
        and (
            'compliance period' in headers[3]
            or 'timeline' in headers[3]
            or 'date' in headers[3]
        )
    ):
        return 'findings'
    return False


def _canonical_table_label(value):
    cleaned = re.sub(r'\s+', ' ', str(value or '')).strip().casefold()
    cleaned = re.sub(r'[^a-z0-9]+', ' ', cleaned)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip()
    if 'project' in cleaned:
        return 'project'
    if 'finding' in cleaned:
        return 'finding'
    if 'recommendation' in cleaned:
        return 'recommendation'
    if 'allocation' in cleaned:
        return 'allocation'
    if 'budget' in cleaned:
        return 'budget'
    if 'amount' in cleaned:
        return 'amount'
    if 'nadai' in cleaned:
        return 'nadai'
    if 'compliance' in cleaned and 'period' in cleaned:
        return 'compliance period'
    if 'timeline' in cleaned:
        return 'timeline'
    if 'date' in cleaned:
        return 'date'
    return cleaned


def _ensure_docx_template(notice_type):
    template_dir = os.path.join(current_app.root_path, 'templates', 'docx')
    os.makedirs(template_dir, exist_ok=True)
    template_path = os.path.join(template_dir, f'{notice_type}.docx')
    if os.path.exists(template_path):
        return template_path

    builder = {
        'annex_e': _build_annex_e_docx_template,
        'annex_m': _build_annex_m_docx_template,
        'annex_n': _build_annex_n_docx_template,
        'annex_o': _build_annex_o_docx_template,
    }.get(notice_type)
    if not builder:
        raise NoticeError(f'No DOCX template builder is available for {notice_type}.', 500)
    builder(template_path)
    return template_path


def _build_annex_e_docx_template(path):
    Document = _import_docx_document()
    doc = Document()
    doc.add_paragraph('Republic of the Philippines')
    doc.add_paragraph('{{ lgu_office_heading }}')
    doc.add_paragraph('{{ lgu_display }}')
    doc.add_paragraph('{{ notice_label | upper }}')
    doc.add_paragraph('{{ notice_date_display }}')
    doc.add_paragraph('Dear Sir / Madam:')
    doc.add_paragraph('{{ default_body_text }}')
    table = doc.add_table(rows=2, cols=4)
    headers = ['Findings / Observations', 'Recommendations', 'Agreements / Timelines', 'Remarks']
    for index, label in enumerate(headers):
        table.cell(0, index).text = label
    table.cell(1, 0).text = '{{ findings_text }}'
    table.cell(1, 1).text = '{{ recommendations_text }}'
    table.cell(1, 2).text = '{{ timelines_text }}'
    table.cell(1, 3).text = '{{ remarks_text }}'
    doc.add_paragraph('{{ custom_body_text }}')
    doc.add_paragraph('Very truly yours,')
    doc.add_paragraph('{{ signatory_name }}')
    doc.add_paragraph('{{ signatory_title }}')
    doc.save(path)


def _build_annex_m_docx_template(path):
    Document = _import_docx_document()
    doc = _base_notice_doc()
    doc.add_paragraph('{{ default_body_text }}')
    _add_notice_project_table(doc)
    doc.add_paragraph('{{ custom_body_text }}')
    doc.add_paragraph('Very truly yours,')
    doc.add_paragraph('{{ signatory_name }}')
    doc.add_paragraph('{{ signatory_title }}')
    doc.save(path)


def _build_annex_n_docx_template(path):
    Document = _import_docx_document()
    doc = _base_notice_doc()
    doc.add_paragraph('{{ default_body_text }}')
    _add_notice_project_table(doc)
    doc.add_paragraph('{{ custom_body_text }}')
    doc.add_paragraph('Very truly yours,')
    doc.add_paragraph('{{ signatory_name }}')
    doc.add_paragraph('{{ signatory_title }}')
    doc.add_paragraph('cc: {{ mlgoo_name }}')
    doc.save(path)


def _build_annex_o_docx_template(path):
    Document = _import_docx_document()
    doc = _base_notice_doc()
    doc.add_paragraph('{{ default_body_text }}')
    _add_notice_project_table(doc)
    doc.add_paragraph('{{ custom_body_text }}')
    doc.add_paragraph('Very truly yours,')
    doc.add_paragraph('{{ signatory_name }}')
    doc.add_paragraph('{{ signatory_title }}')
    doc.add_paragraph('cc: {{ mlgoo_name }}')
    doc.add_paragraph('DILG – Provincial Office')
    doc.add_paragraph('DILG – OUSLG')
    doc.add_paragraph('Regional COA')
    doc.save(path)


def _base_notice_doc():
    Document = _import_docx_document()
    doc = Document()
    doc.add_paragraph('{{ notice_date_display }}')
    doc.add_paragraph('HON. {{ lce_name_upper }}')
    doc.add_paragraph('{{ lce_position }}')
    doc.add_paragraph('{{ lgu_display }}')
    doc.add_paragraph('Subject: {{ subject_line }}')
    doc.add_paragraph('Dear Hon. {{ lce_surname }};')
    return doc


def _add_notice_project_table(doc):
    table = doc.add_table(rows=2, cols=4)
    headers = ['PROJECT(S)', 'FINDING(S)', 'RECOMMENDATION(S)', 'COMPLIANCE PERIOD']
    for index, label in enumerate(headers):
        table.cell(0, index).text = label
    table.cell(1, 0).text = '{{ project_names_text }}'
    table.cell(1, 1).text = '{{ findings_text }}'
    table.cell(1, 2).text = '{{ recommendations_text }}'
    table.cell(1, 3).text = '{{ compliance_period_text }}'
    return table


def _replace_legal_findings(notice, findings_payload):
    if findings_payload is None:
        return

    for finding in notice.legal_findings.all():
        db.session.delete(finding)

    for item in findings_payload:
        payload = item or {}
        finding = LegalFinding(
            notice_id=notice.id,
            informal_notes=(payload.get('informal_notes') or '').strip() or None,
            formal_findings=(payload.get('formal_findings') or '').strip() or None,
            recommendations=(payload.get('recommendations') or '').strip() or None,
            agreements_timelines=(payload.get('agreements_timelines') or '').strip() or None,
            ai_model_used=(payload.get('ai_model_used') or '').strip() or None,
            ai_prompt_used=(payload.get('ai_prompt_used') or '').strip() or None,
            is_edited=True,
        )
        db.session.add(finding)


def _default_legal_findings(requirement, deadline):
    if not requirement:
        return []

    timeline_days = deadline.sla_working_days if deadline and deadline.sla_working_days else requirement.sla_days or 15
    return [{
        'formal_findings': f'The required submission for "{requirement.requirement_name}" remains outstanding or non-compliant.',
        'recommendations': f'Submit or rectify the requirements for "{requirement.requirement_name}" and provide documentary proof of compliance.',
        'agreements_timelines': f'Comply within {timeline_days} working days from receipt of this notice.',
        'informal_notes': 'Auto-generated baseline finding for review and editing.',
    }]


def _context_findings(notice, parent_notice, notice_chain):
    findings = notice.legal_findings.order_by(LegalFinding.id.asc()).all()
    if findings:
        return findings

    for candidate in (parent_notice, notice_chain.get('annex_e')):
        if not candidate:
            continue
        inherited = candidate.legal_findings.order_by(LegalFinding.id.asc()).all()
        if inherited:
            return inherited
    return []


def _notice_chain_date_display(notice_chain, parent_notice, notice_type):
    notice = notice_chain.get(notice_type)
    if not notice and parent_notice and parent_notice.notice_type == notice_type:
        notice = parent_notice
    return _format_notice_date(notice.generated_at) if notice and notice.generated_at else ''


def _prior_nnc_dates_display(notice_chain, parent_notice):
    dates = [
        _notice_chain_date_display(notice_chain, parent_notice, 'annex_m'),
        _notice_chain_date_display(notice_chain, parent_notice, 'annex_n'),
    ]
    return ' and '.join(item for item in dates if item)


def _serialize_legal_finding(finding):
    if isinstance(finding, dict):
        return {
            'id': finding.get('id'),
            'informal_notes': finding.get('informal_notes') or '',
            'formal_findings': finding.get('formal_findings') or '',
            'recommendations': finding.get('recommendations') or '',
            'agreements_timelines': finding.get('agreements_timelines') or '',
            'remarks': finding.get('remarks') or finding.get('informal_notes') or '',
            'ai_model_used': finding.get('ai_model_used'),
            'ai_prompt_used': finding.get('ai_prompt_used'),
            'is_edited': finding.get('is_edited', False),
        }
    return {
        'id': finding.id,
        'informal_notes': finding.informal_notes or '',
        'formal_findings': finding.formal_findings or '',
        'recommendations': finding.recommendations or '',
        'agreements_timelines': finding.agreements_timelines or '',
        'remarks': finding.informal_notes or '',
        'ai_model_used': finding.ai_model_used,
        'ai_prompt_used': finding.ai_prompt_used,
        'is_edited': finding.is_edited,
    }


def _serialize_notice_snapshot(notice):
    return {
        'id': notice.id,
        'project_id': notice.project_id,
        'notice_type': notice.notice_type,
        'notice_number': notice.notice_number,
        'title': notice.title,
        'body_text': notice.body_text,
        'status': notice.status,
        'trigger_sla_id': notice.trigger_sla_id,
        'trigger_project_requirement_id': notice.trigger_project_requirement_id,
        'generated_at': notice.generated_at.isoformat() if notice.generated_at else None,
        'sent_at': notice.sent_at.isoformat() if notice.sent_at else None,
        'acknowledged_at': notice.acknowledged_at.isoformat() if notice.acknowledged_at else None,
        'final_text': notice.final_text,
    }


def _get_project_with_access(project_id, user_id=None):
    try:
        if user_id is not None:
            return access_service.ensure_project_access(project_id, user_id)
    except access_service.AccessError as exc:
        raise NoticeError(exc.message, exc.status_code)

    project = db.session.get(Project, project_id)
    if not project:
        raise NoticeError('Project not found', 404)
    return project


def _get_deadline_with_access(sla_id, user_id=None):
    deadline = db.session.get(SlaDeadline, sla_id)
    if not deadline:
        raise NoticeError('SLA deadline not found', 404)
    _get_project_with_access(deadline.project_id, user_id=user_id)
    return deadline


def _get_requirement_with_access(requirement_id, user_id=None):
    try:
        if user_id is not None:
            return access_service.ensure_project_requirement_access(requirement_id, user_id)
    except access_service.AccessError as exc:
        raise NoticeError(exc.message, exc.status_code)

    requirement = db.session.get(ProjectRequirement, requirement_id)
    if not requirement:
        raise NoticeError('Project requirement not found', 404)
    return requirement


def _get_notice_with_access(notice_id, user_id=None):
    notice = db.session.get(EscalationNotice, notice_id)
    if not notice:
        raise NoticeError('Escalation notice not found', 404)
    _get_project_with_access(notice.project_id, user_id=user_id)
    return notice


def _build_recipient(notice_id, lgu):
    if not lgu:
        return None

    recipient_name = lgu.mayor_name or lgu.contact_person or lgu.name
    recipient_role = 'Mayor' if (lgu.mayor_name or lgu.mayor_email) else 'LGU Contact'
    return NoticeRecipient(
        notice_id=notice_id,
        lgu_id=lgu.id,
        recipient_name=recipient_name,
        recipient_email=lgu.mayor_email or lgu.contact_email,
        recipient_role=recipient_role,
        delivery_method='pending_dispatch',
        delivery_status='pending',
    )


def _get_notice_chain_lookup(notice):
    notices = EscalationNotice.query.filter(
        EscalationNotice.project_id == notice.project_id,
        EscalationNotice.trigger_sla_id == notice.trigger_sla_id,
        EscalationNotice.notice_type.in_(VALID_NOTICE_TYPES),
    ).all()
    return {item.notice_type: item for item in notices}


def _next_notice_number(project_id, notice_type):
    current = db.session.query(db.func.max(EscalationNotice.notice_number)).filter(
        EscalationNotice.project_id == project_id,
        EscalationNotice.notice_type == notice_type,
    ).scalar()
    return (current or 0) + 1


def _coerce_optional_user_id(user_id):
    if user_id in (None, ''):
        return None
    try:
        return int(user_id)
    except (TypeError, ValueError):
        raise NoticeError('Invalid user identity', 422)


def _parse_optional_datetime(value, field_name='acknowledged_at'):
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.replace('Z', '+00:00'))
        except ValueError as exc:
            raise NoticeError(f'{field_name} must be an ISO timestamp.', 422) from exc
    else:
        raise NoticeError(f'{field_name} must be an ISO timestamp.', 422)
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _to_aware_datetime(value):
    if value is None:
        return None
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def _clean_text(value):
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def _default_release_recipient_name(notice):
    recipient = notice.recipients.order_by(NoticeRecipient.id.asc()).first()
    if recipient:
        return recipient.recipient_name
    lgu = notice.project.lgu if notice.project else None
    if not lgu:
        return None
    return lgu.contact_person or lgu.mayor_name or lgu.name


def _default_release_recipient_email(notice, channel):
    if channel != 'email':
        return None
    recipient = notice.recipients.order_by(NoticeRecipient.id.asc()).first()
    if recipient and recipient.recipient_email:
        return recipient.recipient_email
    lgu = notice.project.lgu if notice.project else None
    if not lgu:
        return None
    return lgu.contact_email or lgu.mayor_email


def _serialize_communication(communication):
    return {
        'id': communication.id,
        'communication_type': communication.communication_type,
        'channel': communication.channel,
        'subject': communication.subject,
        'body': communication.body,
        'recipient_email': communication.recipient_email,
        'recipient_name': communication.recipient_name,
        'recipient_destination': communication.recipient_destination,
        'recipient_office': communication.recipient_office,
        'reference_no': communication.reference_no,
        'related_notice_id': communication.related_notice_id,
        'delivery_status': communication.delivery_status,
        'sent_at': communication.sent_at.isoformat() if communication.sent_at else None,
        'posted_at': communication.posted_at.isoformat() if communication.posted_at else None,
        'remarks': communication.remarks,
        'created_at': communication.created_at.isoformat() if communication.created_at else None,
    }


def _format_notice_date(value):
    if not value:
        return ''
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.strftime('%B %d, %Y')


def _build_filename_base(notice, context):
    prefix = NOTICE_FILENAME_PREFIXES.get(notice.notice_type, notice.notice_type)
    slug = _slugify(context.get('project_name') or 'project')
    return f'{prefix}-{slug}-{notice.id}'


def _slugify(value):
    cleaned = ''.join(char.lower() if char.isalnum() else '-' for char in value)
    while '--' in cleaned:
        cleaned = cleaned.replace('--', '-')
    return cleaned.strip('-') or 'notice'


def _split_paragraphs(text):
    if not text:
        return []
    return [part.strip() for part in text.split('\n\n') if part.strip()]


def _build_fallback_pdf_lines(context):
    sections = [
        context.get('notice_label'),
        context.get('notice_date_display'),
        f"To: {context.get('lce_name_upper') or context.get('lce_name') or 'LOCAL CHIEF EXECUTIVE'}",
        context.get('lce_position'),
        context.get('lgu_display'),
        f"Subject: {context.get('subject_line') or context.get('notice_title')}",
        '',
    ]
    for paragraph in context.get('default_body_paragraphs') or []:
        sections.extend(textwrap.wrap(paragraph, width=88) or [''])
        sections.append('')
    if context.get('project_rows_text'):
        sections.append('Projects / Findings:')
        sections.extend(textwrap.wrap(context['project_rows_text'], width=88))
        sections.append('')
    if context.get('custom_body_paragraphs'):
        sections.append('Supplemental Notes:')
        for paragraph in context['custom_body_paragraphs']:
            sections.extend(textwrap.wrap(paragraph, width=88) or [''])
            sections.append('')
    if context.get('findings_rows'):
        sections.append('Findings:')
        for item in context['findings_rows']:
            finding_text = item.get('formal_findings') or item.get('informal_notes') or ''
            sections.extend(textwrap.wrap(f'- {finding_text}', width=88) or ['-'])
        sections.append('')
    sections.append('Very truly yours,')
    sections.append(context.get('signatory_name') or '')
    sections.append(context.get('signatory_title') or '')

    normalized = []
    for line in sections:
        text = (line or '').replace('–', '-').replace('—', '-').replace('\r', ' ').replace('\n', ' ')
        normalized.append(text[:120])
    return normalized[:45]


def _default_lce_position(lgu):
    if not lgu:
        return 'Local Chief Executive'
    locality = f'{lgu.name or ""} {lgu.municipality or ""}'.casefold()
    if 'province of' in locality:
        return 'Provincial Governor'
    if 'city' in locality:
        return 'City Mayor'
    return 'Municipal Mayor'


def _default_lgu_office_heading(lgu):
    position = _default_lce_position(lgu)
    if position == 'City Mayor':
        return 'OFFICE OF THE CITY MAYOR'
    if position == 'Provincial Governor':
        return 'OFFICE OF THE PROVINCIAL GOVERNOR'
    return 'OFFICE OF THE MUNICIPAL MAYOR'


def _build_location_line(lgu, region_name=None):
    if not lgu:
        return ', '.join(item for item in ['Laguna', region_name or 'Philippines'] if item)
    return _join_unique_location_parts(lgu.name, lgu.province, region_name)


def _default_signatory_name(notice_type):
    if notice_type == 'annex_o':
        return _default_regional_director_name()
    return _default_provincial_director_name()


def _default_signatory_title(notice_type, province):
    if notice_type == 'annex_o':
        return current_app.config.get('DILG_REGIONAL_DIRECTOR_TITLE') or 'Regional Director'
    return f'Provincial Director, {province}'


def _default_provincial_director_name():
    return current_app.config.get('DILG_PROVINCIAL_DIRECTOR_NAME') or 'DILG Provincial Director'


def _default_regional_director_name():
    return current_app.config.get('DILG_REGIONAL_DIRECTOR_NAME') or 'Regional Director'


def _join_unique_location_parts(*parts):
    values = []
    seen = set()
    for part in parts:
        cleaned = str(part or '').strip()
        if not cleaned:
            continue
        key = cleaned.casefold()
        if key in seen:
            continue
        seen.add(key)
        values.append(cleaned)
    return ', '.join(values)


def _subject_line(notice):
    if notice.notice_type == 'annex_e':
        return 'Notice of Deficiencies'
    if notice.notice_type == 'annex_m':
        return 'Notice of Non-Compliance No. 1'
    if notice.notice_type == 'annex_n':
        return 'Notice of Non-Compliance No. 2'
    if notice.notice_type == 'annex_o':
        return 'Notice of Non-Compliance No. 3'
    return notice.title


def _extract_surname(name):
    parts = [part for part in (name or '').replace(',', ' ').split() if part]
    if not parts:
        return 'Sir / Madam'
    return parts[-1].rstrip('.')


def _build_project_row(project, requirement, deadline, findings):
    serialized_findings = [_serialize_legal_finding(finding) for finding in findings]
    finding_summary = _join_context_values(
        [item['formal_findings'] for item in serialized_findings if item['formal_findings']],
        [f'Outstanding compliance issue on {requirement.requirement_name}' if requirement else 'Outstanding compliance issue'],
    )
    recommendation_summary = _join_context_values(
        [item['recommendations'] for item in serialized_findings if item['recommendations']],
        [
            f'Submit or rectify the requirements for "{requirement.requirement_name}" and provide documentary proof of compliance.'
            if requirement else 'Submit corrective action and documentary proof of compliance.'
        ],
    )
    compliance_period = (
        _format_notice_date(deadline.deadline_date)
        if deadline and deadline.deadline_date else
        _join_context_values(
            [item['agreements_timelines'] for item in serialized_findings if item['agreements_timelines']],
            [],
        )
    )
    return {
        'project_id': project.id if project else None,
        'project_name': project.project_name if project else 'Project',
        'program_title': project.program.name if project and project.program else 'Program',
        'requirement_name': requirement.requirement_name if requirement else 'General compliance requirement',
        'budget_allocation_display': _format_money_amount(project.budget_amount) if project else '',
        'finding_summary': finding_summary,
        'recommendation_summary': recommendation_summary,
        'deadline_date_display': _format_notice_date(deadline.deadline_date) if deadline and deadline.deadline_date else '',
        'compliance_period_display': compliance_period,
    }


def _format_project_row_text(item):
    return (
        f"{item['project_name']} | {item['program_title']} | "
        f"{item.get('budget_allocation_display') or ''} | "
        f"{item['requirement_name']} | {item['finding_summary']} | "
        f"{item.get('recommendation_summary') or ''} | {item.get('compliance_period_display') or ''}"
    )


def _build_notice_project_rows(notice, project, requirement, deadline, findings):
    rows = [_build_project_row(project, requirement, deadline, findings)]
    for related_notice in _related_program_notice_rows(notice, project):
        if related_notice.id == notice.id:
            continue
        related_project = db.session.get(Project, related_notice.project_id)
        related_requirement = (
            db.session.get(ProjectRequirement, related_notice.trigger_project_requirement_id)
            if related_notice.trigger_project_requirement_id else None
        )
        related_deadline = (
            db.session.get(SlaDeadline, related_notice.trigger_sla_id)
            if related_notice.trigger_sla_id else None
        )
        related_parent = (
            db.session.get(EscalationNotice, related_notice.parent_notice_id)
            if related_notice.parent_notice_id else None
        )
        related_chain = _get_notice_chain_lookup(related_notice)
        related_findings = _context_findings(related_notice, related_parent, related_chain)
        rows.append(_build_project_row(related_project, related_requirement, related_deadline, related_findings))

    return _combine_project_rows(rows)


def _related_program_notice_rows(notice, project):
    if not project or notice.notice_type == 'annex_e' or not notice.generated_at:
        return []

    notice_day = _to_aware_datetime(notice.generated_at).date()
    candidates = (
        EscalationNotice.query
        .join(Project, EscalationNotice.project_id == Project.id)
        .filter(
            EscalationNotice.notice_type == notice.notice_type,
            EscalationNotice.status != 'expired',
            Project.lgu_id == project.lgu_id,
            Project.program_id == project.program_id,
            Project.fiscal_year == project.fiscal_year,
        )
        .order_by(Project.project_name.asc(), EscalationNotice.id.asc())
        .all()
    )
    return [
        item for item in candidates
        if item.generated_at and _to_aware_datetime(item.generated_at).date() == notice_day
    ]


def _combine_project_rows(rows):
    combined = []
    by_project = {}
    for row in rows:
        key = row.get('project_id') or row.get('project_name')
        if key not in by_project:
            by_project[key] = dict(row)
            combined.append(by_project[key])
            continue
        current = by_project[key]
        current['requirement_name'] = _join_unique_lines(current.get('requirement_name'), row.get('requirement_name'))
        current['finding_summary'] = _join_unique_lines(current.get('finding_summary'), row.get('finding_summary'))
        current['recommendation_summary'] = _join_unique_lines(
            current.get('recommendation_summary'),
            row.get('recommendation_summary'),
        )
        current['budget_allocation_display'] = _join_unique_lines(
            current.get('budget_allocation_display'),
            row.get('budget_allocation_display'),
        )
        current['compliance_period_display'] = _join_unique_lines(
            current.get('compliance_period_display'),
            row.get('compliance_period_display'),
        )
    return combined


def _join_unique_lines(*values):
    seen = set()
    lines = []
    for value in values:
        for line in str(value or '').splitlines():
            cleaned = line.strip()
            if not cleaned or cleaned in seen:
                continue
            seen.add(cleaned)
            lines.append(cleaned)
    return '\n'.join(lines)


def _join_context_values(primary, fallback):
    primary_text = _join_unique_lines(*(primary or []))
    if primary_text:
        return primary_text
    return _join_unique_lines(*(fallback or []))


def _format_money_amount(value):
    if value is None:
        return ''
    try:
        return f'PHP {value:,.2f}'
    except (TypeError, ValueError):
        return str(value)


def _format_recipient_line(recipient):
    if recipient.recipient_email:
        return f'{recipient.recipient_name} ({recipient.recipient_email})'
    return recipient.recipient_name


def _import_docxtemplate():
    try:
        from docxtpl import DocxTemplate
    except Exception as exc:
        raise NoticeError(f'docxtpl is required for DOCX notice rendering: {exc}', 500)
    return DocxTemplate


def _import_docx_document():
    try:
        from docx import Document
    except Exception as exc:
        raise NoticeError(f'python-docx is required for DOCX notice rendering: {exc}', 500)
    return Document


def _import_weasyprint_html():
    try:
        from weasyprint import HTML
    except Exception as exc:
        raise NoticeError(f'WeasyPrint is required for PDF notice rendering: {exc}', 500)
    return HTML


def _escape_pdf_text(value):
    return value.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')


def _build_pdf_document(content_stream):
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
        b'<< /Length %d >>\nstream\n%s\nendstream' % (len(content_stream), content_stream),
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]

    pdf = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f'{index} 0 obj\n'.encode('ascii'))
        pdf.extend(obj)
        pdf.extend(b'\nendobj\n')

    xref_start = len(pdf)
    pdf.extend(f'xref\n0 {len(objects) + 1}\n'.encode('ascii'))
    pdf.extend(b'0000000000 65535 f \n')
    for offset in offsets[1:]:
        pdf.extend(f'{offset:010d} 00000 n \n'.encode('ascii'))
    pdf.extend(
        f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF'.encode('ascii')
    )
    return bytes(pdf)
